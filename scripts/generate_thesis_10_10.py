"""Build the final graduation thesis from the verified project evidence.

Design system: narrative_proposal with the university_thesis_arabic override:
A4, 3 cm right margin, 2.54 cm remaining margins, Simplified Arabic 13.5 pt,
Times New Roman 12 pt for English, and 1.15 line spacing.
"""

from __future__ import annotations

import re
import json
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "THESIS_GRADUATION_CURRENT.md"
CATALOG_SOURCE = ROOT / "docs" / "THESIS_REQUIREMENTS_CATALOG_CURRENT.md"
TRUTH_REGISTER_SOURCE = ROOT / "docs" / "THESIS_CAPABILITY_TRUTH_REGISTER_CURRENT.md"
OUT_DIR = ROOT / "docs" / "output"
ASSET_DIR = ROOT / "tmp" / "thesis_assets"
CURATED_DIAGRAM_DIR = ROOT / "docs" / "thesis_diagrams" / "rendered"
DRAWIO_DIAGRAM_DIR = ROOT / "docs" / "thesis_diagrams" / "drawio_rendered"
USECASE_DRAWIO_DIR = ROOT / "docs" / "thesis_diagrams" / "usecase_drawio_rendered"
USECASE_ANALYSIS_SOURCE = ROOT / "docs" / "USE_CASES_ANALYSIS.json"
DATABASE_INVENTORY_SOURCE = ROOT / "docs" / "DATABASE_SCHEMA_INVENTORY.json"
DATABASE_DIAGRAM_DIR = ROOT / "docs" / "thesis_diagrams" / "rendered"
DATABASE_FULL_ERD_DIR = ROOT / "docs" / "thesis_diagrams" / "database_erdds" / "rendered"
DATABASE_FULL_ERD_MANIFEST = ROOT / "docs" / "thesis_diagrams" / "database_erdds" / "manifest.json"
DATABASE_COMPACT_ERD_DIR = ROOT / "docs" / "thesis_diagrams" / "database_erdds" / "compact_rendered"
OUTPUT = OUT_DIR / "HRManagement_Graduation_Thesis_2026_Current.docx"

ARABIC_FONT = "Simplified Arabic"
ENGLISH_FONT = "Times New Roman"
CODE_FONT = "Consolas"
ARIAL_PATH = Path(r"C:\Windows\Fonts\arial.ttf")
ARIAL_BOLD_PATH = Path(r"C:\Windows\Fonts\arialbd.ttf")

INK = "162B3D"
TEAL = "0F6B73"
ACCENT = "DCEEF0"
LIGHT = "F4F7F8"
GRAY = "5E6A71"
TABLE_HEADER = "E8EEF5"
RISK = "FCE8E6"
SUCCESS = "E3F2E8"
TABLE_COUNTERS = {1: 0, 2: 0, 3: 0}
FIGURE_COUNTERS: dict[str, int] = {}


@dataclass(frozen=True)
class Figure:
    label: str
    title: str
    filename: str
    chapter: int


FIGURES = {
    "institution": Figure("الشكل (1-1)", "نموذج العمل بين الإدارة المركزية والإدارات الفرعية", "01_institution.png", 1),
    "generic_architecture": Figure("الشكل (2-1)", "المعمارية العامة للخدمات المصغرة", "02_generic_architecture.png", 2),
    "layers": Figure("الشكل (2-2)", "طبقات الخدمة وفق فصل المسؤوليات", "03_layers.png", 2),
    "events": Figure("الشكل (2-3)", "نمط التكامل القائم على الأحداث", "04_events.png", 2),
    "solution": Figure("الشكل (3-1)", "البنية المعمارية للحل البرمجي HRManagement", "05_solution.png", 3),
    "data": Figure("الشكل (3-2)", "ملكية البيانات وعزل قواعدها بين المجالات", "06_data_ownership.png", 3),
    "scope": Figure("الشكل (3-3)", "نطاق الصلاحيات التنظيمي للمركز والفروع", "07_scope.png", 3),
    "lifecycle": Figure("الشكل (3-4)", "مسار التعيين والتكامل بين الخدمات", "08_lifecycle.png", 3),
    "deployment": Figure("الشكل (3-5)", "نموذج النشر والتشغيل المحلي والإنتاجي", "09_deployment.png", 3),
    "tests": Figure("الشكل (3-6)", "ملخص دليل التحقق البرمجي", "10_tests.png", 3),
    "usecase": Figure("الشكل (1-2)", "مخطط حالات الاستخدام العام للنظام", "11_usecase.png", 1),
    "dfd": Figure("الشكل (1-3)", "مخطط تدفق البيانات بين المركز والفروع", "12_dfd.png", 1),
    "activity_hire": Figure("الشكل (3-7)", "مخطط النشاط لمسار التعيين", "13_activity_hire.png", 3),
    "sequence_hire": Figure("الشكل (3-8)", "مخطط التسلسل لمسار التعيين والأحداث", "14_sequence_hire.png", 3),
    "domain": Figure("الشكل (3-9)", "مخطط نموذج المجال في النواة", "15_domain.png", 3),
    "rbac": Figure("الشكل (3-10)", "مصفوفة الصلاحيات ونطاق المركز والفروع", "16_rbac.png", 3),
    "trace": Figure("الشكل (3-11)", "مصفوفة تتبع المتطلبات والأدلة", "17_trace.png", 3),
    "usecase_org": Figure("الشكل (1-4)", "حالات استخدام الهيكل التنظيمي ونطاق الفرع", "18_usecase_org.png", 1),
    "usecase_lifecycle": Figure("الشكل (1-5)", "حالات استخدام دورة حياة الموظف والعقود", "19_usecase_lifecycle.png", 1),
    "usecase_payroll": Figure("الشكل (1-6)", "حالات استخدام الرواتب والآثار المالية", "20_usecase_payroll.png", 1),
    "usecase_people": Figure("الشكل (1-7)", "حالات استخدام الصحة والتدريب والحركات", "21_usecase_people.png", 1),
    "usecase_portals": Figure("الشكل (1-8)", "حالات استخدام البوابة العامة والخدمة الذاتية", "22_usecase_portals.png", 1),
    "usecase_official_01": Figure("الشكل (1-9)", "حالات استخدام الكادر الإشرافي: الخدمة والعقود", "79_usecase_official_01.png", 1),
    "usecase_official_02": Figure("الشكل (1-10)", "حالات استخدام الكادر الإشرافي: الشؤون الشخصية والصحية", "80_usecase_official_02.png", 1),
    "usecase_official_03": Figure("الشكل (1-11)", "حالات استخدام الكادر الإشرافي: التنقلات والتوزيع", "81_usecase_official_03.png", 1),
    "usecase_official_04": Figure("الشكل (1-12)", "حالات استخدام الكادر الإشرافي: تخطيط الدورات وتأهيل القوى البشرية", "82_usecase_official_04.png", 1),
    "usecase_official_05": Figure("الشكل (1-13)", "حالات استخدام الموظفين التنفيذيين: الخدمة والعقود", "83_usecase_official_05.png", 1),
    "usecase_official_06": Figure("الشكل (1-14)", "حالات استخدام الموظفين التنفيذيين: الشؤون الشخصية والصحية", "84_usecase_official_06.png", 1),
    "usecase_official_07": Figure("الشكل (1-15)", "حالات استخدام الموظفين التنفيذيين: التنقلات والتوزيع", "85_usecase_official_07.png", 1),
    "usecase_official_08": Figure("الشكل (1-16)", "حالات استخدام الموظفين التنفيذيين: تخطيط الدورات وتأهيل القوى البشرية", "86_usecase_official_08.png", 1),
    "usecase_official_09": Figure("الشكل (1-17)", "حالات استخدام السجلات المؤسسية: السيرة الوظيفية", "87_usecase_official_09.png", 1),
    "usecase_official_10": Figure("الشكل (1-18)", "حالات استخدام السجلات المؤسسية: الرواتب والتعويضات", "88_usecase_official_10.png", 1),
    "usecase_official_11": Figure("الشكل (1-19)", "حالات استخدام السجلات المؤسسية: أرشفة الملفات والوثائق", "89_usecase_official_11.png", 1),
    "usecase_official_12": Figure("الشكل (1-20)", "حالات استخدام الفروع المركزية: بطاقات التعريف والهويات", "90_usecase_official_12.png", 1),
    "usecase_official_13": Figure("الشكل (1-21)", "حالات استخدام الفروع المركزية: الرقابة", "91_usecase_official_13.png", 1),
    "usecase_official_14": Figure("الشكل (1-22)", "حالات استخدام الفروع المركزية: التنظيم والتجهيز", "92_usecase_official_14.png", 1),
    "usecase_official_15": Figure("الشكل (1-23)", "حالات استخدام الفروع المركزية: التخطيط والإحصاء", "93_usecase_official_15.png", 1),
    "usecase_official_16": Figure("الشكل (1-24)", "حالات استخدام الفروع المركزية: الفرع الإداري", "94_usecase_official_16.png", 1),
    "usecase_official_17": Figure("الشكل (1-25)", "حالات استخدام الفروع المركزية: الشؤون القانونية والانضباط", "95_usecase_official_17.png", 1),
    "usecase_official_18": Figure("الشكل (1-26)", "حالات استخدام الفروع المركزية: الديوان العام", "96_usecase_official_18.png", 1),
    "usecase_official_19": Figure("الشكل (1-27)", "حالات استخدام الفروع المركزية: البعثات", "97_usecase_official_19.png", 1),
    "usecase_official_20": Figure("الشكل (1-28)", "حالات استخدام الفروع المركزية: تقنية المعلومات", "98_usecase_official_20.png", 1),
    "usecase_official_21": Figure("الشكل (1-29)", "حالات استخدام الفروع المركزية: التطوير البرمجي وتحليل المعلومات", "99_usecase_official_21.png", 1),
    "dfd_main": Figure("الشكل (1-30)", "مخطط تدفق البيانات DFD للمستوى الأول: العمليات الرئيسية ومخازن البيانات", "38_dfd_main_level1.png", 1),
    "dfd_new_employee": Figure("الشكل (1-31)", "مخطط تدفق البيانات للمستوى الثاني: موظف جديد", "39_dfd_new_employee_level2.png", 1),
    "dfd_transfer": Figure("الشكل (1-32)", "مخطط تدفق البيانات للمستوى الثاني: نقل موظف", "40_dfd_transfer_level2.png", 1),
    "dfd_medical": Figure("الشكل (1-33)", "مخطط تدفق البيانات للمستوى الثاني: الملف الطبي", "41_dfd_medical_file_level2.png", 1),
    "dfd_training": Figure("الشكل (1-34)", "مخطط تدفق البيانات للمستوى الثاني: الدورة أو البعثة", "42_dfd_training_scholarship_level2.png", 1),
    "dfd_payroll": Figure("الشكل (1-35)", "مخطط تدفق البيانات للمستوى الثاني: الرواتب والتعويضات", "43_dfd_payroll_compensation_level2.png", 1),
    "dfd_complaint": Figure("الشكل (1-36)", "مخطط تدفق البيانات للمستوى الثاني: الشكوى أو المخالفة", "44_dfd_complaint_violation_level2.png", 1),
    "dfd_separation": Figure("الشكل (1-37)", "مخطط تدفق البيانات للمستوى الثاني: إنهاء الخدمة", "45_dfd_separation_level2.png", 1),
    "dfd_dictionary": Figure("الشكل (1-38)", "قاموس بيانات DFD: توحيد أسماء التدفقات والمخازن", "46_dfd_data_dictionary.png", 1),
    "activity_contract": Figure("الشكل (3-12)", "مخطط النشاط لتفعيل العقد", "23_activity_contract.png", 3),
    "activity_leave": Figure("الشكل (3-13)", "مخطط النشاط للإجازة والسلفة", "24_activity_leave.png", 3),
    "activity_movement": Figure("الشكل (3-14)", "مخطط النشاط للنقل والانتداب", "25_activity_movement.png", 3),
    "activity_separation": Figure("الشكل (3-15)", "مخطط النشاط للفصل وإنهاء الخدمة", "26_activity_separation.png", 3),
    "activity_payroll": Figure("الشكل (3-16)", "مخطط النشاط لدورة الرواتب", "27_activity_payroll.png", 3),
    "sequence_contract": Figure("الشكل (3-17)", "مخطط التسلسل للإجازة والسلفة", "28_sequence_leave.png", 3),
    "sequence_movement": Figure("الشكل (3-18)", "مخطط التسلسل لتنفيذ النقل", "29_sequence_movement.png", 3),
    "sequence_separation": Figure("الشكل (3-19)", "مخطط التسلسل للفصل وآثاره", "30_sequence_separation.png", 3),
    "sequence_payroll": Figure("الشكل (3-20)", "مخطط التسلسل لترحيل دورة الرواتب", "31_sequence_payroll.png", 3),
    "erd": Figure("الشكل (3-21)", "ERD منطقي لخدمات النواة والعقود والرواتب", "32_erd.png", 3),
    "component": Figure("الشكل (3-22)", "مخطط المكونات والبنية المؤسسية", "33_component.png", 3),
    "deployment_formal": Figure("الشكل (3-23)", "مخطط النشر التشغيلي", "34_deployment_formal.png", 3),
    "security": Figure("الشكل (3-24)", "مخطط الأمن والنطاق التنظيمي", "35_security.png", 3),
    "event_flow": Figure("الشكل (3-25)", "مخطط تدفق الأحداث بين الخدمات", "36_event_flow.png", 3),
    "ownership": Figure("الشكل (3-26)", "مخطط ملكية البيانات بين الخدمات", "37_ownership.png", 3),
    "database_landscape": Figure("الشكل (3-27)", "منظومة قواعد البيانات المملوكة للخدمات", "47_database_landscape.png", 3),
    "database_hrcore": Figure("الشكل (3-28)", "ERD مفاهيمي مختصر لقاعدة HR.Core", "48_database_hrcore.png", 3),
    "database_contracts": Figure("الشكل (3-29)", "ERD مفاهيمي مختصر لقاعدة HR.Contracts", "49_database_contracts.png", 3),
    "database_payroll": Figure("الشكل (3-30)", "ERD مفاهيمي مختصر لقاعدة Payroll", "50_database_payroll.png", 3),
    "database_integration": Figure("الشكل (3-31)", "تواصل قواعد البيانات عبر الأحداث والمعرفات", "51_database_integration.png", 3),
    "database_full_gateway": Figure("الشكل (م-6)", "ERD الكامل لقاعدة Gateway Audit", "52_erd_gateway_audit.png", 3),
    "database_full_admin": Figure("الشكل (م-7)", "ERD الكامل لقاعدة Admin.Logistics", "53_erd_admin_logistics.png", 3),
    "database_full_auth": Figure("الشكل (م-8)", "ERD الكامل لقاعدة Auth.Directory", "54_erd_auth_directory.png", 3),
    "database_full_compliance": Figure("الشكل (م-9)", "ERD الكامل لقاعدة Compliance", "55_erd_compliance.png", 3),
    "database_full_analytics": Figure("الشكل (م-10)", "ERD الكامل لقاعدة Dev.Analytics", "56_erd_analytics.png", 3),
    "database_full_files": Figure("الشكل (م-11)", "ERD الكامل لقاعدة File.Storage", "57_erd_file_storage.png", 3),
    "database_full_contracts": Figure("الشكل (م-12)", "ERD الكامل لقاعدة HR.Contracts", "58_erd_contracts.png", 3),
    "database_full_core": Figure("الشكل (م-13)", "ERD الكامل لقاعدة HR.Core", "59_erd_hr_core.png", 3),
    "database_full_health": Figure("الشكل (م-14)", "ERD الكامل لقاعدة HR.Health", "60_erd_health.png", 3),
    "database_full_movements": Figure("الشكل (م-15)", "ERD الكامل لقاعدة HR.Movements", "61_erd_movements.png", 3),
    "database_full_training": Figure("الشكل (م-16)", "ERD الكامل لقاعدة HR.Training", "62_erd_training.png", 3),
    "database_full_idcard": Figure("الشكل (م-17)", "ERD الكامل لقاعدة IDCard", "63_erd_idcard.png", 3),
    "database_full_it": Figure("الشكل (م-18)", "ERD الكامل لقاعدة IT.Infra", "64_erd_it_infra.png", 3),
    "database_full_payroll": Figure("الشكل (م-19)", "ERD الكامل لقاعدة Payroll", "65_erd_payroll.png", 3),
    "database_full_archive": Figure("الشكل (م-20)", "ERD الكامل لقاعدة Records.Archive", "66_erd_archive.png", 3),
    "database_compact_gateway": Figure("الشكل (3-32)", "ERD مختصر لقاعدة Gateway Audit", "67_erd_compact_gateway.png", 3),
    "database_compact_admin": Figure("الشكل (3-33)", "ERD مختصر لقاعدة Admin.Logistics", "68_erd_compact_admin.png", 3),
    "database_compact_auth": Figure("الشكل (3-34)", "ERD مختصر لقاعدة Auth.Directory", "69_erd_compact_auth.png", 3),
    "database_compact_compliance": Figure("الشكل (3-35)", "ERD مختصر لقاعدة Compliance", "70_erd_compact_compliance.png", 3),
    "database_compact_analytics": Figure("الشكل (3-36)", "ERD مختصر لقاعدة Dev.Analytics", "71_erd_compact_analytics.png", 3),
    "database_compact_files": Figure("الشكل (3-37)", "ERD مختصر لقاعدة File.Storage", "72_erd_compact_files.png", 3),
    "database_compact_health": Figure("الشكل (3-38)", "ERD مختصر لقاعدة HR.Health", "73_erd_compact_health.png", 3),
    "database_compact_movements": Figure("الشكل (3-39)", "ERD مختصر لقاعدة HR.Movements", "74_erd_compact_movements.png", 3),
    "database_compact_training": Figure("الشكل (3-40)", "ERD مختصر لقاعدة HR.Training", "75_erd_compact_training.png", 3),
    "database_compact_idcard": Figure("الشكل (3-41)", "ERD مختصر لقاعدة IDCard", "76_erd_compact_idcard.png", 3),
    "database_compact_it": Figure("الشكل (3-42)", "ERD مختصر لقاعدة IT.Infra", "77_erd_compact_it.png", 3),
    "database_compact_archive": Figure("الشكل (3-43)", "ERD مختصر لقاعدة Records.Archive", "78_erd_compact_archive.png", 3),
    "tenant_context": Figure("الشكل (3-44)", "سياق المنصة متعددة الجهات والإدارات", "100_tenant_context.png", 3),
    "usecase_platform": Figure("الشكل (3-45)", "حالات استخدام حوكمة المنصة والجهة", "101_usecase_platform.png", 3),
    "dfd_governance": Figure("الشكل (3-46)", "DFD لحوكمة الجهة والسياسات والرموز", "102_dfd_governance.png", 3),
    "activity_onboarding": Figure("الشكل (3-47)", "مخطط نشاط تهيئة جهة جديدة", "103_activity_onboarding.png", 3),
    "activity_policy": Figure("الشكل (3-48)", "مخطط نشاط دورة حياة السياسة", "104_activity_policy.png", 3),
    "sequence_policy": Figure("الشكل (3-49)", "مخطط تسلسل إنفاذ سياسة حية", "105_sequence_policy.png", 3),
    "sequence_code": Figure("الشكل (3-50)", "مخطط تسلسل توليد رمز مؤسسي", "106_sequence_code.png", 3),
    "activity_asset": Figure("الشكل (3-51)", "مخطط نشاط عهدة الأصل والفقد والإرجاع", "107_activity_asset.png", 3),
    "sequence_asset": Figure("الشكل (3-52)", "مخطط تسلسل اعتماد عهدة أصل", "108_sequence_asset.png", 3),
    "component_governance": Figure("الشكل (3-53)", "مكونات حوكمة المنصة والخدمات المشتركة", "109_component_governance.png", 3),
    "security_rls": Figure("الشكل (3-54)", "طبقات عزل الجهة والنطاق التنظيمي", "110_security_rls.png", 3),
    "reliability_journeys": Figure("الشكل (3-55)", "الاعتمادية حسب رحلات الأعمال العابرة للخدمات", "111_reliability_journeys.png", 3),
    "usecase_current_master": Figure("الشكل (3-56)", "حالات الاستخدام المرجعية لخط الأساس الحالي", "112_usecase_current_master.png", 3),
    "usecase_tenant_assignment": Figure("الشكل (3-57)", "حالات استخدام إسناد الجهة والحجر والمصالحة", "113_usecase_tenant_assignment.png", 3),
    "usecase_temporal_org": Figure("الشكل (3-58)", "حالات استخدام الهيكل المؤرخ والوظائف الممولة", "114_usecase_temporal_org.png", 3),
    "usecase_policy_workflow": Figure("الشكل (3-59)", "حالات استخدام السياسة وسير العمل والتوقيع والرموز", "115_usecase_policy_workflow.png", 3),
    "usecase_advanced_payroll": Figure("الشكل (3-60)", "حالات استخدام الرواتب والخزينة المتقدمة", "116_usecase_advanced_payroll.png", 3),
    "usecase_logistics_lineage": Figure("الشكل (3-61)", "حالات استخدام التوريد وتتبع الأصل والعهدة", "117_usecase_logistics_lineage.png", 3),
    "usecase_security_files": Figure("الشكل (3-62)", "حالات استخدام الخصوصية والملفات وثقة الخدمات", "118_usecase_security_files.png", 3),
    "dfd_current_level0": Figure("الشكل (3-63)", "DFD المرجعي للمستوى صفر في خط الأساس الحالي", "119_dfd_current_level0.png", 3),
    "dfd_tenant_assignment": Figure("الشكل (3-64)", "DFD لإسناد الجهة والحجر والمصالحة", "120_dfd_tenant_assignment.png", 3),
    "dfd_policy_workflow": Figure("الشكل (3-65)", "DFD للسياسة وسير العمل والرموز", "121_dfd_policy_workflow.png", 3),
    "dfd_advanced_payroll": Figure("الشكل (3-66)", "DFD للرواتب والتمويل والدفع والتسوية", "122_dfd_advanced_payroll.png", 3),
    "dfd_logistics_lineage": Figure("الشكل (3-67)", "DFD للتوريد والاستلام وتتبع الأصل والعهدة", "123_dfd_logistics_lineage.png", 3),
    "dfd_file_security": Figure("الشكل (3-68)", "DFD لحجر الملفات والفحص والمنح والخصوصية", "124_dfd_file_security.png", 3),
    "sequence_payroll_resumable": Figure("الشكل (3-69)", "تسلسل تشغيل الرواتب القابل للاستئناف والصرف", "125_sequence_payroll_resumable.png", 3),
    "activity_tenant_reconciliation": Figure("الشكل (3-70)", "نشاط إسناد بيانات الجهة والحجر والمصالحة", "126_activity_tenant_reconciliation.png", 3),
}


CURATED_DIAGRAM_MAP = {
    "usecase": "02_usecase_master.png",
    "dfd": "01_dfd_context.png",
    "usecase_org": "03_usecase_organization.png",
    "usecase_lifecycle": "04_usecase_hire_contracts.png",
    "usecase_people": "05_usecase_people_development.png",
    "usecase_payroll": "06_usecase_payroll_governance.png",
    "usecase_portals": "07_usecase_ess_public.png",
    "activity_hire": "08_activity_hire.png",
    "activity_contract": "09_activity_contract.png",
    "activity_leave": "10_activity_leave_advance.png",
    "activity_movement": "11_activity_transfer.png",
    "activity_separation": "12_activity_separation.png",
    "activity_payroll": "13_activity_payroll.png",
    "sequence_hire": "14_sequence_hire.png",
    "sequence_contract": "15_sequence_leave.png",
    "sequence_movement": "16_sequence_transfer.png",
    "sequence_separation": "17_sequence_separation.png",
    "sequence_payroll": "18_sequence_payroll.png",
    "domain": "19_domain_hrcore.png",
    "erd": "20_erd_logical.png",
    "component": "21_component_architecture.png",
    "deployment_formal": "22_deployment.png",
    "security": "23_security_scope.png",
    "event_flow": "24_event_flow.png",
    "ownership": "25_data_ownership.png",
    "tenant_context": "31_tenant_context.png",
    "usecase_platform": "32_usecase_platform_governance.png",
    "dfd_governance": "33_dfd_governance_level1.png",
    "activity_onboarding": "34_activity_agency_onboarding.png",
    "activity_policy": "35_activity_policy_lifecycle.png",
    "sequence_policy": "36_sequence_policy_enforcement.png",
    "sequence_code": "37_sequence_code_generation.png",
    "activity_asset": "38_activity_asset_custody.png",
    "sequence_asset": "39_sequence_asset_assignment.png",
    "component_governance": "40_component_platform_governance.png",
    "security_rls": "41_security_tenant_rls.png",
    "reliability_journeys": "42_reliability_business_journeys.png",
    "usecase_current_master": "43_usecase_current_master.png",
    "usecase_tenant_assignment": "44_usecase_tenant_assignment.png",
    "usecase_temporal_org": "45_usecase_temporal_org.png",
    "usecase_policy_workflow": "46_usecase_policy_workflow.png",
    "usecase_advanced_payroll": "47_usecase_advanced_payroll.png",
    "usecase_logistics_lineage": "48_usecase_logistics_lineage.png",
    "usecase_security_files": "49_usecase_security_files.png",
    "dfd_current_level0": "50_dfd_current_level0.png",
    "dfd_tenant_assignment": "51_dfd_tenant_assignment.png",
    "dfd_policy_workflow": "52_dfd_policy_workflow.png",
    "dfd_advanced_payroll": "53_dfd_advanced_payroll.png",
    "dfd_logistics_lineage": "54_dfd_logistics_lineage.png",
    "dfd_file_security": "55_dfd_file_security.png",
    "sequence_payroll_resumable": "56_sequence_payroll_resumable.png",
    "activity_tenant_reconciliation": "57_activity_tenant_reconciliation.png",
}


DRAWIO_DIAGRAM_MAP = {
    "dfd": "01_context_level0.png",
    "dfd_main": "02_main_level1.png",
    "dfd_new_employee": "03_new_employee_level2.png",
    "dfd_transfer": "04_employee_transfer_level2.png",
    "dfd_medical": "05_medical_file_level2.png",
    "dfd_training": "06_training_scholarship_level2.png",
    "dfd_payroll": "07_payroll_compensation_level2.png",
    "dfd_complaint": "08_complaint_violation_level2.png",
    "dfd_separation": "09_separation_level2.png",
    "dfd_dictionary": "10_data_dictionary.png",
}


USECASE_DRAWIO_MAP = {
    f"usecase_official_{index:02d}": f"{index:02d}_usecase.png"
    for index in range(1, 22)
}


DATABASE_DIAGRAM_MAP = {
    "database_landscape": "26_database_landscape.png",
    "database_hrcore": "27_database_hrcore.png",
    "database_contracts": "28_database_contracts.png",
    "database_payroll": "29_database_payroll.png",
    "database_integration": "30_database_integration.png",
}


DATABASE_FULL_ERD_MAP = {
    "database_full_gateway": "01_gatewayaudit.png",
    "database_full_admin": "02_adminlogistics.png",
    "database_full_auth": "03_authdirectory.png",
    "database_full_compliance": "04_compliance.png",
    "database_full_analytics": "05_analytics.png",
    "database_full_files": "06_filestorage.png",
    "database_full_contracts": "07_contracts.png",
    "database_full_core": "08_hrcore.png",
    "database_full_health": "09_health.png",
    "database_full_movements": "10_movements.png",
    "database_full_training": "11_training.png",
    "database_full_idcard": "12_idcard.png",
    "database_full_it": "13_itinfra.png",
    "database_full_payroll": "14_payroll.png",
    "database_full_archive": "15_archive.png",
}


DATABASE_FULL_ERD_KEYS = list(DATABASE_FULL_ERD_MAP)


DATABASE_COMPACT_ERD_MAP = {
    "database_compact_gateway": "01_gatewayaudit_compact.png",
    "database_compact_admin": "02_adminlogistics_compact.png",
    "database_compact_auth": "03_authdirectory_compact.png",
    "database_compact_compliance": "04_compliance_compact.png",
    "database_compact_analytics": "05_analytics_compact.png",
    "database_compact_files": "06_filestorage_compact.png",
    "database_compact_health": "09_health_compact.png",
    "database_compact_movements": "10_movements_compact.png",
    "database_compact_training": "11_training_compact.png",
    "database_compact_idcard": "12_idcard_compact.png",
    "database_compact_it": "13_itinfra_compact.png",
    "database_compact_archive": "15_archive_compact.png",
    "database_contracts": "07_contracts_compact.png",
    "database_hrcore": "08_hrcore_compact.png",
    "database_payroll": "14_payroll_compact.png",
}


def relabel_figures(keys: list[str], prefix: str, chapter: int = 3) -> None:
    for index, key in enumerate(keys, 1):
        current = FIGURES[key]
        FIGURES[key] = Figure(f"الشكل ({prefix}-{index})", current.title, current.filename, chapter)


relabel_figures(["generic_architecture", "layers", "events"], "2", 2)
relabel_figures(
    [
        "solution", "tenant_context", "scope", "usecase", "usecase_platform", "dfd", "dfd_governance",
        "usecase_org", "usecase_lifecycle", "usecase_payroll", "usecase_people", "usecase_portals",
        "activity_hire", "activity_contract", "activity_leave", "activity_movement", "activity_separation",
        "activity_payroll", "activity_onboarding", "activity_policy", "activity_asset", "sequence_hire",
        "sequence_contract", "sequence_movement", "sequence_separation", "sequence_payroll", "sequence_policy",
        "sequence_code", "sequence_asset", "domain", "erd", "component", "component_governance",
        "deployment_formal", "security", "security_rls", "event_flow", "ownership", "reliability_journeys",
        "rbac", "trace", "usecase_current_master", "usecase_tenant_assignment", "usecase_temporal_org",
        "usecase_policy_workflow", "usecase_advanced_payroll", "usecase_logistics_lineage",
        "usecase_security_files", "dfd_current_level0", "dfd_tenant_assignment", "dfd_policy_workflow",
        "dfd_advanced_payroll", "dfd_logistics_lineage", "dfd_file_security", "sequence_payroll_resumable",
        "activity_tenant_reconciliation", "database_landscape", "database_hrcore", "database_contracts", "database_payroll",
        "database_integration", "database_compact_gateway", "database_compact_admin", "database_compact_auth",
        "database_compact_compliance", "database_compact_analytics", "database_compact_files",
        "database_compact_health", "database_compact_movements", "database_compact_training",
        "database_compact_idcard", "database_compact_it", "database_compact_archive", "deployment", "tests",
    ],
    "3",
)
relabel_figures(list(USECASE_DRAWIO_MAP), "ج")
relabel_figures(
    ["dfd_main", "dfd_new_employee", "dfd_transfer", "dfd_medical", "dfd_training", "dfd_payroll", "dfd_complaint", "dfd_separation", "dfd_dictionary"],
    "د",
)
relabel_figures(DATABASE_FULL_ERD_KEYS, "ح")

FIGURE_CATALOG_KEYS = [
    "generic_architecture", "events", "solution", "tenant_context", "scope", "usecase", "usecase_platform",
    "dfd", "dfd_governance", "usecase_org", "usecase_lifecycle", "usecase_payroll", "usecase_people",
    "usecase_portals", "activity_hire", "activity_contract", "activity_leave", "activity_movement",
    "activity_separation", "activity_payroll", "activity_onboarding", "activity_policy", "activity_asset",
    "sequence_hire", "sequence_contract", "sequence_movement", "sequence_separation", "sequence_payroll",
    "sequence_policy", "sequence_code", "sequence_asset", "domain", "erd", "component", "component_governance",
    "deployment_formal", "security", "security_rls", "event_flow", "ownership", "reliability_journeys", "rbac",
    "trace", "usecase_current_master", "usecase_tenant_assignment", "usecase_temporal_org",
    "usecase_policy_workflow", "usecase_advanced_payroll", "usecase_logistics_lineage",
    "usecase_security_files", "dfd_current_level0", "dfd_tenant_assignment", "dfd_policy_workflow",
    "dfd_advanced_payroll", "dfd_logistics_lineage", "dfd_file_security", "sequence_payroll_resumable",
    "activity_tenant_reconciliation", "database_landscape", "database_hrcore", "database_contracts", "database_payroll",
    "database_integration", *DATABASE_COMPACT_ERD_MAP.keys(), "deployment", "tests", *USECASE_DRAWIO_MAP.keys(),
    "dfd_main", "dfd_new_employee", "dfd_transfer", "dfd_medical", "dfd_training", "dfd_payroll",
    "dfd_complaint", "dfd_separation", "dfd_dictionary", *DATABASE_FULL_ERD_KEYS,
]


FIGURE_AFTER_HEADING = {
    "## 2.4": "generic_architecture",
    "## 2.5": "events",
    "## 3.2": "solution",
    "## 3.6": "tenant_context",
    "## 3.7": "component_governance",
    "## 3.11": "security_rls",
    "## 3.12": "deployment",
}


CITATION_NOTES = {
    "## 2.1": "تُعرض المفاهيم في هذا القسم بوصفها إطاراً لإدارة دورة حياة الموظف والبيانات التنظيمية، مع الاستناد إلى مبادئ النمذجة الموجهة بالمجال [1].",
    "## 2.3": "يميز النموذج بين عزل الجهة ونطاق الوحدة التنظيمية، ويستند دفاع قاعدة البيانات إلى مرشحات EF Core وسياسات PostgreSQL [5][10].",
    "## 2.4": "يعتمد هذا العرض على أدبيات الخدمات المصغرة والتصميم الموجه بالمجال التي تشدد على استقلالية النشر وحدود الملكية [1][2][3].",
    "## 2.5": "تستند مفاهيم الرسائل وإعادة المحاولة والمستهلكين إلى وثائق RabbitMQ وMassTransit الرسمية [7][8].",
    "## 2.6": "تستند الخيارات المتعلقة بملكية البيانات وEF Core وPostgreSQL إلى مراجع البيانات الموزعة والوثائق الرسمية [5][10][22].",
    "## 2.8": "تُعرض مفاهيم OpenID Connect والأدوار والسياسات بالاستناد إلى وثائق Keycloak وإطار التحكم بالوصول القائم على السمات [9][14][35][36].",
    "## 2.9": "تستند مبادئ الاستمرارية والتعافي إلى إرشادات NIST ومراجع هندسة الاعتمادية [25][32].",
    "## 2.10": "يستند اختيار محاور الجودة والتقييم إلى نموذج ISO/IEC 25010، مع تمييز الاختبار الوظيفي عن قياس الأداء [13].",
    "## 3.1": "يعرض هذا الفصل بيئة التنفيذ والبنية الناتجة عن مراحل التحليل والتصميم، ويستند توصيف المكونات إلى وثائق المشروع وتقارير التحقق المعتمدة [15][16].",
}


TABLE_TITLES = {
    "## 1.9": ["المتطلبات الوظيفية الأساسية", "المتطلبات غير الوظيفية"],
    "## 2.2": ["مقارنة المعمارية الأحادية والخدمات المصغرة"],
    "## 2.7": ["الأدوار وسياسات الوصول"],
    "## 3.1": ["بيئة التطوير والتشغيل"],
    "## 3.2": ["مؤشرات بنية النظام المنفذ"],
    "## 3.3": ["خدمات المنظومة وقواعد بياناتها"],
    "## 3.4": ["مؤشرات تنفيذ الخدمات ومكونات العرض والبيانات"],
    "## 3.10": ["عينة من عقود أحداث التكامل"],
    "### 3.18.1": ["الكيانات الأساسية في خدمة HR.Core"],
    "### 3.18.2": ["الكيانات الأساسية في خدمة HR.Contracts"],
    "### 3.18.3": ["الكيانات الأساسية في خدمة Payroll"],
    "### 3.18.4": ["الكيانات الأساسية في خدمة HR.Health"],
    "### 3.18.5": ["الكيانات الأساسية في خدمة HR.Movements"],
    "### 3.18.6": ["الكيانات الأساسية في خدمة HR.Training"],
    "### 3.18.7": ["الكيانات الأساسية في خدمة IDCard"],
    "### 3.18.9": ["الكيانات الأساسية في خدمة Admin.Logistics"],
    "### 3.18.10": ["الكيانات الأساسية في خدمة IT.Infra"],
    "### 3.18.11": ["الكيانات الأساسية في خدمة Dev.Analytics"],
    "### 3.18.12": ["الكيانات الأساسية في خدمة Compliance"],
    "### 3.18.13": ["الكيانات الأساسية في خدمة Auth.Directory"],
    "### 3.18.14": ["الكيانات الأساسية في خدمة File.Storage"],
    "## 3.16": ["تحليل المخاطر والضوابط"],
}


def _font(path: Path, size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidate = ARIAL_BOLD_PATH if bold else path
    return ImageFont.truetype(str(candidate), size)


def _centered(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, font, fill: str = "FFFFFF") -> None:
    left, top, right, bottom = box
    bounds = draw.multiline_textbbox((0, 0), text, font=font, spacing=6, align="center")
    width = bounds[2] - bounds[0]
    height = bounds[3] - bounds[1]
    color = fill if fill.startswith("#") else f"#{fill}"
    draw.multiline_text(((left + right - width) / 2, (top + bottom - height) / 2), text, font=font, fill=color, spacing=6, align="center")


def _box(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], title: str, fill: str = TEAL, font_size: int = 30) -> None:
    draw.rounded_rectangle(box, radius=24, fill=f"#{fill}", outline="#0D3B44", width=3)
    _centered(draw, box, title, _font(ARIAL_PATH, font_size, True))


def _arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], fill: str = INK, width: int = 7) -> None:
    draw.line([start, end], fill=f"#{fill}", width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 >= x1 else -1
        points = [(x2, y2), (x2 - 20 * direction, y2 - 13), (x2 - 20 * direction, y2 + 13)]
    else:
        direction = 1 if y2 >= y1 else -1
        points = [(x2, y2), (x2 - 13, y2 - 20 * direction), (x2 + 13, y2 - 20 * direction)]
    draw.polygon(points, fill=f"#{fill}")


def _canvas(title: str) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (1800, 1020), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1800, 86), fill=f"#{INK}")
    draw.text((58, 22), title, font=_font(ARIAL_PATH, 34, True), fill="#FFFFFF")
    return image, draw


def make_figures() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)

    image, draw = _canvas("Institutional model: central governance and branch operations")
    _box(draw, (650, 135, 1150, 285), "Central HR\nGovernance", INK, 38)
    for x, label in [(120, "Branch A\nOperations"), (650, "Branch B\nOperations"), (1180, "Branch C\nOperations")]:
        _box(draw, (x, 510, x + 500, 670), label, TEAL, 36)
        _arrow(draw, (900, 285), (x + 250, 510))
    draw.text((468, 770), "Central view: all units | Branch view: own unit and descendants only", font=_font(ARIAL_PATH, 35, True), fill=f"#{GRAY}")
    image.save(ASSET_DIR / FIGURES["institution"].filename)

    image, draw = _canvas("Microservices architecture: independent domains behind one gateway")
    _box(draw, (85, 350, 355, 525), "Users\n& ESS", "356B8C", 34)
    _box(draw, (500, 350, 820, 525), "YARP\nGateway", INK, 38)
    _box(draw, (970, 170, 1350, 315), "Keycloak\nIdentity", "6A4C93", 34)
    _box(draw, (970, 395, 1350, 540), "HR & Support\nServices (14)", TEAL, 32)
    _box(draw, (1450, 395, 1720, 540), "RabbitMQ\nEvents", "B86622", 32)
    _box(draw, (970, 700, 1350, 845), "PostgreSQL\nService Databases", "2D7A4B", 32)
    _box(draw, (1450, 700, 1720, 845), "MinIO\nFiles", "2D7A4B", 32)
    _arrow(draw, (355, 438), (500, 438))
    _arrow(draw, (820, 438), (970, 468))
    _arrow(draw, (820, 400), (970, 242))
    _arrow(draw, (1350, 468), (1450, 468))
    _arrow(draw, (1160, 540), (1160, 700))
    _arrow(draw, (1350, 500), (1450, 770))
    image.save(ASSET_DIR / FIGURES["generic_architecture"].filename)

    image, draw = _canvas("Service layering: responsibility is separated inside each bounded context")
    layers = [(180, "Web / MVC\nControllers and Razor Views", "356B8C"), (365, "Application\nUse cases and policies", "0F6B73"), (550, "Domain\nEntities and rules", "6A4C93"), (735, "Infrastructure\nEF Core, adapters, messaging", "2D7A4B")]
    for y, label, color in layers:
        _box(draw, (330, y, 1470, y + 130), label, color, 35)
    draw.text((440, 915), "Shared defaults: authentication, health checks, correlation IDs, audit, integration contracts", font=_font(ARIAL_PATH, 31), fill=f"#{GRAY}")
    image.save(ASSET_DIR / FIGURES["layers"].filename)

    image, draw = _canvas("Event-driven integration: services exchange facts, not shared tables")
    labels = ["HR.Core\nEmployeeCreated", "RabbitMQ\nMessage bus", "Contracts\nPrepare contract", "Payroll\nPay profile", "Archive\nCase archive"]
    colors = ["356B8C", "B86622", "0F6B73", "2D7A4B", "6A4C93"]
    xs = [70, 410, 750, 1090, 1430]
    for x, label, color in zip(xs, labels, colors):
        _box(draw, (x, 400, x + 290, 590), label, color, 30)
    for x in xs[:-1]:
        _arrow(draw, (x + 290, 495), (x + 340, 495))
    _centered(draw, (180, 735, 1620, 890), "Outbox, retry and idempotent consumers\nreduce duplicate processing and improve traceability", _font(ARIAL_PATH, 31, True), GRAY)
    image.save(ASSET_DIR / FIGURES["events"].filename)

    image, draw = _canvas("HRManagement solution architecture")
    _box(draw, (85, 420, 345, 585), "Gateway", INK, 38)
    groups = [(470, "Core HR\nHR.Core\nHR.Contracts\nPayroll", "356B8C"), (820, "People Services\nHealth\nMovements\nTraining\nIDCard", TEAL), (1170, "Support Services\nArchive\nLogistics\nIT\nAnalytics", "6A4C93"), (1520, "Governance\nCompliance\nDirectory\nFiles", "B86622")]
    for x, label, color in groups:
        _box(draw, (x, 315, x + 245, 690), label, color, 27)
    _arrow(draw, (345, 502), (470, 502))
    _arrow(draw, (715, 502), (820, 502))
    _arrow(draw, (1065, 502), (1170, 502))
    _arrow(draw, (1415, 502), (1520, 502))
    _centered(draw, (250, 790, 1550, 900), "14 domain services | 15 application databases | separate Keycloak store", _font(ARIAL_PATH, 28, True), GRAY)
    image.save(ASSET_DIR / FIGURES["solution"].filename)

    image, draw = _canvas("Data ownership: each service owns its schema and exchanges identifiers or events")
    items = [(140, "HR.Core\nEmployee + OrgUnit", "356B8C"), (590, "Contracts\nContract + Leave", TEAL), (1040, "Payroll\nPay profile + Run", "2D7A4B"), (1490, "Health\nMedical records", "6A4C93")]
    for x, label, color in items:
        _box(draw, (x, 320, x + 260, 530), label, color, 24)
        _box(draw, (x, 680, x + 260, 810), "Own PostgreSQL\ndatabase", "6D7B83", 26)
        _arrow(draw, (x + 130, 530), (x + 130, 680))
    draw.text((335, 895), "No cross-service database joins: references and messages maintain the integration boundary", font=_font(ARIAL_PATH, 32, True), fill=f"#{GRAY}")
    image.save(ASSET_DIR / FIGURES["data"].filename)

    image, draw = _canvas("Organizational scope: central roles are global; branch managers are subtree scoped")
    _box(draw, (665, 140, 1135, 285), "Central HR\nGlobal scope", INK, 38)
    _box(draw, (160, 445, 600, 600), "Branch A\nManager scope root", TEAL, 34)
    _box(draw, (670, 445, 1130, 600), "Branch B\nIndependent scope", "356B8C", 34)
    _box(draw, (1200, 445, 1640, 600), "Branch C\nIndependent scope", "6A4C93", 34)
    _box(draw, (210, 760, 550, 885), "A.1 and A.2\nallowed descendants", "2D7A4B", 28)
    _arrow(draw, (900, 285), (380, 445))
    _arrow(draw, (900, 285), (900, 445))
    _arrow(draw, (900, 285), (1420, 445))
    _arrow(draw, (380, 600), (380, 760))
    draw.text((620, 925), "scope_org_unit_id + OrgScopeService + visibility policies", font=_font(ARIAL_PATH, 31, True), fill=f"#{GRAY}")
    image.save(ASSET_DIR / FIGURES["scope"].filename)

    image, draw = _canvas("Hiring lifecycle: one business fact drives bounded operational effects")
    nodes = [(55, "Public\nApplication", "356B8C"), (335, "HR.Core\nEmployee", TEAL), (615, "Contracts\nContract", "6A4C93"), (895, "Directory\nAccount", "B86622"), (1175, "Payroll\nPay profile", "2D7A4B"), (1455, "Archive\nCase", "6D7B83")]
    for x, label, color in nodes:
        _box(draw, (x, 400, x + 230, 585), label, color, 29)
    for x, _, _ in nodes[:-1]:
        _arrow(draw, (x + 230, 492), (x + 280, 492))
    _centered(draw, (160, 720, 1640, 885), "EmployeeCreated coordinates contract, identity, payroll\nand archive activities through events", _font(ARIAL_PATH, 29, True), GRAY)
    image.save(ASSET_DIR / FIGURES["lifecycle"].filename)

    image, draw = _canvas("Deployment model: reproducible local stack and guarded production configuration")
    _box(draw, (90, 300, 350, 500), "Gateway\nPublic entry", INK, 34)
    _box(draw, (510, 180, 825, 330), "Keycloak\nOIDC / JWT", "6A4C93", 30)
    _box(draw, (510, 420, 825, 570), "14 Services\nContainers", TEAL, 30)
    _box(draw, (985, 180, 1300, 330), "RabbitMQ\nMessages", "B86622", 30)
    _box(draw, (985, 420, 1300, 570), "PostgreSQL\nDatabases", "2D7A4B", 30)
    _box(draw, (1460, 300, 1720, 500), "MinIO\nFiles", "6D7B83", 32)
    _arrow(draw, (350, 400), (510, 495))
    _arrow(draw, (350, 380), (510, 255))
    _arrow(draw, (825, 495), (985, 495))
    _arrow(draw, (825, 495), (985, 255))
    _arrow(draw, (1300, 495), (1460, 400))
    _centered(draw, (200, 720, 1600, 885), "Docker Compose, health endpoints, production overrides\nand disabled automatic migrations", _font(ARIAL_PATH, 29, True), GRAY)
    image.save(ASSET_DIR / FIGURES["deployment"].filename)

    image, draw = _canvas("Verification evidence: current results and declared remaining limits")
    _box(draw, (135, 235, 735, 585), "Release build\n0 errors\n2416 unit tests passed\n5 failed", "2D7A4B", 39)
    _box(draw, (1065, 235, 1665, 585), "Integration\n158 passed / 7 failed\n47 skipped\nE2E: 2 passed", TEAL, 33)
    _arrow(draw, (735, 410), (1065, 410))
    draw.rounded_rectangle((360, 750, 1440, 890), radius=22, fill="#FCE8E6", outline="#9B1C1C", width=3)
    _centered(draw, (360, 750, 1440, 890), "Docker-dependent RLS, Testcontainers and performance\nremain outside this measured run", _font(ARIAL_PATH, 26, True), "9B1C1C")
    image.save(ASSET_DIR / FIGURES["tests"].filename)

    make_formal_figures()
    # PlantUML sources in docs/thesis_diagrams are the editable academic source of truth.
    for key, source_name in CURATED_DIAGRAM_MAP.items():
        source_path = CURATED_DIAGRAM_DIR / source_name
        if source_path.exists():
            shutil.copyfile(source_path, ASSET_DIR / FIGURES[key].filename)
    # Retain the earlier Draw.io DFD package for historical traceability.
    for key, source_name in DRAWIO_DIAGRAM_MAP.items():
        source_path = DRAWIO_DIAGRAM_DIR / source_name
        if source_path.exists():
            shutil.copyfile(source_path, ASSET_DIR / FIGURES[key].filename)
    # Retain the earlier 21-sector use-case workbook for historical traceability.
    for key, source_name in USECASE_DRAWIO_MAP.items():
        source_path = USECASE_DRAWIO_DIR / source_name
        if source_path.exists():
            shutil.copyfile(source_path, ASSET_DIR / FIGURES[key].filename)
    # Database diagrams are generated from the repository's DbContext and migration evidence.
    for key, source_name in DATABASE_DIAGRAM_MAP.items():
        source_path = DATABASE_DIAGRAM_DIR / source_name
        if source_path.exists():
            shutil.copyfile(source_path, ASSET_DIR / FIGURES[key].filename)
    # Complete key/FK ERDs for every database context are appendix evidence.
    for key, source_name in DATABASE_FULL_ERD_MAP.items():
        source_path = DATABASE_FULL_ERD_DIR / source_name
        if source_path.exists():
            shutil.copyfile(source_path, ASSET_DIR / FIGURES[key].filename)
    for key, source_name in DATABASE_COMPACT_ERD_MAP.items():
        source_path = DATABASE_COMPACT_ERD_DIR / source_name
        if source_path.exists():
            shutil.copyfile(source_path, ASSET_DIR / FIGURES[key].filename)


def make_formal_figures() -> None:
    """Create compact UML/DFD-style figures from the documented project model."""
    image, draw = _canvas("Use-case model: actors interact through central governance and scoped branch operations")
    actors = [(80, 245, "Central HR"), (80, 675, "Branch Manager"), (1510, 245, "Employee"), (1510, 675, "Citizen")]
    for x, y, label in actors:
        draw.ellipse((x + 55, y, x + 105, y + 50), outline=f"#{INK}", width=5)
        draw.line((x + 80, y + 50, x + 80, y + 130), fill=f"#{INK}", width=5)
        draw.line((x + 35, y + 75, x + 125, y + 75), fill=f"#{INK}", width=5)
        draw.line((x + 80, y + 130, x + 40, y + 185), fill=f"#{INK}", width=5)
        draw.line((x + 80, y + 130, x + 120, y + 185), fill=f"#{INK}", width=5)
        draw.text((x - 5, y + 205), label, font=_font(ARIAL_PATH, 25, True), fill=f"#{INK}")
    draw.rounded_rectangle((390, 130, 1410, 850), radius=20, outline=f"#{TEAL}", width=6)
    draw.text((470, 160), "HRManagement system boundary", font=_font(ARIAL_PATH, 30, True), fill=f"#{TEAL}")
    usecases = [(530, 285, "Manage employee"), (930, 285, "Manage hire"), (530, 495, "Approve movement"), (930, 495, "Process payroll"), (530, 705, "Scoped reports"), (930, 705, "Self service")]
    for x, y, label in usecases:
        draw.ellipse((x, y, x + 300, y + 105), outline=f"#{INK}", width=4, fill="#F4F7F8")
        _centered(draw, (x, y, x + 300, y + 105), label, _font(ARIAL_PATH, 25, True), INK)
    for start, end in [((205, 365), (530, 338)), ((205, 795), (530, 548)), ((1595, 365), (1230, 758)), ((1595, 795), (1230, 338))]:
        draw.line([start, end], fill=f"#{GRAY}", width=4)
    image.save(ASSET_DIR / FIGURES["usecase"].filename)

    image, draw = _canvas("DFD context: requests, approvals, records and events across organizational levels")
    _box(draw, (650, 380, 1150, 560), "HRManagement\nSystem", TEAL, 38)
    entities = [(80, 170, "Central HR"), (80, 700, "Branch HR"), (1450, 170, "Employee"), (1450, 700, "Citizen")]
    for x, y, label in entities:
        _box(draw, (x, y, x + 260, y + 120), label, INK if "Central" in label else "356B8C", 28)
    for start, end, label in [((340, 230), (650, 420), "policies"), ((340, 760), (650, 500), "requests"), ((1450, 230), (1150, 420), "self service"), ((1450, 760), (1150, 500), "applications")]:
        _arrow(draw, start, end)
        draw.text(((start[0] + end[0]) // 2 - 35, (start[1] + end[1]) // 2), label, font=_font(ARIAL_PATH, 23, True), fill=f"#{GRAY}")
    _box(draw, (690, 720, 1110, 875), "Events + archive\n+ audit trail", "6A4C93", 30)
    _arrow(draw, (900, 560), (900, 720))
    image.save(ASSET_DIR / FIGURES["dfd"].filename)

    image, draw = _canvas("Activity diagram: hire case with validation, approval and downstream effects")
    steps = [(160, "Receive application"), (500, "Validate candidate"), (840, "Approve and activate contract"), (1180, "Create employee and publish events"), (1520, "Archive and issue effects")]
    for x, label in steps:
        _box(draw, (x, 420, x + 230, 570), label, TEAL if x in (500, 1180) else INK, 24)
    for x, _ in steps[:-1]:
        _arrow(draw, (x + 230, 495), (x + 285, 495))
    draw.ellipse((220, 185, 330, 295), fill="#E3F2E8", outline="#2D7A4B", width=5)
    _centered(draw, (220, 185, 330, 295), "Start", _font(ARIAL_PATH, 24, True), "2D7A4B")
    draw.ellipse((1470, 700, 1580, 810), fill="#FCE8E6", outline="#9B1C1C", width=5)
    _centered(draw, (1470, 700, 1580, 810), "End", _font(ARIAL_PATH, 24, True), "9B1C1C")
    _arrow(draw, (275, 295), (275, 420))
    _arrow(draw, (1635, 570), (1525, 700))
    draw.text((370, 720), "Alternative path: reject -> notify candidate -> retain audit record", font=_font(ARIAL_PATH, 28, True), fill=f"#{GRAY}")
    image.save(ASSET_DIR / FIGURES["activity_hire"].filename)

    image, draw = _canvas("Sequence diagram: hire command, events and idempotent consumers")
    lifelines = [(170, "Gateway"), (510, "HR.Core"), (850, "Contracts"), (1190, "Payroll"), (1530, "Archive")]
    for x, label in lifelines:
        draw.text((x, 150), label, font=_font(ARIAL_PATH, 27, True), fill=f"#{INK}")
        draw.line((x + 70, 220, x + 70, 850), fill=f"#{GRAY}", width=3)
    messages = [(300, 170, 510, "POST candidate"), (410, 510, 850, "HireCaseCompleted"), (520, 850, 1190, "EmployeeCreated"), (630, 850, 1530, "Archive case"), (740, 1190, 850, "Pay profile"), (820, 850, 510, "Ack / receipt")]
    for y, x1, x2, label in messages:
        _arrow(draw, (x1 + 70, y), (x2 + 70, y), fill=TEAL, width=4)
        draw.text(((x1 + x2) // 2, y - 27), label, font=_font(ARIAL_PATH, 21), fill=f"#{GRAY}")
    image.save(ASSET_DIR / FIGURES["sequence_hire"].filename)

    image, draw = _canvas("Domain model: employee is linked to organization, employment and service records")
    _box(draw, (690, 360, 1110, 535), "Employee", TEAL, 36)
    domains = [(150, 150, "OrgUnit"), (150, 650, "EmploymentRecord"), (1320, 150, "Contract"), (1320, 650, "ServiceTransaction")]
    for x, y, label in domains:
        _box(draw, (x, y, x + 330, y + 125), label, INK, 28)
        _arrow(draw, (x + 165, y + 125 if y < 360 else y), (900, 360 if y < 360 else 535))
    _centered(draw, (430, 760, 1370, 900), "EmployeeId and PrimaryOrgUnitId are references; cross-service facts travel through events", _font(ARIAL_PATH, 26, True), GRAY)
    image.save(ASSET_DIR / FIGURES["domain"].filename)

    image, draw = _canvas("RBAC and organizational scope: role permission is combined with branch scope")
    headers = [(100, "Role"), (480, "Central"), (820, "Own branch"), (1160, "Descendants"), (1500, "Sibling branch")]
    for x, label in headers:
        _box(draw, (x, 150, x + 250, 250), label, INK, 24)
    rows = [(320, "Central HR", ["ALLOW", "ALLOW", "ALLOW", "ALLOW"]), (470, "Branch manager", ["DENY", "ALLOW", "ALLOW", "DENY"]), (620, "Employee", ["DENY", "SELF", "DENY", "DENY"])]
    for y, role, values in rows:
        _box(draw, (100, y, 350, y + 100), role, TEAL, 22)
        for x, value in zip((480, 820, 1160, 1500), values):
            color = "2D7A4B" if value in ("ALLOW", "SELF") else "9B1C1C"
            _box(draw, (x, y, x + 250, y + 100), value, color, 24)
    draw.text((260, 825), "scope_org_unit_id -> OrgScopeService -> allowed unit set -> query/policy check", font=_font(ARIAL_PATH, 26, True), fill=f"#{GRAY}")
    image.save(ASSET_DIR / FIGURES["rbac"].filename)

    image, draw = _canvas("Traceability: requirement evidence is linked from business intent to test result")
    stages = [(80, "Requirement"), (440, "Use case"), (800, "Service / DB"), (1160, "Test"), (1520, "Evidence")]
    for x, label in stages:
        _box(draw, (x, 400, x + 240, 560), label, TEAL if x in (80, 1160) else INK, 28)
    for x, _ in stages[:-1]:
        _arrow(draw, (x + 240, 480), (x + 300, 480))
    _centered(draw, (210, 700, 1590, 850), "Example: FR-ORG-05 -> UC-SCOPE -> OrgScopeService -> TS-01..04 -> test result / screenshot", _font(ARIAL_PATH, 24, True), GRAY)
    image.save(ASSET_DIR / FIGURES["trace"].filename)

    sector_specs = {
        "usecase_org": ("Use-case sector: organizational structure and branch scope", "Central HR", "Branch manager", ["Manage OrgUnit", "Assign scope", "View descendants", "Deny sibling access"]),
        "usecase_lifecycle": ("Use-case sector: employee lifecycle and contracts", "HR officer", "Approver", ["Create employee", "Activate contract", "Approve leave", "Separate employee"]),
        "usecase_payroll": ("Use-case sector: payroll and financial effects", "Payroll officer", "Central approver", ["Open run", "Calculate pay", "Approve run", "Post adjustments"]),
        "usecase_people": ("Use-case sector: health, training and movements", "People officer", "Employee", ["Record medical case", "Enroll training", "Request transfer", "Issue certificate"]),
        "usecase_portals": ("Use-case sector: public careers and employee self-service", "Citizen", "Employee", ["Browse jobs", "Submit application", "Track request", "Apply internally"]),
    }
    for key, (title, left_actor, right_actor, actions) in sector_specs.items():
        image, draw = _canvas(title)
        for x, label in [(120, left_actor), (1510, right_actor)]:
            draw.ellipse((x + 60, 270, x + 120, 330), outline=f"#{INK}", width=5)
            draw.line((x + 90, 330, x + 90, 435), fill=f"#{INK}", width=5)
            draw.line((x + 45, 370, x + 135, 370), fill=f"#{INK}", width=5)
            draw.line((x + 90, 435, x + 50, 490), fill=f"#{INK}", width=5)
            draw.line((x + 90, 435, x + 130, 490), fill=f"#{INK}", width=5)
            draw.text((x, 530), label, font=_font(ARIAL_PATH, 26, True), fill=f"#{INK}")
        draw.rounded_rectangle((420, 160, 1380, 825), radius=20, outline=f"#{TEAL}", width=6)
        for row, action in enumerate(actions):
            x = 520 if row % 2 == 0 else 920
            y = 290 + (row // 2) * 245
            draw.ellipse((x, y, x + 300, y + 110), outline=f"#{INK}", width=4, fill="#F4F7F8")
            _centered(draw, (x, y, x + 300, y + 110), action, _font(ARIAL_PATH, 24, True), INK)
            draw.line(((210 if row % 2 == 0 else 1590), 405 + (row // 2) * 220, (x if row % 2 == 0 else x + 300), y + 55), fill=f"#{GRAY}", width=3)
        image.save(ASSET_DIR / FIGURES[key].filename)

    activity_specs = {
        "activity_contract": ("Activity diagram: contract activation", ["Draft contract", "Validate terms", "Approve", "Activate", "Publish contract effect"]),
        "activity_leave": ("Activity diagram: leave and advance request", ["Submit request", "Check policy", "Approve or reject", "Create service transaction", "Publish payroll effect"]),
        "activity_movement": ("Activity diagram: transfer and secondment", ["Request movement", "Review scope", "Approve", "Execute movement", "Update employee unit"]),
        "activity_separation": ("Activity diagram: separation and final checklist", ["Open separation", "Complete checklist", "Approve", "Publish separation", "Revoke downstream assets"]),
        "activity_payroll": ("Activity diagram: payroll run lifecycle", ["Open draft", "Calculate lines", "Review", "Approve", "Post payroll run"]),
    }
    for key, (title, steps) in activity_specs.items():
        image, draw = _canvas(title)
        x_positions = [80, 425, 770, 1115, 1460]
        for x, step in zip(x_positions, steps):
            _box(draw, (x, 420, x + 260, 575), step, TEAL if x in (425, 1115) else INK, 23)
        for x in x_positions[:-1]:
            _arrow(draw, (x + 260, 498), (x + 345, 498))
        draw.ellipse((150, 175, 280, 305), fill="#E3F2E8", outline="#2D7A4B", width=5)
        _centered(draw, (150, 175, 280, 305), "Start", _font(ARIAL_PATH, 23, True), "2D7A4B")
        _arrow(draw, (215, 305), (215, 420))
        draw.ellipse((1510, 720, 1640, 850), fill="#FCE8E6", outline="#9B1C1C", width=5)
        _centered(draw, (1510, 720, 1640, 850), "End", _font(ARIAL_PATH, 23, True), "9B1C1C")
        _arrow(draw, (1590, 575), (1575, 720))
        image.save(ASSET_DIR / FIGURES[key].filename)

    sequence_specs = {
        "sequence_contract": ("Sequence diagram: contract activation", ["Gateway", "Contracts", "HR.Core", "Payroll", "Archive"], ["Activate command", "Validate employee", "ContractActivated", "Create pay profile", "Archive reference"]),
        "sequence_movement": ("Sequence diagram: transfer execution", ["Gateway", "Movements", "HR.Core", "Payroll", "Directory"], ["Execute transfer", "TransferExecuted", "Update OrgUnit", "Apply effect", "Sync directory"]),
        "sequence_separation": ("Sequence diagram: separation effects", ["Gateway", "Contracts", "Payroll", "IDCard", "Archive"], ["Separate command", "SeparationCompleted", "Final adjustment", "Revoke card", "Archive case"]),
        "sequence_payroll": ("Sequence diagram: payroll run posting", ["Gateway", "Payroll", "Rules", "Analytics", "Archive"], ["Post run", "Calculate rules", "PayrollProcessed", "Update KPI", "Archive report"]),
    }
    for key, (title, actors, messages) in sequence_specs.items():
        image, draw = _canvas(title)
        xs = [110, 470, 830, 1190, 1550]
        for x, actor in zip(xs, actors):
            draw.text((x, 150), actor, font=_font(ARIAL_PATH, 24, True), fill=f"#{INK}")
            draw.line((x + 65, 220, x + 65, 850), fill=f"#{GRAY}", width=3)
        for i, message in enumerate(messages[:len(xs) - 1]):
            y = 300 + i * 115
            x1 = xs[i] + 65
            x2 = xs[i + 1] + 65
            _arrow(draw, (x1, y), (x2, y), fill=TEAL, width=4)
            draw.text(((x1 + x2) // 2 - 40, y - 28), message, font=_font(ARIAL_PATH, 20), fill=f"#{GRAY}")
        image.save(ASSET_DIR / FIGURES[key].filename)

    image, draw = _canvas("Logical ERD: service-owned records share stable identifiers, not cross-database foreign keys")
    entities = [(100, 255, "Employee\nEmployeeId"), (480, 255, "OrgUnit\nOrgUnitId"), (860, 255, "Contract\nEmployeeId"), (1240, 255, "PayrollRun\nPeriod"), (480, 650, "LeaveRequest\nEmployeeId"), (860, 650, "SeparationRecord\nEmployeeId"), (1240, 650, "ArchiveItem\nReferenceNumber")]
    for x, y, label in entities:
        _box(draw, (x, y, x + 260, y + 130), label, TEAL if "Employee" in label else INK, 23)
    for start, end in [((360, 320), (480, 320)), ((360, 340), (860, 320)), ((1120, 320), (1240, 320)), ((610, 385), (610, 650)), ((990, 385), (990, 650)), ((1120, 710), (1240, 710))]:
        _arrow(draw, start, end, fill=GRAY, width=4)
    _centered(draw, (250, 875, 1550, 965), "Logical links are represented by identifiers and events across service database boundaries", _font(ARIAL_PATH, 22, True), GRAY)
    image.save(ASSET_DIR / FIGURES["erd"].filename)


def ensure_rfonts(run, arabic: bool | None = None) -> None:
    text = run.text
    has_arabic = bool(re.search(r"[\u0600-\u06FF]", text)) if arabic is None else arabic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    primary = ARABIC_FONT if has_arabic else ENGLISH_FONT
    rfonts.set(qn("w:ascii"), ENGLISH_FONT)
    rfonts.set(qn("w:hAnsi"), ENGLISH_FONT)
    rfonts.set(qn("w:cs"), primary)
    rfonts.set(qn("w:eastAsia"), primary)
    run.font.name = primary


def set_run_direction(run, rtl: bool | None = None) -> None:
    if rtl is None:
        return
    rpr = run._element.get_or_add_rPr()
    rtl_node = rpr.find(qn("w:rtl"))
    if rtl_node is None:
        rtl_node = OxmlElement("w:rtl")
        rpr.append(rtl_node)
    rtl_node.set(qn("w:val"), "1" if rtl else "0")


def set_run(run, size: float | None = None, bold: bool | None = None, color: str | None = None, italic: bool | None = None, font: str | None = None, rtl: bool | None = None) -> None:
    ensure_rfonts(run, arabic=font == ARABIC_FONT if font else None)
    if font:
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(key), font)
        run.font.name = font
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    set_run_direction(run, rtl)


def bidi(paragraph, alignment: WD_ALIGN_PARAGRAPH | None = WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        ppr.append(OxmlElement("w:bidi"))
    if alignment is not None:
        paragraph.alignment = alignment


def rtl_numbered_text(text: str) -> str:
    """Keep a leading Arabic section number on the right side in Word RTL text."""

    match = re.match(r"^(\d+(?:\.\d+)*)(\s+)(.+)$", text)
    if not match:
        return text
    return f"\u200f{match.group(1)}\u200f {match.group(3)}"


def rtl_list_item(number: str, text: str) -> str:
    """Render an explicit list number without inheriting Word's global list state."""

    return f"\u200f{number}.\u200f {text}"


def keep_with_next(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:keepNext")) is None:
        insert_before_first(
            ppr,
            OxmlElement("w:keepNext"),
            ("w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl", "w:numPr", "w:pBdr", "w:shd", "w:tabs", "w:bidi", "w:spacing", "w:ind", "w:jc"),
        )


def insert_before_first(parent, element, tags: tuple[str, ...]) -> None:
    for tag in tags:
        reference = parent.find(qn(tag))
        if reference is not None:
            parent.insert(parent.index(reference), element)
            return
    parent.append(element)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        reference = tc_pr.find(qn("w:tcMar"))
        if reference is None:
            tc_pr.append(shd)
        else:
            tc_pr.insert(tc_pr.index(reference), shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 85, left: int = 115, bottom: int = 85, right: int = 115) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        insert_before_first(tbl_pr, tbl_w, ("w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout"))
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        insert_before_first(tbl_pr, tbl_layout, ("w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"))
    tbl_layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    grid_cols = list(grid.gridCol_lst)
    for column, width in zip(grid_cols, widths):
        column.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table, color: str = "B8C2C9") -> None:
    table.style = "Table Grid"
    tbl_pr = table._tbl.tblPr
    bidi_visual = tbl_pr.find(qn("w:bidiVisual"))
    if bidi_visual is None:
        bidi_visual = OxmlElement("w:bidiVisual")
        insert_before_first(
            tbl_pr,
            bidi_visual,
            (
                "w:tblStyleRowBandSize",
                "w:tblStyleColBandSize",
                "w:tblW",
                "w:jc",
                "w:tblCellSpacing",
                "w:tblInd",
                "w:tblBorders",
                "w:shd",
                "w:tblLayout",
                "w:tblCellMar",
                "w:tblLook",
                "w:tblCaption",
                "w:tblDescription",
            ),
        )


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)


def set_paragraph_shading(paragraph, fill: str, border: str = "A7BBC5") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        insert_before_first(ppr, borders, ("w:shd", "w:tabs", "w:bidi", "w:spacing", "w:ind", "w:jc"))
    for edge in ("top", "left", "bottom", "right"):
        item = borders.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "6")
        item.set(qn("w:space"), "8")
        item.set(qn("w:color"), border)
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        insert_before_first(ppr, shd, ("w:tabs", "w:bidi", "w:spacing", "w:ind", "w:jc"))
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def add_field(paragraph, instruction: str, cached: str = "") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = cached
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run = OxmlElement("w:r")
        run.append(element)
        paragraph._p.append(run)


def set_page_number_start(section, start: int) -> None:
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        columns = sect_pr.find(qn("w:cols"))
        if columns is None:
            sect_pr.append(pg_num)
        else:
            sect_pr.insert(sect_pr.index(columns), pg_num)
    pg_num.set(qn("w:start"), str(start))


def continue_page_numbering(section) -> None:
    """Let a new section inherit the document's continuous page sequence."""
    pg_num = section._sectPr.find(qn("w:pgNumType"))
    if pg_num is not None:
        section._sectPr.remove(pg_num)


def add_page_number_footer(section) -> None:
    section.footer.is_linked_to_previous = False
    footer = section.footer
    p = footer.paragraphs[0]
    p.clear()
    bidi(p, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("الصفحة ")
    set_run(run, 10, color=GRAY)
    add_field(p, " PAGE ", "1")
    for run in p.runs:
        set_run(run, 10, color=GRAY)


def add_running_header(section) -> None:
    section.header.is_linked_to_previous = False
    header = section.header
    p = header.paragraphs[0]
    p.clear()
    bidi(p, WD_ALIGN_PARAGRAPH.RIGHT)
    p.paragraph_format.space_after = Pt(0)
    add_directional_text(p, "HRManagement | جامعة الشمال الخاصة | كلية الهندسة المعلوماتية", 9.5, GRAY)


def configure_section(section) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.right_margin = Cm(3)
    section.left_margin = Cm(2.54)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.header_distance = Cm(1.15)
    section.footer_distance = Cm(1.15)


def begin_landscape_section(doc: Document, horizontal_margin: float = 1.5) -> object:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.right_margin = Cm(horizontal_margin)
    section.left_margin = Cm(horizontal_margin)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.75)
    continue_page_numbering(section)
    add_running_header(section)
    add_page_number_footer(section)
    return section


def begin_portrait_section(doc: Document) -> object:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section)
    continue_page_numbering(section)
    add_running_header(section)
    add_page_number_footer(section)
    return section


ASCII_WORD_CHARS = set("ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789")
ZWSP = "\u200b"
LTR_INLINE_CHARS = set("ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789._:/#+'-") | {ZWSP}
LTR_LEADING_CHARS = set(".#")
LRM = "\u200e"


def has_arabic_text(text: str) -> bool:
    return bool(re.search(r"[\u0600-\u06FF]", text))


def is_ltr_start(text: str, index: int) -> bool:
    char = text[index]
    if char in ASCII_WORD_CHARS:
        return True
    return char in LTR_LEADING_CHARS and index + 1 < len(text) and text[index + 1] in ASCII_WORD_CHARS


def split_directional_text(text: str) -> list[tuple[str, bool]]:
    """Split Arabic prose into RTL text and small LTR islands for Latin words/numbers."""

    if not text:
        return []
    mixed_arabic = has_arabic_text(text)
    if not mixed_arabic:
        return [(text, True)]

    pieces: list[tuple[str, bool]] = []
    start = 0
    index = 0
    while index < len(text):
        if not is_ltr_start(text, index):
            index += 1
            continue

        if start < index:
            pieces.append((text[start:index], False))

        span_start = index
        index += 1
        while index < len(text):
            char = text[index]
            if char in LTR_INLINE_CHARS:
                index += 1
                continue
            if char == " ":
                lookahead = index + 1
                while lookahead < len(text) and text[lookahead] == " ":
                    lookahead += 1
                if lookahead < len(text) and is_ltr_start(text, lookahead):
                    index = lookahead
                    continue
            break
        pieces.append((text[span_start:index], True))
        start = index

    if start < len(text):
        pieces.append((text[start:], False))
    return pieces


def protect_ltr_terms(text: str, isolate_ltr: bool = False) -> str:
    if not text or not (isolate_ltr or has_arabic_text(text)):
        return text
    protected: list[str] = []
    for segment, ltr in split_directional_text(text):
        protected.append(f"{LRM}{segment}{LRM}" if ltr else segment)
    return "".join(protected)


def add_directional_text(paragraph, text: str, size: float, color: str, bold: bool | None = None, italic: bool | None = None, font: str | None = None, isolate_ltr: bool | None = None) -> None:
    isolate = has_arabic_text(text) if isolate_ltr is None else isolate_ltr
    run_text = protect_ltr_terms(text, isolate)
    run = paragraph.add_run(run_text)
    rtl_run = has_arabic_text(text)
    run_font = font if font else (None if rtl_run else ENGLISH_FONT)
    set_run(run, size, bold=bold, italic=italic, color=color, font=run_font, rtl=rtl_run)


def add_inline(paragraph, text: str, size: float = 13.5, color: str = INK) -> None:
    isolate_ltr = has_arabic_text(text)
    pieces = re.split(r"(\*\*.*?\*\*|`[^`]+`|\[[0-9]+(?:\]\[[0-9]+)*\])", text)
    for piece in pieces:
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            add_directional_text(paragraph, piece[2:-2], size, color, bold=True, isolate_ltr=isolate_ltr)
        elif piece.startswith("`") and piece.endswith("`"):
            add_directional_text(paragraph, piece[1:-1], min(size, 10.5), TEAL, font=CODE_FONT, isolate_ltr=isolate_ltr)
        elif piece.startswith("[") and piece.endswith("]"):
            add_directional_text(paragraph, piece, 11.5, TEAL, font=ENGLISH_FONT, isolate_ltr=isolate_ltr)
        else:
            add_directional_text(paragraph, piece, size, color, isolate_ltr=isolate_ltr)


def add_table_wrap_points(text: str) -> str:
    """Add invisible wrap opportunities to long technical identifiers in table cells."""
    text = re.sub(r"([,;/._-])(?=[A-Za-z0-9])", rf"\1{ZWSP}", text)
    return re.sub(r"(?<=[a-z0-9])(?=[A-Z])", ZWSP, text)


def add_para(doc: Document, text: str = "", style: str | None = None, alignment: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.RIGHT, before: float = 0, after: float = 7, size: float = 13.5) -> object:
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    bidi(p, alignment)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        add_inline(p, text, size=size)
    return p


def add_heading(doc: Document, text: str, level: int) -> object:
    style_name = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
    p = doc.add_paragraph(style=style_name)
    bidi(p, WD_ALIGN_PARAGRAPH.RIGHT)
    p.paragraph_format.keep_with_next = True
    numbered = re.match(r"^(\d+(?:\.\d+)*)(\s+)(.+)$", text)
    if numbered:
        number_run = p.add_run(f"\u200f{numbered.group(1)}\u200f")
        title_text = f" {numbered.group(3)}"
    else:
        number_run = None
        title_text = text
    if level == 1:
        size, color = 18, INK
    elif level == 2:
        size, color = 16, TEAL
    else:
        size, color = 14.5, INK
    if number_run is not None:
        set_run(number_run, size, bold=True, color=color, font=ENGLISH_FONT, rtl=False)
    add_directional_text(p, title_text, size, color, bold=True)
    return p


def add_code(doc: Document, code: str) -> None:
    p = doc.add_paragraph()
    bidi(p, WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.right_indent = Cm(0.45)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    for index, line in enumerate(code.splitlines()):
        run = p.add_run(line)
        set_run(run, 9, color=INK, font=CODE_FONT)
        if index < len(code.splitlines()) - 1:
            run.add_break()


def add_callout(doc: Document, title: str, text: str, fill: str = LIGHT) -> None:
    p = doc.add_paragraph()
    bidi(p, WD_ALIGN_PARAGRAPH.RIGHT)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.right_indent = Cm(0.2)
    set_paragraph_shading(p, fill)
    add_directional_text(p, title + " ", 13, INK, bold=True)
    add_inline(p, text, 12.5, INK)


def add_caption(doc: Document, label: str, title: str, table: bool = False) -> None:
    p = doc.add_paragraph()
    bidi(p, WD_ALIGN_PARAGRAPH.RIGHT if table else WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(5 if table else 3)
    p.paragraph_format.space_after = Pt(4)
    keep_with_next(p) if table else None
    add_directional_text(p, f"{label}: ", 12, INK, bold=True)
    add_directional_text(p, title, 12, GRAY, italic=True)


def add_figure(doc: Document, figure_key: str, max_width: float = 5.92, max_height: float = 7.05) -> None:
    figure = FIGURES[figure_key]
    if figure_key in USECASE_DRAWIO_MAP:
        prefix = "ج"
    elif figure_key in DRAWIO_DIAGRAM_MAP and figure_key != "dfd":
        prefix = "د"
    elif figure_key in DATABASE_FULL_ERD_MAP:
        prefix = "ح"
    else:
        prefix = str(figure.chapter)
    FIGURE_COUNTERS[prefix] = FIGURE_COUNTERS.get(prefix, 0) + 1
    label = f"الشكل ({prefix}-{FIGURE_COUNTERS[prefix]})"
    path = ASSET_DIR / figure.filename
    with Image.open(path) as bitmap:
        width_px, height_px = bitmap.size
    width_in = max_width
    height_in = width_in * height_px / width_px
    if height_in > max_height:
        height_in = max_height
        width_in = height_in * width_px / height_px
    p = doc.add_paragraph()
    bidi(p, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in), height=Inches(height_in))
    shape = doc.inline_shapes[-1]
    shape._inline.docPr.set("descr", figure.title)
    shape._inline.docPr.set("title", label)
    add_caption(doc, label, figure.title, table=False)
    source_text = (
        "المصدر: إعداد الباحثين استناداً إلى نموذج البيانات النهائي المعتمد لكل سياق خدمة."
        if figure_key in DATABASE_DIAGRAM_MAP or figure_key in DATABASE_FULL_ERD_MAP or figure_key in DATABASE_COMPACT_ERD_MAP
        else
        "المصدر: إعداد الباحثين استناداً إلى تحليل المتطلبات وحالات الاستخدام المعتمدة."
        if figure_key in USECASE_DRAWIO_MAP
        else
        "المصدر: إعداد الباحثين استناداً إلى تحليل العمليات وتدفقات البيانات المعتمدة."
        if figure_key in DRAWIO_DIAGRAM_MAP
        else
        "المصدر: إعداد الباحثين استناداً إلى نماذج التحليل والتصميم المعتمدة."
        if figure_key in CURATED_DIAGRAM_MAP
        else "المصدر: إعداد الباحثين استناداً إلى تصميم النظام وتنفيذه."
    )
    source = add_para(doc, source_text, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=8, size=10.5)
    for run in source.runs:
        run.italic = True
        set_run(run, 10.5, color=GRAY)


def add_table(
    doc: Document,
    rows: list[list[str]],
    title: str,
    chapter: int,
    table_no: int,
    label_override: str | None = None,
    column_weights: list[float] | None = None,
) -> None:
    if not rows:
        return
    width_total = 14000 if doc.sections[-1].orientation == WD_ORIENT.LANDSCAPE else 8300
    columns = len(rows[0])
    weights = column_weights or [1] * columns
    if column_weights is None:
        if columns == 2:
            weights = [3, 7]
        elif columns == 3:
            weights = [2, 4, 4]
        elif columns == 4:
            weights = [2, 3, 3, 2]
        elif columns >= 5:
            weights = [2] * columns
            weights[0] = 1
    elif len(column_weights) != columns:
        raise ValueError("column_weights must match the table column count")
    widths = [max(850, int(width_total * item / sum(weights))) for item in weights]
    widths[-1] += width_total - sum(widths)
    if label_override is None:
        TABLE_COUNTERS[chapter] = TABLE_COUNTERS.get(chapter, 0) + 1
        label = f"الجدول ({chapter}-{TABLE_COUNTERS[chapter]})"
    else:
        label = label_override
    add_caption(doc, label, title, table=True)
    table = doc.add_table(rows=len(rows), cols=columns)
    set_table_geometry(table, widths)
    set_table_borders(table)
    for row_index, values in enumerate(rows):
        for cell, value in zip(table.rows[row_index].cells, values):
            p = cell.paragraphs[0]
            bidi(p, WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline(p, add_table_wrap_points(value), 10.5 if row_index else 11)
            if row_index == 0:
                set_cell_shading(cell, TABLE_HEADER)
                for run in p.runs:
                    set_run(run, 11, bold=True, color=INK)
    set_repeat_table_header(table.rows[0])
    source = add_para(doc, "المصدر: إعداد الباحثين استناداً إلى تحليل النظام وتصميمه ونتائج التحقق، ما لم يذكر خلاف ذلك.", alignment=WD_ALIGN_PARAGRAPH.RIGHT, after=8, size=9.8)
    for run in source.runs:
        run.italic = True
        set_run(run, 9.8, color=GRAY)


def add_metadata_table(doc: Document, rows: list[list[str]]) -> None:
    """Use a compact, uncatalogued key-value table for detailed appendix cards."""
    rows = [["الحقل", "التفصيل"]] + rows
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [2150, 6630])
    set_table_borders(table, color="C5D0D7")
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.rows[row_index].cells[col_index]
            p = cell.paragraphs[0]
            bidi(p, WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline(p, add_table_wrap_points(value), 10.4 if col_index else 10.8)
            if row_index == 0:
                set_cell_shading(cell, TABLE_HEADER)
                for run in p.runs:
                    set_run(run, 10.8, bold=True, color=INK)
            elif col_index == 0:
                set_cell_shading(cell, "EEF3F5")
                for run in p.runs:
                    set_run(run, 10.8, bold=True, color=INK)
    set_repeat_table_header(table.rows[0])


def markdown_rows(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            rows.append(cells)
    return rows


def strip_markdown(value: str) -> str:
    value = value.replace("**", "")
    value = re.sub(r"`([^`]+)`", r"\1", value)
    value = re.sub(r"<[^>]+>", "", value)
    return value.strip()


def apply_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = ARABIC_FONT
    normal.font.size = Pt(13.5)
    rpr = normal.element.rPr
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ENGLISH_FONT)
    rfonts.set(qn("w:hAnsi"), ENGLISH_FONT)
    rfonts.set(qn("w:cs"), ARABIC_FONT)

    for style_name, size, color, before, after in [
        ("Heading 1", 18, INK, 18, 10),
        ("Heading 2", 16, TEAL, 13, 6),
        ("Heading 3", 14.5, INK, 10, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = ARABIC_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15

    for list_style in ("List Bullet", "List Number"):
        style = doc.styles[list_style]
        style.font.name = ARABIC_FONT
        style.font.size = Pt(13.5)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.15

    if "Thesis Caption" not in [style.name for style in doc.styles]:
        caption = doc.styles.add_style("Thesis Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = ARABIC_FONT
        caption.font.size = Pt(12)
        caption.font.italic = True


def add_cover(doc: Document, english: bool = False) -> None:
    if english:
        lines = [
            ("Syrian Arab Republic", 14, True),
            ("North Private University", 16, True),
            ("Faculty of Informatics Engineering", 14, True),
            ("Department of Software Engineering and Information Systems", 13, True),
            ("", 12, False),
            ("AN INTEGRATED HUMAN RESOURCES PLATFORM", 22, True),
            ("FOR CENTRAL AND BRANCH ADMINISTRATIONS", 18, True),
            ("HRManagement", 16, True),
            ("", 12, False),
            ("Graduation Project submitted in partial fulfillment of the requirements", 13, False),
            ("for the Bachelor degree in Informatics Engineering", 13, False),
            ("", 12, False),
            ("Prepared by", 13, True),
            ("Omar Ahmad Al-Khattab | Ali Saleh Ahmad Al-Dahou", 13, False),
            ("Ali Mustafa Al-Saeed | Abdulrahman Haitham Tahhan", 13, False),
            ("", 12, False),
            ("Supervised by", 13, True),
            ("Dr. Maher Dubsawi", 14, True),
            ("", 12, False),
            ("Academic Year 2025 - 2026", 13, True),
        ]
    else:
        lines = [
            ("الجمهورية العربية السورية", 14, True),
            ("جامعة الشمال الخاصة", 17, True),
            ("كلية الهندسة المعلوماتية", 14, True),
            ("قسم هندسة البرمجيات ونظم المعلومات", 13, True),
            ("", 12, False),
            ("منصة موحدة لإدارة الموارد البشرية", 23, True),
            ("في الإدارة المركزية والإدارات الفرعية", 19, True),
            ("HRManagement", 16, True),
            ("", 12, False),
            ("مشروع تخرج مقدم لنيل درجة الإجازة في الهندسة المعلوماتية", 13.5, False),
            ("باختصاص هندسة البرمجيات ونظم المعلومات", 13.5, False),
            ("", 12, False),
            ("إعداد الطلاب", 13.5, True),
            ("عمر أحمد الخطاب | علي صالح أحمد الدحو", 13.5, False),
            ("علي مصطفى السعيد | عبدالرحمن هيثم طحان", 13.5, False),
            ("", 12, False),
            ("بإشراف", 13.5, True),
            ("د. ماهر دبساوي", 14.5, True),
            ("", 12, False),
            ("العام الدراسي 2025 - 2026", 13.5, True),
        ]
    for index, (text, size, bold) in enumerate(lines):
        p = doc.add_paragraph()
        bidi(p, WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6 if index < 4 else 10)
        p.paragraph_format.line_spacing = 1.0
        add_directional_text(p, text, size, INK, bold=bold)
    doc.add_page_break()


def add_front_matter(doc: Document) -> None:
    add_cover(doc, english=False)
    add_cover(doc, english=True)

    add_heading(doc, "كلمة الشكر", 1)
    add_para(doc, "نتقدم بخالص الشكر والتقدير إلى الدكتور ماهر دبساوي على إشرافه العلمي وملاحظاته المنهجية التي ساعدت الفريق على تحويل مشروع واسع المجالات إلى منظومة هندسية ذات حدود واضحة ودليل تحقق قابل للمراجعة.")
    add_para(doc, "كما نتقدم بالشكر إلى جامعة الشمال الخاصة، وكلية الهندسة المعلوماتية، وقسم هندسة البرمجيات ونظم المعلومات، وإلى كل من ساهم في مراجعة المتطلبات أو اختبار الواجهات أو مناقشة القرارات المعمارية.")
    doc.add_page_break()

    add_heading(doc, "الملخص", 1)
    add_para(doc, "تعالج هذه الأطروحة مشكلة تشتت بيانات وإجراءات الموارد البشرية في المؤسسات التي تجمع إدارة مركزية وإدارات فرعية. تحتاج الإدارة المركزية إلى رؤية شاملة واعتماد القرارات ومراقبة الامتثال، بينما تحتاج الفروع إلى تنفيذ العمليات اليومية ضمن نطاق تنظيمي مضبوط لا يتيح الاطلاع على بيانات الفروع الأخرى.")
    add_para(doc, "يقدم المشروع منصة HRManagement لإدارة الموارد البشرية وفق نموذج يجمع مركزية الحوكمة ولا مركزية التنفيذ. صيغت 108 متطلبات وظيفية و30 متطلباً غير وظيفي، ثم وزعت المسؤوليات على 14 خدمة مجال خلف بوابة موحدة و15 سياق بيانات تطبيقي تشمل تدقيق البوابة، مع مخزن هوية مستقل عبر Keycloak. يتضمن نموذج البيانات 274 DbSet، وتتبادل الخدمات 72 واجهة لعقود الرسائل. طُورت المنصة باستخدام .NET 10 وASP.NET Core MVC وPostgreSQL وRabbitMQ/MassTransit وYARP، مع ملفات Docker Compose للتشغيل.")
    add_para(doc, "تتمثل المساهمة الأولى في الجمع بين عزل الجهة ونطاقها التنظيمي: يحدد tenant_id المؤسسة، ويحدد scope_org_unit_id جذر الفرع داخلها. ويكمل ذلك مسار لإسناد البيانات القديمة وحجر الملتبس منها ومصالحته. وتشمل المساهمة الثانية هيكلاً تنظيمياً مؤرخاً ووظائف ممولة وخطة قوى عاملة، ومحرك سياسات قابلاً للإصدار والمحاكاة والإنفاذ والربط بسير عمل وتوقيع، وتوليداً مركزياً للرموز، وحزم تهيئة وطنية ومحلية. كما توسع نموذج الرواتب إلى إصدارات ثابتة ولقطات مدخلات وحساب قابل للاستئناف وتمويل وصرف جزئي ومتأخرات وأوامر دفع وتسوية وحجوزات، مع التصريح بأن الاعتماد القانوني والتكامل المالي الحي خارج نتائج الدراسة.")
    add_para(doc, "نجح بناء الحل في وضع Release من دون أخطاء. وفي تشغيل التحقق الحالي نجح 2416 من 2421 اختبار وحدة، ونجح 158 اختبار تكامل وفشل 7 وتخطى 47، كما نجح اختباران من اختبارات النهاية إلى النهاية وتخطى 55. كشفت النتائج فجوات معلنة في نطاق جزء من اللوجستيات وسلامة الربط، وفي تهيئة بيئة بعض اختبارات السياسات والرموز، إضافة إلى انتهاء دليل اختبار مؤقت. لم يتوفر Docker daemon لتشغيل Testcontainers وRLS الحية، ولم ينفذ قياس أداء أو قبول ميداني؛ لذلك تفرق الأطروحة بين القدرة المنفذة والتحقق المكتمل ولا تعرض المنصة كمنتج إنتاجي معتمد.")
    add_para(doc, "الكلمات المفتاحية: إدارة الموارد البشرية، الإدارة المركزية، الإدارات الفرعية، تعدد الجهات، الخدمات المصغرة، السياسات الحية، النطاق التنظيمي، Keycloak، RabbitMQ، PostgreSQL، .NET.", after=0)
    doc.add_page_break()

    add_heading(doc, "Abstract", 1)
    english = [
        "This thesis addresses fragmented human-resources data and workflows in institutions that operate through a central administration and multiple branches. Central HR requires global visibility and governance, while branch administrations need to execute daily work without seeing data that belongs to other branches.",
        "HRManagement is an integrated platform that combines centralized governance with scoped branch operations. The study defines 108 functional and 30 non-functional requirements and allocates them across 14 bounded domain services, 15 application data contexts, 274 declared DbSets, and 72 message-contract interfaces. The solution uses .NET 10, ASP.NET Core MVC, PostgreSQL, RabbitMQ/MassTransit, YARP, Keycloak, and Docker Compose.",
        "The contribution combines tenant isolation with organizational scope and adds governed tenant assignment and quarantine, dated organizational structures, funded positions, versioned policies linked to workflows, institutional code generation, and repeatable agency onboarding. The payroll model covers immutable configuration versions, input snapshots, resumable processing, partial disbursement, arrears, payment orders, reconciliation, fraud alerts, and payment holds, without claiming legal certification or live treasury integration.",
        "The current Release build completed with no errors. Verification passed 2,416 of 2,421 unit tests, passed 158 integration tests with 7 failures and 47 skips, and passed 2 end-to-end tests with 55 skips. The disclosed gaps concern logistics scope and binding guards, test-environment substitutes for policy and code-generation clients, and an expired temporary evidence fixture. Docker-dependent RLS and Testcontainers runs, performance benchmarks, disaster-recovery exercises, and field acceptance remain explicitly unverified.",
    ]
    for paragraph in english:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(paragraph)
        set_run(run, 12, color=INK, font=ENGLISH_FONT, rtl=False)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("Keywords: Human Resources Management, central administration, branch administration, multi-tenancy, microservices, live policies, organizational scope, Keycloak, RabbitMQ, PostgreSQL, .NET.")
    set_run(run, 12, bold=True, color=INK, font=ENGLISH_FONT, rtl=False)
    doc.add_page_break()

    add_heading(doc, "فهرس المحتويات", 1)
    add_toc_outline(doc)
    doc.add_page_break()

    add_heading(doc, "قائمة الأشكال", 1)
    figure_marker = doc.add_paragraph("[[FIGURE_CATALOG_PLACEHOLDER]]")
    bidi(figure_marker, WD_ALIGN_PARAGRAPH.RIGHT)
    doc.add_page_break()

    add_heading(doc, "قائمة الجداول", 1)
    marker = doc.add_paragraph("[[TABLE_CATALOG_PLACEHOLDER]]")
    bidi(marker, WD_ALIGN_PARAGRAPH.RIGHT)
    return
    table_catalog = [
        ("الجدول (1-1)", "كتالوج المتطلبات الوظيفية التفصيلي"),
        ("الجدول (1-2)", "المتطلبات غير الوظيفية ومعايير تحققها"),
        ("الجدول (1-3)", "تغطية مخططات حالات الاستخدام وربطها بعائلات المتطلبات"),
        ("الجدول (1-4)", "التفسير المنهجي لعلاقات حالات الاستخدام"),
        ("الجدول (1-5)", "خلاصة تتبع المتطلبات إلى التنفيذ والتحقق"),
        ("الجدول (1-6)", "تحليل جدوى التكلفة والمنفعة والقيود التشغيلية"),
        ("الجدول (1-7)", "مصفوفة أسئلة الدراسة وأدوات التقويم"),
        ("الجدول (2-1)", "مقارنة المعمارية الأحادية والخدمات المصغرة"),
        ("الجدول (2-2)", "الأدوار وسياسات الوصول"),
        ("الجدول (2-3)", "مقارنة نقدية لنماذج نظم الموارد البشرية"),
        ("الجدول (3-1)", "بيئة التطوير والتشغيل"),
        ("الجدول (3-2)", "مؤشرات بنية النظام المنفذ"),
        ("الجدول (3-3)", "خدمات المنظومة وقواعد بياناتها"),
        ("الجدول (3-4)", "مؤشرات تنفيذ الخدمات ومكونات العرض والبيانات"),
        ("الجدول (3-5)", "توصيف قواعد البيانات وسياقات ملكية البيانات"),
        ("الجدول (3-6)", "ERD منطقي مختصر وملكية البيانات في المجالات الجوهرية"),
        ("الجدول (3-7)", "عينة من عقود أحداث التكامل"),
        ("الجدول (3-8)", "نتيجة تشغيل اختبارات الوحدة"),
        ("الجدول (3-9)", "حالات التحقق من نموذج المركز والفروع"),
        ("الجدول (3-10)", "حدود تفسير نتائج المشروع"),
        ("الجدول (3-11)", "الكيانات الأساسية في خدمة HR.Core"),
        ("الجدول (3-12)", "الكيانات الأساسية في خدمة HR.Contracts"),
        ("الجدول (3-13)", "الكيانات الأساسية في خدمة Payroll"),
        ("الجدول (3-14)", "الكيانات الأساسية في خدمة HR.Health"),
        ("الجدول (3-15)", "الكيانات الأساسية في خدمة HR.Movements"),
        ("الجدول (3-16)", "الكيانات الأساسية في خدمة HR.Training"),
        ("الجدول (3-17)", "الكيانات الأساسية في خدمة IDCard"),
        ("الجدول (3-18)", "الكيانات الأساسية في خدمة Admin.Logistics"),
        ("الجدول (3-19)", "الكيانات الأساسية في خدمة IT.Infra"),
        ("الجدول (3-20)", "الكيانات الأساسية في خدمة Dev.Analytics"),
        ("الجدول (3-21)", "الكيانات الأساسية في خدمة Compliance"),
        ("الجدول (3-22)", "الكيانات الأساسية في خدمة Auth.Directory"),
        ("الجدول (3-23)", "الكيانات الأساسية في خدمة File.Storage"),
        ("الجدول (3-24)", "تحليل المخاطر والضوابط"),
        ("الجدول (3-25)", "القيود المعلنة وخطة المعالجة"),
        ("الجدول (خ-1)", "الاختصارات المستخدمة"),
        ("الجدول (م-1)", "مصفوفة التتبع الكاملة للمتطلبات الوظيفية"),
        ("الجدول (م-2)", "دليل الخدمات ومجالات ملكية البيانات"),
        ("الجدول (م-3)", "مطابقة البنية مع منهجية إعداد مشاريع التخرج"),
        ("الجدول (م-4)", "قائمة تدقيق نزاهة النتائج والبيانات"),
        ("الجدول (م-5)", "سجل النتائج القابلة لإعادة التشغيل"),
    ]
    table_catalog = [
        ("الجدول (1-1)", "أصحاب المصلحة ونطاق الوصول"),
        ("الجدول (1-2)", "تصنيف حالة المتطلبات"),
        ("الجدول (1-3)", "كتالوج المتطلبات الوظيفية التفصيلي"),
        ("الجدول (1-4)", "المتطلبات غير الوظيفية ومعايير تحققها"),
        ("الجدول (1-5)", "خلاصة تتبع المتطلبات إلى التنفيذ والتحقق"),
        ("الجدول (1-6)", "تحليل الجدوى والكلفة التشغيلية"),
        ("الجدول (1-7)", "أسئلة الدراسة وأدوات التقويم"),
        ("الجدول (2-1)", "مقارنة التطبيق الأحادي وخدمات المجال"),
        ("الجدول (2-2)", "مقارنة نقدية لنماذج نظم الموارد البشرية"),
        ("الجدول (3-1)", "مؤشرات خط الأساس الحالي"),
        ("الجدول (3-2)", "مجموعات خدمات المنظومة"),
        ("الجدول (3-3)", "الرحلات الحرجة وآثارها"),
        ("الجدول (3-4)", "حالات الوصول في نموذج الجهة والمركز والفروع"),
        ("الجدول (3-5)", "توصيف قواعد البيانات وسياقات ملكية البيانات"),
        ("الجدول (3-6)", "نتائج البناء والاختبارات الحالية"),
        ("الجدول (3-7)", "تحليل المخاطر والإجراءات المطلوبة"),
        ("الجدول (ب-1)", "تغطية حزمة حالات الاستخدام المؤسسية"),
        ("الجدول (ب-2)", "تفسير علاقات حالات الاستخدام"),
        ("الجدول (م-1)", "مصفوفة التتبع الكاملة"),
        ("الجدول (م-2)", "دليل الخدمات ومجالات الملكية"),
        ("الجدول (م-3)", "مطابقة الأطروحة مع المنهجية المرجعية"),
        ("الجدول (م-4)", "قائمة تدقيق نزاهة النتائج"),
        ("الجدول (م-5)", "سجل النتائج القابلة لإعادة التشغيل"),
    ]
    add_catalog(doc, table_catalog)


def add_catalog(doc: Document, items: Iterable[tuple[str, str]]) -> None:
    rows = [["الرقم", "العنوان"]] + [[label, title] for label, title in items]
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [2100, 6680])
    set_table_borders(table)
    for row_index, values in enumerate(rows):
        for cell, value in zip(table.rows[row_index].cells, values):
            p = cell.paragraphs[0]
            bidi(p, WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value, 11)
            if row_index == 0:
                set_cell_shading(cell, TABLE_HEADER)
                for run in p.runs:
                    set_run(run, 11, bold=True, color=INK)
    set_repeat_table_header(table.rows[0])


def materialize_table_catalog(doc: Document) -> None:
    marker = next((p for p in doc.paragraphs if p.text.strip() == "[[TABLE_CATALOG_PLACEHOLDER]]"), None)
    if marker is None:
        return
    items: list[tuple[str, str]] = []
    pattern = re.compile(r"^(الجدول\s*\([^)]+\)):\s*(.+)$")
    for paragraph in doc.paragraphs:
        match = pattern.match(paragraph.text.strip())
        if match:
            items.append((match.group(1), match.group(2)))
    table = doc.add_table(rows=len(items) + 1, cols=2)
    set_table_geometry(table, [2100, 6680])
    set_table_borders(table)
    values = [["الرقم", "العنوان"], *[[label, title] for label, title in items]]
    for row_index, row_values in enumerate(values):
        for cell, value in zip(table.rows[row_index].cells, row_values):
            paragraph = cell.paragraphs[0]
            bidi(paragraph, WD_ALIGN_PARAGRAPH.RIGHT)
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value, 11)
            if row_index == 0:
                set_cell_shading(cell, TABLE_HEADER)
                for run in paragraph.runs:
                    set_run(run, 11, bold=True, color=INK)
    set_repeat_table_header(table.rows[0])
    marker._p.addprevious(table._tbl)
    marker._element.getparent().remove(marker._element)


def materialize_figure_catalog(doc: Document) -> None:
    marker = next((p for p in doc.paragraphs if p.text.strip() == "[[FIGURE_CATALOG_PLACEHOLDER]]"), None)
    if marker is None:
        return
    items: list[tuple[str, str]] = []
    pattern = re.compile(r"^(الشكل\s*\([^)]+\)):\s*(.+)$")
    for paragraph in doc.paragraphs:
        match = pattern.match(paragraph.text.strip())
        if match:
            items.append((match.group(1), match.group(2)))
    table = doc.add_table(rows=len(items) + 1, cols=2)
    set_table_geometry(table, [2100, 6680])
    set_table_borders(table)
    values = [["الرقم", "العنوان"], *[[label, title] for label, title in items]]
    for row_index, row_values in enumerate(values):
        for cell, value in zip(table.rows[row_index].cells, row_values):
            paragraph = cell.paragraphs[0]
            bidi(paragraph, WD_ALIGN_PARAGRAPH.RIGHT)
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value, 11)
            if row_index == 0:
                set_cell_shading(cell, TABLE_HEADER)
                for run in paragraph.runs:
                    set_run(run, 11, bold=True, color=INK)
    set_repeat_table_header(table.rows[0])
    marker._p.addprevious(table._tbl)
    marker._element.getparent().remove(marker._element)


def add_toc_outline(doc: Document) -> None:
    p = doc.add_paragraph()
    bidi(p, WD_ALIGN_PARAGRAPH.RIGHT)
    add_field(p, ' TOC \\o "1-3" \\h \\z \\u ', "يُحدّث فهرس المحتويات تلقائياً عند فتح المستند أو قبل تصديره إلى PDF.")
    return
    entries = [
        ("الفصل الأول: المقدمة", 0),
        ("1.1 تمهيد", 1), ("1.2 مشكلة المشروع", 1), ("1.3 فكرة المشروع", 1),
        ("1.4 أهداف المشروع", 1), ("1.5 أهمية المشروع", 1),
        ("1.6 الجهة المستفيدة ونموذج المركز والفروع", 1),
        ("1.7 حدود المشروع", 1), ("1.8 منهجية التنفيذ", 1),
        ("1.9 المتطلبات الوظيفية وغير الوظيفية", 1), ("1.9.1 كتالوج المتطلبات التفصيلي", 2),
        ("1.9.2 مخططات حالات الاستخدام وتدفق البيانات", 2),
        ("1.9.2.1 حزمة مخططات حالات الاستخدام المعتمدة", 3),
        ("1.9.2.2 تفسير العلاقات ونماذج الحالات الحرجة", 3),
        ("1.9.2.3 حزمة مخططات DFD المعتمدة", 3), ("1.9.3 مصفوفة التتبع", 2),
        ("1.10 دراسة الجدوى", 1), ("1.11 دراسة جدوى موسعة", 1),
        ("1.12 أسئلة الدراسة وإطار التقييم", 1),
        ("الفصل الثاني: الدراسة النظرية", 0),
        ("2.1 نظم إدارة الموارد البشرية", 1), ("2.2 معمارية الخدمات المصغرة", 1),
        ("2.3 التصميم الموجه بالمجال", 1), ("2.4 التكامل القائم على الأحداث", 1),
        ("2.5 قاعدة بيانات لكل خدمة", 1), ("2.6 البوابة المركزية", 1),
        ("2.7 المصادقة والتفويض", 1), ("2.8 قواعد البيانات وORM", 1),
        ("2.9 النشر والتشغيل", 1), ("2.10 الاختبارات والجودة", 1),
        ("2.11 الدراسات والأنظمة المقارنة", 1), ("2.12 الأمن والحوكمة كمتطلبات تصميم", 1),
        ("الفصل الثالث: التطبيق العملي", 0),
        ("3.1 بيئة العمل", 1), ("3.2 بنية الحل", 1), ("3.3 تحليل النظام", 1),
        ("3.4 تصميم الخدمات", 1), ("3.5 تصميم قواعد البيانات", 1),
        ("3.5.1 منهجية تصميم قواعد البيانات وملكية البيانات", 2),
        ("3.5.2 آلية التواصل بين قواعد البيانات", 2),
        ("3.5.3 خرائط ERD المختصرة للخدمات المساندة", 2),
        ("3.6 نموذج المركز والفروع", 1), ("3.7 سير العمل", 1),
        ("3.7.8 مخططات التحليل والتصميم الرسمية", 2),
        ("3.8 واجهات المستخدم", 1), ("3.9 التكامل بين الخدمات", 1),
        ("3.10 الأمن والامتثال", 1), ("3.11 النشر والتشغيل", 1),
        ("3.12 الاختبارات والتحقق", 1), ("3.13 التحقق من نموذج المركز والفروع", 1),
        ("3.14 النتائج", 1), ("3.15 التحديات", 1), ("3.16 التوصيات", 1),
        ("3.17 شرح تفصيلي للخدمات", 1), ("3.18 تحليل المخاطر والضوابط", 1),
        ("3.19 حوكمة البيانات والقيود المتبقية", 1),
        ("الخاتمة", 0), ("قائمة الاختصارات", 0), ("المراجع", 0), ("الملاحق", 0),
    ]
    for text, depth in entries:
        p = doc.add_paragraph()
        bidi(p, WD_ALIGN_PARAGRAPH.RIGHT)
        p.paragraph_format.right_indent = Cm(0.65 * depth)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        add_directional_text(p, rtl_numbered_text(text), 12.2 if depth else 13.2, INK if depth == 0 else GRAY, bold=depth == 0)


def begin_chapter(doc: Document, name: str, subtitle: str, first: bool = False) -> None:
    if first:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        configure_section(section)
        set_page_number_start(section, 1)
        add_running_header(section)
        add_page_number_footer(section)
    else:
        doc.add_page_break()
    p = doc.add_paragraph()
    bidi(p, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(150)
    p.paragraph_format.space_after = Pt(18)
    add_directional_text(p, name, 38, INK, bold=True)
    p2 = doc.add_paragraph()
    bidi(p2, WD_ALIGN_PARAGRAPH.CENTER)
    p2.paragraph_format.space_after = Pt(0)
    add_directional_text(p2, subtitle, 18, TEAL, bold=True)
    doc.add_page_break()
    add_heading(doc, f"{name}: {subtitle}" if subtitle else name, 1)


def add_evaluation_framework(doc: Document) -> None:
    add_heading(doc, "1.12 أسئلة الدراسة وإطار التقييم", 2)
    add_para(doc, "تندرج الدراسة ضمن البحوث الهندسية التطبيقية؛ لذلك يركز التقويم على ملاءمة الحل للمشكلة المؤسسية، واتساق قرارات التصميم، وقابلية التحقق من السلوك. وصيغت ثلاثة أسئلة: إلى أي مدى يدعم النموذج التنظيمي عمل المركز والفروع؟ هل تحافظ المعمارية على حدود واضحة لملكية البيانات والمسؤوليات؟ وهل تحقق الاختبارات المنفذة معايير القبول المحددة ضمن نطاق الإصدار الأول؟")
    rows = [
        ["سؤال الدراسة", "أداة التقويم", "حدود الاستدلال"],
        ["هل يدعم النظام المركز والفروع؟", "نموذج الأدوار والنطاق التنظيمي وحالات قبول المركز والفرع", "تحقق وظيفي آلي؛ لا يمثل دراسة استخدام ميدانية"],
        ["هل تحترم الخدمات حدود البيانات؟", "نماذج ملكية البيانات وعقود الأحداث وآليات Outbox/Inbox", "تقويم تصميمي وتشغيلي لا يقيس زمن الاستجابة تحت حمل"],
        ["هل تحقق الوظائف معايير القبول؟", "نتائج الوحدة والتكامل وE2E الحالية مع إظهار النجاح والفشل والتخطي", "لا يشمل Testcontainers وRLS الحية أو اختبار الأداء"],
    ]
    add_table(doc, rows, "مصفوفة أسئلة الدراسة وأدوات التقويم", 1, 7)
    add_callout(doc, "حدود التقويم:", "تثبت الاختبارات الآلية صحة السلوك الذي تغطيه ضمن بيئة التحقق، ولا تُستخدم بديلاً عن اختبار الأداء أو القبول الميداني أو الاعتماد القانوني.")


def add_feasibility_detail(doc: Document) -> None:
    add_heading(doc, "1.11 دراسة جدوى موسعة", 2)
    add_para(doc, "الجدوى لا تعني أن البرمجيات مفتوحة المصدر بلا تكلفة. فاختيار PostgreSQL، RabbitMQ، Keycloak، وDocker يخفض رسوم الترخيص الابتدائية، لكنه ينقل جزءاً من الكلفة إلى الاستضافة، المراقبة، النسخ الاحتياطي، إدارة الهوية، تدريب فريق التشغيل، واختبار التحديثات. لذلك يفصل التحليل بين كلفة التملك وبين قرار شراء منتج HRMS جاهز.")
    rows = [
        ["البعد", "الفائدة في النسخة الحالية", "الكلفة أو القيد الذي يجب إدارته"],
        ["تقني", "حدود خدمة واضحة وقاعدة بيانات لكل خدمة", "تشغيل 14 خدمة وبنية رسائل ومراقبة"],
        ["اقتصادي", "تراخيص أولية منخفضة للتقنيات الأساسية", "خوادم، نسخ احتياطي، مراقبة، ودعم فني"],
        ["تشغيلي", "مركزية الحوكمة مع تنفيذ الفروع ضمن نطاق", "تدريب المستخدمين وحوكمة الأدوار والبيانات"],
        ["أمني", "Keycloak وسياسات وعزل مجالات حساسة", "تدوير الأسرار، مراجعة الصلاحيات، وتمارين الاستعادة"],
        ["قانوني", "تهيئة قواعد الرواتب بحسب نظام قانوني قابل للضبط", "لا تدعي النسخة الحالية اعتماد محرك رواتب قانوني نهائي"],
    ]
    add_table(doc, rows, "تحليل جدوى التكلفة والمنفعة والقيود التشغيلية", 1, 6)
    add_callout(doc, "قرار النطاق:", "يبرر اتساع النظام الحاجة إلى معمارية قابلة للتطور، لكنه لا يبرر ادعاء النشر الإنتاجي قبل استكمال اختبارات الأداء والتكامل الحي وسياسات المؤسسة.", fill="FFF7E6")


def add_related_work(doc: Document) -> None:
    add_heading(doc, "2.11 الدراسات والأنظمة المقارنة", 2)
    add_para(doc, "تظهر أدبيات نظم معلومات الموارد البشرية أن القيمة لا تأتي من حفظ بيانات الموظف فقط، بل من تكامل دورة العمل، جودة البيانات، والقدرة على إنتاج معلومات موثوقة للقرار [18][19]. كما تؤكد أدبيات الموارد البشرية الإلكترونية أن إدخال التقنية يعيد توزيع المسؤوليات بين المركز والوحدات التشغيلية؛ وهي فكرة تتقاطع مباشرة مع نموذج المركز والفروع في هذا المشروع [20].")
    rows = [
        ["النموذج", "نقطة القوة", "القيد عند تطبيقه على المؤسسة متعددة الفروع", "موقف HRManagement"],
        ["نظام مركزي أحادي", "بساطة النشر والمعاملات المشتركة", "تداخل المجالات وصعوبة توسع فرق العمل", "يفصل المجالات مع احتفاظ المركز بالرؤية"],
        ["HRMS سحابي عام", "جاهزية الواجهة والوظائف النمطية", "قيود على قواعد العمل والسيادة على البيانات", "يبني قواعد قابلة للتهيئة داخل المؤسسة"],
        ["خدمات مصغرة", "استقلال المجال والنشر وملكية البيانات", "كلفة الرسائل والمراقبة والاتساق النهائي", "اختارها بسبب اتساع المجالات ووثق كلفتها"],
        ["بوابة فرعية مستقلة", "استقلال تنفيذ الفرع", "خطر تشتت البيانات وتجاوز الصلاحية", "نطاق تنظيمي موحد مع منع الفروع الشقيقة"],
    ]
    add_table(doc, rows, "مقارنة نقدية لنماذج نظم الموارد البشرية", 2, 2)
    add_heading(doc, "2.11.1 مقارنة مع منتجات HRMS فعلية", 3)
    add_para(doc, "تستخدم المقارنة الآتية صفحات المنتجات الرسمية لتحديد النطاق المعلن، ولا تمثل اختبار شراء أو تقييماً مستقلاً لجودة المنتج. كما لا تقارن الأسعار أو الأداء أو الامتثال المحلي، لأن ذلك يحتاج عروضاً تعاقدية وبيئات تجريبية متكافئة. الغرض هو تحديد موضع المشروع بين حزم سحابية مؤسسية ومنصات معيارية قابلة للتخصيص [43]-[48].")
    begin_landscape_section(doc)
    products = [
        ["المنتج", "النطاق الذي تعلنه الجهة المنتجة", "نمط المنصة", "صلة المقارنة بالمشروع", "حد المقارنة"],
        ["SAP SuccessFactors Employee Central", "الموارد البشرية الأساسية، العمليات الموحدة، الخدمة الذاتية وسير العمل", "خدمة سحابية مؤسسية ضمن حزمة HCM", "مرجع لاتساع Core HR والعمل متعدد البلدان", "لم تختبر الأطروحة التهيئة المحلية أو التكلفة أو الأداء"],
        ["Oracle Fusion Cloud HCM", "منصة موحدة للموارد البشرية والمواهب والقوى العاملة مع منتج رواتب ضمن الحزمة", "حزمة سحابية مؤسسية مترابطة", "مرجع لتوحيد البيانات والعمليات المالية والبشرية", "لا تقارن الدراسة موصلات Oracle أو التغطية القانونية السورية"],
        ["Workday HCM", "نواة موارد بشرية وعمليات وأتمتة ومواهب وتعلم وتخطيط قوى عاملة", "منصة HCM سحابية موحدة", "مرجع للتكوين والتحليلات ودورة القوى العاملة", "لا تتوافر تجربة تشغيل متكافئة تسمح بحكم أداء أو سهولة استخدام"],
        ["Odoo Employees and HR", "تطبيقات معيارية للموظفين والتوظيف والإجازات والتقييمات ضمن مجموعة أعمال قابلة للتوسعة", "منصة تطبيقات معيارية ومتكاملة", "مرجع للتركيب التدريجي والتخصيص", "لا يعني التشابه الوظيفي تطابق الحوكمة أو العزل أو العمق المالي"],
        ["OrangeHRM", "إدارة الأفراد والتوظيف والتهيئة والأداء والإجازات، مع إصدار Starter مفتوح المصدر", "حل HR معياري بإصدارات مختلفة", "مرجع لبدء منخفض الكلفة ووظائف HR المعتادة", "لا تقارن الدراسة الإضافات التجارية أو قابلية التوسع المؤسسي عملياً"],
        ["HRManagement", "14 خدمة مجال، تعدد جهات، مركز وفروع، سياسة حية، رموز، بيانات مرجعية، رواتب ولوجستيات وملفات", "منصة مخصصة بخدمات وقواعد مملوكة للمجالات", "يركز على الحوكمة المحلية والعزل القابل للتنفيذ والتوسع الوظيفي", "مشروع v1 بلا اعتماد إنتاجي أو قياس أداء أو تكامل خارجي حي"],
    ]
    add_table(
        doc,
        products,
        "مقارنة نطاقية مع أنظمة HRMS مسماة",
        2,
        3,
        column_weights=[1.8, 3.4, 2.1, 2.7, 2.4],
    )
    begin_portrait_section(doc)
    add_para(doc, "لا تستنتج الدراسة تفوق HRManagement على هذه المنتجات؛ فالمنتجات المؤسسية تملك نضجاً سوقياً وخدمات دعم وتحديثات لا يثبتها مشروع تخرج. تتمثل خصوصية المشروع في قابلية تمثيل الجهة والمركز والفروع، وربط Tenant بالنطاق التنظيمي، وإتاحة سياسات ورموز وتهيئة محلية ضمن ملكية بيانات مستقلة. وفي المقابل يتحمل المشروع كلفة تشغيل وتحقق وتوطين وصيانة كان المنتج الجاهز سيوزع جزءاً منها على المزود.")
    decision_rows = [
        ["معيار القرار", "منتج HRMS جاهز", "HRManagement مخصص"],
        ["سرعة البدء", "أعلى عادة عند ملاءمة العمليات القياسية", "يحتاج تهيئة واختباراً وتشغيلاً مؤسسياً"],
        ["تخصيص الحوكمة المحلية", "مرتبط بحدود المنتج والعقد", "مرن لأن نموذج الجهة والسياسة والرمز جزء من التصميم"],
        ["المسؤولية التشغيلية", "جزء كبير لدى المزود في SaaS", "تقع على المؤسسة وفريق المشروع"],
        ["السيادة والتكامل", "تحددها خيارات الاستضافة والموصلات", "قابلة للضبط مع كلفة تنفيذ واعتماد أعلى"],
        ["الدليل المتاح هنا", "صفحات رسمية فقط، بلا تجربة مقارنة", "تصميم وتنفيذ واختبارات معلنة مع فجوات صريحة"],
    ]
    add_table(doc, decision_rows, "مصفوفة قرار الشراء أو البناء", 2, 4)
    add_para(doc, "لا تقدم الدراسة الخدمات المصغرة بوصفها غاية بحد ذاتها. فالمراجع المعمارية تشير إلى أن تقسيم الخدمات يعالج حدود المجال واستقلال النشر، لكنه يتطلب انضباطاً في العقود والرصد والاسترداد ومعالجة فشل الاتصال [2][21][22]. لهذا يستخدم المشروع عقود أحداث وإعادة محاولة وCorrelation ID وأنماط idempotency، مع إقرار أن التغطية الكاملة لهذه المخاطر تتجاوز الدليل المقاس في النسخة الحالية.")
    add_para(doc, "وتشير أدبيات هندسة البرمجيات التشغيلية إلى أن قابلية المراقبة والاختبار والتعافي جزء من قرار المعمارية، وليست إضافات لاحقة [29][30][31][32][33]. لذلك تضع الأطروحة نتائج اختبارات الوحدة في موضعها الصحيح، ولا تستبدل بها اختبارات حمل أو قبول ميداني غير منفذة.")
    add_heading(doc, "2.12 الأمن والحوكمة كمتطلبات تصميم", 2)
    add_para(doc, "توصي مراجع التحكم بالوصول بالجمع بين هوية المستخدم، دوره، وسياق الوصول بدلاً من التعامل مع الدور كقائمة صلاحيات مجردة [14][23]. ويطبق المشروع هذا المبدأ عملياً عبر الدور وclaim النطاق التنظيمي؛ وبذلك لا يكفي أن يكون المستخدم مدير فرع، بل يجب أن يحمل جذر الوحدة الذي يسمح له برؤية تفرعات محددة. وتبقى سياسات الاحتفاظ والنسخ والاستعادة قرارات مؤسسية يجب اعتمادها قبل الإنتاج [24][25].")
    add_para(doc, "تدعم معايير رأس المال البشري وأمن المعلومات هذا الفصل بين حوكمة البيانات والأمن التقني [26][34][37]. كما تفسر مراجع OAuth وOpenID Connect سبب فصل مسؤولية الهوية عن خدمات الأعمال في Keycloak والبوابة [35][36]، في حين تستخدم متطلبات OWASP مرجعاً لمراجعة التحقق من المدخلات، الجلسات، الإفصاح عن الأخطاء، وحماية المسارات الإدارية [27][28].")


def add_approved_usecase_package(doc: Document, table_no: int) -> int:
    """Insert the retained sector package as historical analysis evidence."""
    add_heading(doc, "ملحق (ج): الحزمة القطاعية السابقة لحالات الاستخدام", 2)
    try:
        analysis = json.loads(USECASE_ANALYSIS_SOURCE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        add_callout(doc, "تنبيه:", "تعذر تحميل جدول حالات الاستخدام المصحح، لذلك لم تضف مصفوفة التغطية.", fill="FFF7E6")
        return table_no

    include_total = sum(int(item.get("include_count", 0)) for item in analysis.get("groups", []))
    extend_total = sum(int(item.get("extend_count", 0)) for item in analysis.get("groups", []))
    add_para(doc, f"تحتفظ الأطروحة بهذه الحزمة لأنها توثق التحليل القطاعي السابق في {analysis.get('diagram_count', 0)} مخططاً و{analysis.get('case_count', 0)} حالة استخدام ضمن {analysis.get('department_count', 0)} إدارات عليا، وتتضمن {include_total} علاقة Include و{extend_total} علاقة Extend. لكنها ليست المرجع الوحيد لخط الأساس الحالي؛ إذ سبقت قدرات تعدد الجهات وإسنادها وحجرها، والهيكل المؤرخ، والسياسات والرموز، والرواتب والخزينة والملفات الحديثة. لذلك تعتمد الرسوم المرجعية في الفصل الثالث، وتستخدم هذه الحزمة لتتبع أصل العمليات والأدوار فقط.")
    rows = [["الشكل", "الإدارة والفرع", "التغطية التحليلية", "الخدمات وعائلات المتطلبات"]]
    for item, figure_key in zip(analysis.get("groups", []), USECASE_DRAWIO_MAP):
        rows.append([
            FIGURES[figure_key].label,
            f"{item.get('department', '-')}\n{item.get('branch', '-')}",
            f"{item.get('case_count', 0)} حالة | Include: {item.get('include_count', 0)} | Extend: {item.get('extend_count', 0)}",
            f"{item.get('services', '-')}\n{item.get('requirements', '-')}",
        ])
    add_table(doc, rows, "تغطية الحزمة القطاعية السابقة وربطها بعائلات المتطلبات", 3, table_no, label_override="الجدول (ج-1)")
    table_no += 1
    begin_portrait_section(doc)
    add_callout(doc, "قاعدة الاستخدام:", "تمثل الحالات الـ247 تفصيلاً قطاعياً مهماً، لكنها تقرأ مع الحزمة المرجعية الحالية في الفصل الثالث وسجل الحقيقة. لا يعد غياب قدرة حديثة من الرسم القديم دليلاً على غيابها من خط الأساس، ولا يعد وجود حالة قديمة دليلاً على اكتمال تنفيذها.", fill="FFF7E6")
    doc.add_page_break()
    for figure_key in USECASE_DRAWIO_MAP:
        add_figure(doc, figure_key)

    add_heading(doc, "ملحق (ج-2): تفسير العلاقات ونماذج الحالات الحرجة", 3)
    relation_rows = [
        ["العلاقة", "المعنى التحليلي", "قاعدة استخدامها", "مثال من الحزمة"],
        ["Association", "تفاعل الممثل مباشرة مع الحالة", "تستخدم عندما يبدأ الممثل الإجراء أو يراجعه", "قسم العقود يستعرض بيانات المتقدمين"],
        ["Include", "سلوك إلزامي يعاد استخدامه داخل حالة أصلية", "تنفذ الحالة المضمنة دائماً ضمن المسار الأصلي", "إنهاء الخدمة يتضمن إيقاف الأثر المالي والأرشفة"],
        ["Extend", "سلوك شرطي يضاف عند تحقق سبب محدد", "لا ينفذ في كل مرة، ويجب توثيق شرط التفعيل", "الإحالة القانونية تمتد من إنهاء الخدمة عند وجود نزاع"],
        ["لوحة/شاشة رئيسية", "نقطة دخول ومتابعة وليست تجميعاً إلزامياً لكل العمليات", "لا تربط تلقائياً بعلاقات Include مع كل الإجراءات", "متابعة التوظيف والخدمة والعقود"],
    ]
    add_table(doc, relation_rows, "التفسير المنهجي لعلاقات حالات الاستخدام", 3, table_no, label_override="الجدول (ج-2)")
    table_no += 1

    selected = [
        ("إدارة شؤون الكادر الإشرافي", "فرع الخدمة والعقود", "إنشاء إعلان التوظيف ونشره"),
        ("إدارة شؤون الكادر الإشرافي", "فرع الخدمة والعقود", "إنهاء خدمة الموظف"),
        ("إدارة شؤون الكادر الإشرافي", "فرع التنقلات والتوزيع", "تنفيذ نقل موظف"),
        ("إدارة السجلات المؤسسية", "فرع الرواتب والتعويضات", "استعراض تدقيق دورة صرف الرواتب"),
        ("إدارة شؤون الكادر الإشرافي", "فرع الشؤون الشخصية والصحية", "إنشاء ملف صحي للموظف"),
        ("إدارة شؤون الكادر الإشرافي", "فرع تخطيط الدورات وتأهيل القوى البشرية", "ترشيح موظف للدورة"),
        ("الفروع المركزية", "فرع بطاقات التعريف والهويات", "إصدار بطاقة إلكترونية"),
    ]
    cases = analysis.get("cases", [])
    for department, branch, title in selected:
        case = next((item for item in cases if item.get("department") == department and item.get("branch") == branch and item.get("title") == title), None)
        if case is None:
            continue
        fields = case.get("fields", {})
        relations = []
        if case.get("include"):
            relations.append("Include: " + "؛ ".join(case["include"]))
        if case.get("extend"):
            relations.append("Extend: " + "؛ ".join(case["extend"]))
        add_heading(doc, f"نموذج حالة: {title}", 3)
        add_metadata_table(doc, [
            ["الإدارة والفرع", f"{department} - {branch}"],
            ["الممثل", fields.get("المنفذ", "غير محدد")],
            ["نوع الحالة", fields.get("نوع الحالة", "غير محدد")],
            ["المدخلات", fields.get("المدخلات", "-")],
            ["المخرجات", fields.get("المخرجات", "-")],
            ["العلاقات", " | ".join(relations) if relations else "لا توجد علاقة Include/Extend موثقة"],
            ["التفسير", fields.get("الشرح", "-")],
        ])
        add_para(doc, "يرتبط هذا التوصيف بعائلة المتطلبات والخدمة المبينة في مصفوفة التغطية، وتحدد مصفوفة التتبع علاقته بالتنفيذ ومعيار القبول ونتيجة التحقق.", size=11.5, after=8)
    return table_no


def add_dfd_appendix(doc: Document) -> None:
    begin_landscape_section(doc)
    add_heading(doc, "ملحق (د): حزمة DFD القطاعية السابقة", 2)
    add_para(doc, "تحتفظ هذه الحزمة بمستويات DFD للعمليات الجوهرية كما كانت في خط الأساس السابق. وتحافظ على أثر تحليل التعيين والنقل والصحة والتدريب والرواتب والشكاوى والفصل، لكنها لا تمثل وحدها تعدد الجهات والحجر والسياسات والرموز والرواتب والخزينة والملفات الحديثة. حزمة DFD المرجعية الحالية معروضة في الفصل الثالث، ويقرأ هذا الملحق بوصفه تفصيلاً تاريخياً للعمليات لا بديلاً عنها.")
    keys = (
        "dfd_main", "dfd_new_employee", "dfd_transfer", "dfd_medical", "dfd_training",
        "dfd_payroll", "dfd_complaint", "dfd_separation", "dfd_dictionary",
    )
    for index, key in enumerate(keys):
        add_figure(doc, key, max_width=9.5, max_height=4.45)
        if index < len(keys) - 1:
            doc.add_page_break()
    begin_portrait_section(doc)


def capability_truth_rows() -> list[list[str]]:
    if not TRUTH_REGISTER_SOURCE.exists():
        return []
    lines = TRUTH_REGISTER_SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index + 1 < len(lines):
        line = lines[index].strip()
        if line.startswith("|") and re.match(r"^\|?\s*:?-{3,}", lines[index + 1].strip()):
            table_lines = [line]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            return markdown_rows(table_lines)
        index += 1
    return []


def add_capability_truth_register(doc: Document) -> None:
    rows = capability_truth_rows()
    if not rows:
        return
    begin_landscape_section(doc, horizontal_margin=1.1)
    add_heading(doc, "ملحق (ب): سجل الحقيقة الموحد لقدرات النظام", 2)
    add_para(doc, "يربط هذا السجل القدرة بحالتها ومالكها وكياناتها وحد التكامل ودليلها وموضع تحليلها. وهو المرجع الذي يحسم التعارض بين وجود مكون برمجي وبين اكتمال التحقق أو القبول المؤسسي؛ فلا يعرض كيان أو وثيقة تصميم أو اختبار متخطى بوصفه نتيجة ناجحة.")
    add_table(doc, rows, "سجل الحقيقة الموحد للقدرات والحدود والأدلة", 3, 1, label_override="الجدول (ب-1)")
    add_callout(doc, "قاعدة الحالة:", "تستخدم أربع حالات فقط: منفذ ومتحقق، منفذ ولم يكتمل التحقق، منفذ جزئياً، ومخطط أو مؤجل. ويعود تفسير كل حالة إلى الدليل وحدوده في السطر نفسه.", fill="FFF7E6")
    begin_portrait_section(doc)


def requirement_status(requirement_id: str, evidence: str = "") -> str:
    verified = {"FR-EMP-04", "FR-TRN-03", "FR-TEN-05"}
    partial = {"FR-ORG-08", "FR-LOG-02", "FR-LOG-03", "FR-LOG-04", "FR-OPS-05", "FR-OPS-06"}
    if requirement_id in verified:
        return "منفذ ومتحقق"
    if requirement_id in partial:
        return "منفذ جزئياً"
    return "منفذ ولم يكتمل التحقق"


def add_requirements_catalog(doc: Document) -> None:
    """Render the detailed, reviewable requirements catalog as part of chapter one."""
    add_heading(doc, "1.9.1 كتالوج المتطلبات التفصيلي", 3)
    add_para(doc, "يربط هذا القسم كل سلوك جوهري بممثل وقاعدة عمل وخدمة مالكة ومعيار قبول. تستخدم الحالة أربع درجات موحدة: منفذ ومتحقق، منفذ ولم يكتمل التحقق، منفذ جزئياً، ومخطط أو مؤجل. وتظهر فجوات البيئة والتكامل والقبول المؤسسي صراحة في الدليل، فلا تساوي كلمة التنفيذ الجاهزية الإنتاجية.")
    if not CATALOG_SOURCE.exists():
        add_callout(doc, "تنبيه:", "لم يعثر المولد على كتالوج المتطلبات التفصيلي.", fill="FFF7E6")
        return

    doc.add_section(WD_SECTION.NEW_PAGE)
    landscape = doc.sections[-1]
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width = Cm(29.7)
    landscape.page_height = Cm(21)
    landscape.right_margin = Cm(1.5)
    landscape.left_margin = Cm(1.5)
    landscape.top_margin = Cm(1.6)
    landscape.bottom_margin = Cm(1.6)
    landscape.header_distance = Cm(0.75)
    landscape.footer_distance = Cm(0.75)
    continue_page_numbering(landscape)
    add_running_header(landscape)
    add_page_number_footer(landscape)

    lines = CATALOG_SOURCE.read_text(encoding="utf-8").splitlines()
    table_no = 1
    index = 0
    current_title = ""
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith("## "):
            current_title = strip_markdown(line[3:])
            suffix = "1" if current_title == "المتطلبات الوظيفية" else "2"
            add_heading(doc, f"1.9.1.{suffix} {current_title}", 3)
            index += 1
            continue
        if line and not line.startswith("#") and not line.startswith("|") and line != "---":
            add_para(doc, strip_markdown(line), after=5)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-{3,}", lines[index + 1].strip()):
            table_lines = [line]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = markdown_rows(table_lines)
            title = "المتطلبات غير الوظيفية ومعايير تحققها" if "غير الوظيفية" in current_title else "كتالوج المتطلبات الوظيفية التفصيلي"
            if "غير الوظيفية" not in current_title and rows and len(rows[0]) == 7:
                rows = [rows[0][:-1] + ["الحالة", rows[0][-1]]] + [
                    row[:-1] + [requirement_status(row[0], row[-1]), row[-1]] for row in rows[1:]
                ]
            add_table(doc, rows, title, 1, table_no)
            table_no += 1
            continue
        index += 1

    add_heading(doc, "1.9.2 مصفوفة التتبع المختصرة", 3)
    add_para(doc, "تنقل المخططات إلى الفصل الثالث، بينما يثبت هذا الجدول منذ الفصل الأول أن كل متطلب حرج يملك حالة استخدام أو رحلة، ومكوناً مسؤولاً، ودليل تحقق أو قيداً معلناً. وترد المصفوفة الكاملة في الملحق.")
    rows = [
        ["المتطلب", "الحالة أو الرحلة", "التنفيذ", "دليل التحقق الحالي"],
        ["FR-TEN-03", "عزل الجهة", "TenantContext وTenantAwareDbContext", "اختبارات TEN؛ RLS الحية مؤجلة"],
        ["FR-ORG-05", "نطاق الفرع", "OrgScopeService وscope_org_unit_id", "اختبارات النطاق؛ تغطية اللوجستيات قيد الإغلاق"],
        ["FR-POL-04", "إنفاذ السياسة", "Policy API وPolicyEnforceClient", "اختبارات وحدة؛ بديل عميل التكامل مطلوب"],
        ["FR-CODE-03", "توليد الرمز", "Code API وسجل التوليد", "اختبارات وحدة؛ تهيئة التكامل مطلوبة"],
        ["FR-HIRE-05", "اكتمال التعيين", "عقود أحداث ومستهلكون", "اختبارات قواعد ومستهلكين"],
        ["FR-LOG-02", "عهدة الأصل", "AssetAssignment workflow", "اختبارات أمنية كشفت فجوة نطاق معلنة"],
        ["NFR-24", "سلامة الربط", "DTO/Commands واختبارات معمارية", "14 حالة لوجستية ما زالت فاشلة"],
    ]
    add_table(doc, rows, "خلاصة تتبع المتطلبات إلى التنفيذ والتحقق", 1, table_no)
    truth = capability_truth_rows()
    if len(truth) > 1:
        status_counts: dict[str, int] = {}
        for row in truth[1:]:
            if len(row) >= 3:
                status_counts[row[2]] = status_counts.get(row[2], 0) + 1
        add_heading(doc, "1.9.3 سجل الحقيقة الموحد", 3)
        add_para(doc, "يكمل سجل الحقيقة كتالوج المتطلبات بربط كل قدرة عابرة للمجالات بمالكها وكياناتها وواجهتها أو حدثها ودليلها وحد الاستدلال. ويعرض الملحق (ب) السجل كاملاً، بينما تلخص المصفوفة الآتية توزيع الحالات من دون تحويل العدد إلى مقياس جودة.")
        status_rows = [["الحالة", "عدد القدرات", "التفسير"]]
        explanations = {
            "منفذ ومتحقق": "يسنده دليل ناجح ضمن البيئة المعلنة",
            "منفذ ولم يكتمل التحقق": "المسار قائم وتبقى بيئة حية أو تكامل أو قبول",
            "منفذ جزئياً": "توجد قدرة جوهرية مع حلقة أو حارس أو اعتماد ناقص",
            "مخطط أو مؤجل": "لا يعرض بوصفه منجزاً في خط الأساس الحالي",
        }
        for status in ("منفذ ومتحقق", "منفذ ولم يكتمل التحقق", "منفذ جزئياً", "مخطط أو مؤجل"):
            status_rows.append([status, str(status_counts.get(status, 0)), explanations[status]])
        add_table(doc, status_rows, "توزيع قدرات سجل الحقيقة بحسب مستوى الإثبات", 1, table_no + 1)
    doc.add_section(WD_SECTION.NEW_PAGE)
    portrait = doc.sections[-1]
    configure_section(portrait)
    continue_page_numbering(portrait)
    add_running_header(portrait)
    add_page_number_footer(portrait)
    return

    add_heading(doc, "1.9.2 مخططات حالات الاستخدام وتدفق البيانات", 3)
    add_para(doc, "تجسد المخططات الآتية منظور التحليل قبل الانتقال إلى البنية البرمجية. وهي تفصل بين من يطلب الإجراء، ومن يعتمد القرار، ومن ينفذ الأثر داخل نطاق مؤسسي مضبوط.")
    add_figure(doc, "usecase")
    add_figure(doc, "dfd")
    table_no = add_approved_usecase_package(doc, table_no)
    add_heading(doc, "1.9.2.3 حزمة مخططات DFD المعتمدة", 3)
    add_para(doc, "تستكمل الصفحات التالية مخطط السياق بمخطط المستوى الأول ومخططات المستوى الثاني للعمليات الجوهرية، ثم تختم بقاموس البيانات. وتحافظ النماذج على ترابط الجهات الخارجية والعمليات ومخازن البيانات عبر مستويات التحليل المختلفة.", after=5)
    add_figure(doc, "dfd_main")
    for key in (
        "dfd_new_employee",
        "dfd_transfer",
        "dfd_medical",
        "dfd_training",
        "dfd_payroll",
        "dfd_complaint",
        "dfd_separation",
        "dfd_dictionary",
    ):
        add_figure(doc, key)

    add_heading(doc, "1.9.3 مصفوفة التتبع", 3)
    add_para(doc, "تمنع مصفوفة التتبع انفصال التحليل عن التنفيذ: لا يصنف متطلب أساسياً ما لم يمكن ربطه بحالة استخدام وخدمة أو قاعدة بيانات ودليل تحقق. تستخدم المصفوفة الكاملة في الملحق، بينما تعرض الخلاصة الآتية نقاط القوة الجوهرية.")
    rows = [
        ["المتطلب", "حالة الاستخدام", "التنفيذ", "دليل التحقق"],
        ["FR-ORG-05", "UC-SCOPE: نطاق الفرع", "OrgScopeService وscope_org_unit_id", "TS-01 إلى TS-04"],
        ["FR-HIRE-05", "UC-HIRE: استكمال التعيين", "HR.Contracts + أحداث التعيين", "HireWorkflowRules واختبارات المستهلك"],
        ["FR-MOV-03", "UC-MOVE: تنفيذ النقل", "TransferExecutedConsumer", "اختبارات الحركات ودورة الحياة"],
        ["FR-PAY-04", "UC-PAY: أثر معاملة", "PayrollAdjustment والمستهلكون", "PayrollPeriodRulesTests"],
        ["NFR-02", "UC-SCOPE: منع الفرع الشقيق", "سياسات الرؤية التنظيمية", "HrOrgScopeVisibilityTests"],
    ]
    add_table(doc, rows, "خلاصة تتبع المتطلبات إلى التنفيذ والتحقق", 1, table_no)


def catalog_functional_rows() -> list[list[str]]:
    """Return functional requirement rows without duplicating their content in code."""
    if not CATALOG_SOURCE.exists():
        return []
    lines = CATALOG_SOURCE.read_text(encoding="utf-8").splitlines()
    rows: list[list[str]] = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith("## المتطلبات غير الوظيفية"):
            break
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-{3,}", lines[index + 1].strip()):
            table_lines = [line]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows.extend(markdown_rows(table_lines)[1:])
            continue
        index += 1
    return rows


def catalog_nonfunctional_rows() -> list[list[str]]:
    if not CATALOG_SOURCE.exists():
        return []
    lines = CATALOG_SOURCE.read_text(encoding="utf-8").splitlines()
    rows: list[list[str]] = []
    index = 0
    active = False
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith("## المتطلبات غير الوظيفية"):
            active = True
            index += 1
            continue
        if active and line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-{3,}", lines[index + 1].strip()):
            table_lines = [line]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows.extend(markdown_rows(table_lines)[1:])
            continue
        index += 1
    return rows


def add_requirement_cards(doc: Document) -> None:
    add_heading(doc, "ملحق (هـ): بطاقات المتطلبات الوظيفية", 2)
    add_para(doc, "تقدم هذه البطاقات مستوى التحليل الذي لا يمكن أن تحمله الجداول العامة وحدها. كل بطاقة تحدد الممثل والسلوك وقاعدة العمل والخدمة المالكة ودليل القبول، وتتعامل مع غياب الشرط المسبق بوصفه منعاً للاعتماد أو للنشر مع حفظ أثر تدقيق مناسب.")
    current_domain = ""
    for row in catalog_functional_rows():
        if len(row) < 7:
            continue
        requirement_id, domain, actor, behavior, rule, owner, evidence = row[:7]
        if domain != current_domain:
            current_domain = domain
            add_heading(doc, f"مجال: {domain}", 3)
        heading = add_heading(doc, f"{requirement_id}: {behavior}", 3)
        keep_with_next(heading)
        add_metadata_table(doc, [
            ["الحالة", requirement_status(requirement_id, evidence)],
            ["الممثل", actor],
            ["السلوك المطلوب", behavior],
            ["الشرط أو قاعدة العمل", rule],
            ["الخدمة أو الحد المسؤول", owner],
            ["معيار القبول ووسيلة التحقق", evidence],
            ["الحالة البديلة", "عند عدم تحقق القاعدة يمنع الاعتماد أو النشر، وتحفظ المعاملة أو سبب الرفض ضمن الأثر التشغيلي المتاح."],
        ])
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)


def add_quality_cards(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "ملحق (و): بطاقات المتطلبات غير الوظيفية", 2)
    add_para(doc, "المتطلبات غير الوظيفية في هذا الملحق لا تحول أهداف الجودة إلى نتائج مزعومة. عند وجود تنفيذ أو اختبار يشار إليه كدليل، وعند غياب القياس تبقى القيمة هدف قبول أو بنداً مؤجلاً.")
    for row in catalog_nonfunctional_rows():
        if len(row) < 4:
            continue
        nfr_id, axis, requirement, evidence = row[:4]
        heading = add_heading(doc, f"{nfr_id}: {axis}", 3)
        keep_with_next(heading)
        measured = "دليل قائم أو تحقق قابل لإعادة التشغيل" if any(token in evidence for token in ("اختبار", "Health", "Docker", "Correlation", "Keycloak", "Outbox", "Message")) else "هدف أو قيد معلن حتى يجرى القياس"
        add_metadata_table(doc, [
            ["محور الجودة", axis],
            ["المطلب", requirement],
            ["نوع الدليل", measured],
            ["الدليل أو آلية القياس", evidence],
            ["حد التفسير", "لا تعد هذه البطاقة برهاناً على أداء إنتاجي أو امتثال قانوني كامل ما لم تتضمن نتيجة تشغيل موثقة."],
        ])
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)


def add_service_profiles(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "ملحق (ز): بطاقات ملكية البيانات للخدمات", 2)
    add_para(doc, "تعرض البطاقات الآتية ملكية كل خدمة للبيانات ودورها في نموذج المركز والفروع. وهي مكملة لمخططات ERD؛ إذ توضح حدود المجال وسبب منع الوصلات المباشرة بين قواعد بيانات الخدمات.")
    profiles = [
        ("HR.Core", "الجهة والموظف والهيكل المؤرخ والوظائف الممولة والتوظيف والسياسات وسير العمل والرموز والتهيئة", "Tenant, TenantAssignmentRule, TenantQuarantineEntry, Employee, OrgUnitVersion, FundedPosition, WorkforcePlan, PolicyTemplateVersion, WorkflowInstance, CodeIssuance", "EmployeeCreated, PolicyPublished, AgencyOnboarded", "مرجع الجهة والنطاق التنظيمي للمركز والفروع وحوكمة الإسناد والحجر"),
        ("HR.Contracts", "العقود والمباشرة والإجازات والسلف والفصل", "Contract, LeaveRequest, AdvanceRequest, SeparationRecord", "ContractActivated, SeparationCompleted", "تنفيذ معاملات الموظفين ضمن صلاحيات الخدمة"),
        ("Payroll", "إصدارات إعدادات الرواتب واللقطات والحساب القابل للاستئناف والتمويل والمتأخرات والدفع والتسوية ومكافحة الاحتيال", "PayrollConfiguration, PayrollRunInputSnapshot, PayrollRunJob, PayrollRunPartition, SalaryLine, FundingDecision, ArrearsLedgerEntry, PaymentOrder, BankStatementLine, FraudAlert, PaymentHold", "PayrollProcessed وأحداث الدفع والتسوية", "بيانات مالية لا تعرض إلا للدور المخول مع فصل الواجبات"),
        ("HR.Health", "السجل الصحي واللجان والسلامة", "MedicalRecord, MedicalRecommendation, WorkplaceAccident", "MedicalRecommendationIssued", "عزل صحي وصلاحيات مقيدة"),
        ("HR.Movements", "النقل والانتداب والتوزيع", "TransferRequest, SecondmentRequest, DistributionPlan", "TransferExecuted, SecondmentExecuted", "يغير الإسناد التنظيمي عبر حدث موثق"),
        ("HR.Training", "التدريب والتعلم الإلكتروني", "TrainingCourse, TrainingEnrollment, LearningCertificate", "TrainingCompleted", "خطة تدريب مرتبطة بالموظف والوحدة"),
        ("IDCard", "طلبات وإصدار وإبطال البطاقات", "CardRequest, IssuedCard, CardDeliveryRecord", "CardIssued", "يرتبط بالتعيين والفصل دون ملكية ملف الموظف"),
        ("Records.Archive", "الأرشفة والاحتفاظ والإتلاف", "ArchiveItem, RecordDisposal, LoanRecord", "HireCaseCompleted, SeparationCompleted", "حفظ الأثر الإداري للمعاملات"),
        ("Admin.Logistics", "العهد والأصول والمخزون والتوريد وبلاغات الفقد", "Asset, AssetAssignment, AssetLossReport, SupplyOrderLine", "AssetAssignmentApproved, TransferExecuted, SeparationCompleted", "استرداد العهد عند الحركة أو الفصل ضمن نطاق الجهة"),
        ("IT.Infra", "الأجهزة والدعم والنسخ والحوادث", "UserDevice, ITSupportTicket, BackupJob, SecurityIncident", "EmployeeCreated, SeparationCompleted", "تشغيل تقني مع فصل بيانات الموارد البشرية"),
        ("Dev.Analytics", "التقارير ولوحات القياس", "AnalyticsEvent, Dashboard, KPI, Report", "PayrollProcessed, TrainingCompleted", "يدعم رؤية مركزية بمؤشرات مجمعة"),
        ("Compliance", "الشكاوى والتدقيق والانضباط", "Complaint, AuditFinding, DisciplinaryAction", "DisciplinaryActionTaken", "بيانات حساسة بصلاحية امتثال"),
        ("Auth.Directory", "المستخدمون والدليل والأدوار", "AuthUser, DirectoryEntry, Role, AuthAuditLog", "EmployeeCreated, SecondmentExecuted", "ربط الهوية بالدور والنطاق التنظيمي"),
        ("File.Storage", "حجر الملفات والفحص والتنقية والترقية والمراجع والمنح وسجل الوصول", "FileRecord, FileScanWorkItem, FileScanResult, FileReferenceRecord, FileAccessGrant, FileAccessAuditLog", "مرجع FileId وحالة ScanStatus", "لا تملك قرار أعمال؛ تحفظ الملف وتفشل مغلقاً حتى النظافة ومنحة الوصول"),
    ]
    for service, responsibility, entities, events, branch_role in profiles:
        heading = add_heading(doc, service, 3)
        keep_with_next(heading)
        add_metadata_table(doc, [
            ["ملكية المجال", responsibility],
            ["الكيانات المحورية", entities],
            ["الأحداث أو المراجع", events],
            ["علاقته بالمركز والفروع", branch_role],
            ["قاعدة البيانات", "قاعدة مستقلة للخدمة؛ التكامل عبر المعرفات أو العقود وليس عبر FK عابر للخدمات."],
        ])
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)


def add_verification_protocol(doc: Document, new_page: bool = True) -> None:
    if new_page:
        doc.add_page_break()
    add_heading(doc, "ملحق (ط): بروتوكول التحقق وحالات القبول", 2)
    add_para(doc, "يوضح هذا الملحق كيفية تحويل المتطلبات إلى حالات تحقق قابلة لإعادة التشغيل. وتميز خانة الحالة بين اختبار نُفذ وسُجلت نتيجته، وبين حالة قبول مخططة تتطلب بنية حاويات أو بيانات تجريبية معتمدة قبل إدراجها ضمن النتائج.")
    cases = [
        ("VT-01", "بناء Release", "بناء الحل في وضع Release دون أخطاء تجميع", "dotnet build / Release", "ينفذ في بيئة .NET 10 ويوثق الخرج"),
        ("VT-02", "اختبارات الوحدة", "تشغيل قواعد المجال والسياسات والحراس المعمارية", "HRManagement.UnitTests", "2416 ناجح و5 فاشلة؛ لا يقبل الخط قبل إغلاق الفشل الأمني"),
        ("VT-03", "تغطية الاختبار", "جمع تقرير Cobertura لاختبارات الوحدة", "XPlat Code Coverage", "يعاد القياس بعد استقرار الخط؛ لا تستخدم نسبة قديمة"),
        ("VT-04", "اختبارات التكامل", "اختبار البنية والرسائل والعملاء المشتركين", "HRManagement.IntegrationTests", "158 ناجح و7 فاشلة و47 متخطاة"),
        ("VT-05", "تكامل PostgreSQL", "تشغيل Testcontainers وترحيل قاعدة اختبار", "RUN_TESTCONTAINERS=true", "يتطلب Docker daemon"),
        ("VT-06", "E2E البوابة", "تشغيل المسارات بعد إقلاع المكدس", "HRManagement.E2ETests", "2 ناجح و55 متخطاة؛ يتطلب المكدس الحي"),
        ("VT-07", "وصول المركز", "يرى الدور المركزي كل الوحدات وفق الصلاحية", "TS-04 + لقطة واجهة", "قبول عند غياب قيد فرعي"),
        ("VT-08", "نطاق الفرع", "يرى مدير الفرع الجذر والتفرعات فقط", "TS-01", "قبول عند تكوين مجموعة الوحدات الصحيحة"),
        ("VT-09", "منع الفرع الشقيق", "يرفض استعلام وحدة خارج مجموعة الفرع", "TS-03", "قبول عند الرفض وتسجيل الأثر"),
        ("VT-10", "نقل بين فرعين", "يتغير PrimaryOrgUnitId بعد اعتماد النقل", "TransferExecutedConsumer", "قبول عند تحديث المرجع ونشر الأثر"),
        ("VT-11", "تعيين متكامل", "تستكمل قائمة التعيين وتنشر الآثار", "HireWorkflowRules + consumers", "قبول عند وجود سجل وأحداث غير مكررة"),
        ("VT-12", "فصل متكامل", "تنفذ قائمة الفصل وتبلغ الخدمات التابعة", "SeparationCompleted", "قبول عند إغلاق العهد والهوية والأرشيف"),
        ("VT-13", "أداء القراءة", "قياس P50/P95 لقائمة موظفين مقيدة بالنطاق", "k6 أو أداة حمل موثقة", "هدف قبول فقط حتى ظهور نتيجة مقاسة"),
        ("VT-14", "أداء التعيين", "قياس المسار مع بيانات تجريبية معلنة", "سيناريو Gateway", "لا ينسب الأداء للإنتاج قبل القياس"),
        ("VT-15", "استعادة النسخ", "استعادة نسخة قاعدة خدمة في بيئة معزولة", "BACKUP_POSTGRES runbook", "قبول بعد توثيق زمن الاستعادة والبيانات"),
        ("VT-16", "إسناد الجهة والحجر", "تشغيل Dry Run ثم الإسناد والمصالحة على نسخة بيانات تجريبية", "TEN-003 assign/reconcile", "كل سجل منسوب أو محجور، والعدد والمراجع متطابقة"),
        ("VT-17", "استئناف الرواتب", "إيقاف job أثناء قسم ثم استئنافه ومنع سطر مكرر", "PAY-003 payroll job tests", "اكتمال الأقسام وتجميد النتائج مع حفظ checkpoints"),
        ("VT-18", "حجر الملفات", "فشل الماسح ومنع التنزيل ثم ترقية ملف نظيف", "FILE-002 scan lifecycle", "لا رابط قبل Clean ومنحة صالحة؛ تشغيل الماسح الحي مطلوب"),
        ("VT-19", "ثقة الخدمات", "رفض جمهور أو نطاق أو هوية غير مخولة لكل مسار S2S", "IAM-004 trust matrix", "قبول بعد تشغيل مصفوفة الرفض في بيئة Keycloak وقواعد ورسائل حية"),
    ]
    for case_id, name, objective, mechanism, status in cases:
        heading = add_heading(doc, f"{case_id}: {name}", 3)
        keep_with_next(heading)
        add_metadata_table(doc, [
            ["هدف التحقق", objective],
            ["الآلية أو الأداة", mechanism],
            ["معيار القبول", status],
            ["نطاق التحقق", "تشغيلي أو آلي بحسب الحالة؛ ولا يمثل قبولاً ميدانياً ما لم توثق عينة المستخدمين والبيانات وإجراءات التنفيذ."],
            ["الارتباط المنهجي", "يربط المتطلب الوظيفي أو غير الوظيفي بنتيجة قابلة للمراجعة أو بقيد معلن."],
        ])
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)


def add_submission_integrity_appendix(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "ملحق (ك): قائمة نزاهة النتائج والبيانات", 2)
    add_para(doc, "تلخص هذه القائمة ضوابط عرض النتائج وحماية البيانات، وتمنع الالتباس بين الوظائف المنفذة والمتطلبات المخططة، وبين التحقق الآلي والاعتماد الميداني أو القانوني.")
    rows = [
        ["عنصر المراجعة", "قاعدة القبول", "الحالة الموثقة"],
        ["نتائج الاختبار", "تذكر النتيجة والأمر والبيئة، لا العدد وحده", "2416/2421 للوحدة؛ التكامل 158 ناجحاً و7 فاشلة و47 متخطاة؛ E2E اختباران ناجحان و55 متخطاة"],
        ["اختبارات التكامل والأداء", "لا تعرض كمنجزة من دون سجل تشغيل", "معلنة كتحقق لاحق يحتاج Docker أو أداة حمل"],
        ["الرواتب", "يفصل بين قواعد تشغيلية ومحرك قانوني مصدق", "النسخة الحالية تشغيلية وقابلة للتهيئة، لا محرك قانوني نهائي"],
        ["المركز والفروع", "يثبت بالسياسة والنطاق واختبارات المنع", "موثق بـ OrgScopeService وTS-01..04"],
        ["البيانات الحساسة", "لا تستخدم بيانات أشخاص حقيقيين في اللقطات أو الاختبار", "تعتمد بيانات Seed أو بيانات مجهولة"],
        ["المراجع", "كل ادعاء نظري أو معياري يملك استشهاداً", "IEEE مع مراجع تقنية وأكاديمية ومعيارية"],
        ["الرسوم والجداول", "كل رسم بعنوان ومصدر وقابل للتتبع", "مرتبطة بنماذج التحليل والتصميم المعتمدة"],
        ["حدود النسخة الحالية", "تظهر في المتن والخاتمة والملاحق", "موثقة ولا تقدم كميزات مكتملة"],
    ]
    add_table(doc, rows, "قائمة تدقيق نزاهة النتائج والبيانات", 3, 22, label_override="الجدول (ك-1)")


def add_execution_log_appendix(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "ملحق (ل): سجل تشغيل التحقق", 2)
    add_para(doc, "يسجل هذا الملحق إجراءات التحقق التي اعتمدت عليها النتائج، بما يتيح إعادة الاختبار وتمييز النتائج المقاسة عن الجوانب المخططة. نُفذت الإجراءات في بيئة Windows باستخدام .NET 10 وفي وضع Release.")
    rows = [
        ["المعرف", "الأمر أو الأداة", "النتيجة الفعلية", "ملاحظة تفسيرية"],
        ["RUN-01", "dotnet build HRManagement.slnx -c Release --no-restore", "نجح: 0 أخطاء و4 تحذيرات", "خط الأساس الحالي يبنى في وضع Release"],
        ["RUN-02", "HRManagement.UnitTests", "2416 ناجح و5 فاشلة", "الفشل يتضمن حراساً أمنية وfixture مؤقتاً منتهي الصلاحية"],
        ["RUN-03", "HRManagement.IntegrationTests", "158 ناجح و7 فاشلة و47 متخطاة", "تهيئة مصادقة وبدائل Policy/Code مطلوبة"],
        ["RUN-04", "HRManagement.E2ETests", "2 ناجح و55 متخطاة", "المكدس الحي غير متاح"],
        ["RUN-05", "docker version", "العميل موجود وdaemon غير متاح", "لم تشغل Testcontainers أو RLS الحية"],
        ["RUN-06", "قياس الأداء", "غير منفذ", "لا يعرض زمن استجابة أو حمل مفترض"],
        ["RUN-07", "Open XML validation + accessibility audit", "ينفذ بعد توليد النسخة النهائية", "تحقق بنيوي من مستند التسليم"],
        ["RUN-08", "Word Export + PDF review", "ينفذ بعد تحديث الحقول", "تحقق بصري من RTL والجداول والمخططات"],
    ]
    add_table(doc, rows, "سجل النتائج القابلة لإعادة التشغيل", 3, 23, label_override="الجدول (ل-1)")


def add_formal_design_models(doc: Document) -> None:
    """Insert the current analysis and design package in chapter three."""
    add_heading(doc, "3.4.1 مخططات التحليل والتصميم الرسمية", 3)
    add_para(doc, "تأتي هذه المخططات بعد تثبيت المشكلة والمتطلبات والإطار النظري، ولذلك توضع في الفصل التطبيقي لا في الفصل الأول. تعرض الخرائط الشاملة في المتن، بينما تنقل الحزم القطاعية الكبيرة إلى الملاحق مع إحالة صريحة إليها.")
    add_heading(doc, "3.4.1.1 السياق وحالات الاستخدام المرجعية", 3)
    add_para(doc, "أعيد بناء الحزمة المرجعية الحالية بعد توسع النظام. يبدأ المخطط الشامل بالممثلين والقدرات الكبرى، ثم تفصل المخططات القطاعية إسناد الجهة والهيكل المؤرخ والسياسة وسير العمل والرواتب والخزينة واللوجستيات والملفات. وتطابق حزمة DFD هذه الحدود نفسها، لذلك لا تعتمد الأطروحة على الرسوم القطاعية السابقة بوصفها وصفاً كاملاً للنسخة الحالية.")
    for key in (
        "usecase_current_master", "usecase_tenant_assignment", "usecase_temporal_org",
        "usecase_policy_workflow", "usecase_advanced_payroll", "usecase_logistics_lineage",
        "usecase_security_files", "usecase_lifecycle", "usecase_people", "usecase_portals",
    ):
        add_figure(doc, key)
    begin_landscape_section(doc)
    add_heading(doc, "3.4.1.2 حزمة DFD المرجعية الحالية", 3)
    dfd_keys = (
        "dfd_current_level0", "dfd_tenant_assignment", "dfd_policy_workflow",
        "dfd_advanced_payroll", "dfd_logistics_lineage", "dfd_file_security",
    )
    for index, key in enumerate(dfd_keys):
        add_figure(doc, key, max_width=9.45, max_height=5.55)
        if index < len(dfd_keys) - 1:
            doc.add_page_break()
    begin_portrait_section(doc)
    add_callout(doc, "الحزمة التفصيلية:", "تقدم مخططات الفصل الثالث الحزمة المرجعية الحالية. ويعرض الملحق (ج) الحزمة القطاعية السابقة لحالات الاستخدام، بينما يعرض الملحق (د) مستويات DFD السابقة لأغراض التتبع التاريخي.")

    add_heading(doc, "3.4.1.3 مخططات النشاط", 3)
    add_para(doc, "تمثل مخططات النشاط نقاط القرار والتحقق والمسارات البديلة. ولا ينشر الأثر العابر للخدمات قبل اكتمال الشرط أو الاعتماد الذي تحدده رحلة العمل.")
    for key in ("activity_hire", "activity_contract", "activity_leave", "activity_movement", "activity_separation", "activity_payroll", "activity_onboarding", "activity_tenant_reconciliation", "activity_policy", "activity_asset"):
        add_figure(doc, key)

    add_heading(doc, "3.4.1.4 مخططات التسلسل", 3)
    add_para(doc, "تميز مخططات التسلسل بين قرار متزامن يحتاج نتيجة فورية، مثل إنفاذ السياسة أو توليد الرمز، وبين حدث غير متزامن يعلن حقيقة ثبتت محلياً.")
    for key in ("sequence_hire", "sequence_contract", "sequence_movement", "sequence_separation", "sequence_payroll"):
        add_figure(doc, key)
    begin_landscape_section(doc)
    add_figure(doc, "sequence_payroll_resumable", max_width=9.35, max_height=5.45)
    begin_portrait_section(doc)
    for key in ("sequence_policy", "sequence_code", "sequence_asset"):
        add_figure(doc, key)

    add_heading(doc, "3.4.1.5 نموذج المجال والحدود المعمارية", 3)
    add_para(doc, "يوضح نموذج المجال أن Tenant وOrgUnit وEmployee مراجع نواة، بينما تبقى حقائق العقد والراتب والصحة والحركة والعهد في سياقاتها. وتستخدم الخدمات معرفات وأحداثاً من دون مفاتيح أجنبية عابرة للقواعد.")
    for key in ("domain", "erd", "component", "deployment_formal", "security", "event_flow", "ownership", "reliability_journeys", "scope", "rbac", "trace"):
        add_figure(doc, key)
    rows = [
        ["المجال", "الكيانات المحورية", "ملكية البيانات", "مرجع التكامل"],
        ["HR.Core", "Tenant, TenantAssignmentRule, TenantQuarantineEntry, OrgUnitVersion, FundedPosition, PolicyVersion, WorkflowInstance", "hr_core_db", "TenantId وEmployeeId وOrgUnitId وإصدارات الحوكمة"],
        ["HR.Contracts", "Contract, LeaveRequest, SeparationRecord", "hr_contracts_db", "EmployeeId وأحداث التعيين/الفصل"],
        ["Payroll", "PayrollRunJob, InputSnapshot, SalaryLine, FundingDecision, ArrearsLedgerEntry, PaymentOrder, FraudAlert", "payroll_db", "EmployeeId وإصدار القاعدة وأحداث الأثر المالي"],
        ["HR.Health", "MedicalRecord, MedicalRecommendation", "hr_health_db", "EmployeeId بصلاحيات صحية"],
        ["Admin.Logistics", "Asset, AssetAssignment, AssetLossReport", "admin_db", "EmployeeId وAssetTag وأحداث العهد"],
    ]
    add_table(doc, rows, "ملكية البيانات والمراجع التكاملية في المجالات الجوهرية", 3, 5)
    add_callout(doc, "قاعدة تفسير المخططات:", "صلاحية الدور لا تكفي وحدها لإتاحة البيانات. فمدير الفرع يحتاج الدور المناسب مع جذر تنظيمي صالح؛ وبعد ذلك توسع الخدمة النطاق إلى التفرعات وتمنع الوحدات الشقيقة.")


def add_test_evidence(doc: Document) -> None:
    add_heading(doc, "3.13 الاختبارات والتحقق", 2)
    add_para(doc, "أعيد تشغيل البناء وحزم الاختبار على خط الأساس الحالي بتاريخ 2026/08/16. تعرض النتائج كما ظهرت من دون تحويل الاختبار المتخطى إلى ناجح أو إخفاء الاختبار الأمني الفاشل. وتعني نتيجة Release أن الشيفرة بُنيت في وضع الإصدار، ولا تعني وحدها جاهزية التشغيل الإنتاجي.")
    add_code(doc, "dotnet build HRManagement.slnx --configuration Release --no-restore\ndotnet test src/Tests/HRManagement.UnitTests/HRManagement.UnitTests.csproj --configuration Release --no-restore\ndotnet test src/Tests/HRManagement.IntegrationTests/HRManagement.IntegrationTests.csproj --configuration Release --no-restore\ndotnet test src/Tests/HRManagement.E2ETests/HRManagement.E2ETests.csproj --configuration Release --no-restore")
    rows = [
        ["عنصر التحقق", "النتيجة", "الدلالة"],
        ["بناء الحل", "نجح: 0 أخطاء و4 تحذيرات", "يمر خط البناء الحالي في وضع Release"],
        ["اختبارات الوحدة", "2416 ناجح، 5 فاشلة، 0 متخطى؛ المجموع 2421", "الفشل دليل حراسة أمني/معماري أو بيانات اختبار، وليس نتيجة مخفية"],
        ["اختبارات التكامل", "158 ناجح، 7 فاشلة، 47 متخطاة؛ المجموع 212", "تحتاج الإخفاقات تهيئة مخطط مصادقة وبدائل عملاء Policy/Code Generation"],
        ["اختبارات E2E", "2 ناجح، 0 فاشل، 55 متخطاة؛ المجموع 57", "المسارات التي تتطلب المكدس الحي لم تشغل"],
        ["Docker daemon", "غير متاح", "لم تشغل Testcontainers ولا اختبارات PostgreSQL RLS الحية"],
        ["اختبار الأداء", "غير منفذ", "لا يقدم التقرير زمناً أو حملاً مفترضاً"],
        ["اختبار القبول الميداني", "غير منفذ", "لا تستنتج سهولة الاستخدام أو رضا المستخدم من الاختبارات الآلية"],
    ]
    add_table(doc, rows, "نتائج البناء والاختبارات الحالية", 3, 6)
    add_figure(doc, "tests")
    add_para(doc, "كشفت اختبارات الوحدة الحالية غياب حارس النطاق التنظيمي عن استعلام لعهد الأصول، واكتشفت مجموعة اختبار الربط الواسع 14 إجراء في Admin.Logistics تحتاج أوامر أو نماذج ربط أضيق. كما انتهت صلاحية دليل STR-005 المؤقت الذي أعد للاختبار؛ ويعالج بتجديد بيانات الاختبار أو تثبيت الساعة، ولا يصنف وحده كخلل في وظيفة إنتاجية. أما اختبارات التكامل فتحتاج تسجيل بدائل IPolicyEnforceClient وICodeGenerationClient ومخطط مصادقة للاختبار قبل إعادة التشغيل.")
    add_para(doc, "اعتمد الفريق بعض التسهيلات والبدائل مؤقتاً لأغراض الاختبار، ولذلك لا تعاملها الأطروحة كقرارات إنتاجية ولا تستنتج منها ضعفاً وظيفياً مستقلاً. وفي الوقت نفسه تبقى نتيجة الاختبار كما سجلت: لا يحول القصد المؤقت الاختبار الفاشل إلى ناجح، بل يوثق سبب الفشل وشرط إعادة التحقق عند تثبيت بيئة الاختبار المستهدفة.")

    add_heading(doc, "3.13.1 التحقق من نموذج الجهة والمركز والفروع", 3)
    add_para(doc, "يمثل العزل متعدد المستويات القيمة الفارقة للمشروع، ولذلك تفصل المصفوفة بين السلوك الذي نجح في اختبارات النواة وبين المسار الذي ما زال يحتاج إغلاقاً أو بيئة PostgreSQL حية.")
    rows = [
        ["المعرّف", "الحالة", "السلوك المتوقع", "الدليل", "النتيجة"],
        ["TS-01", "إدارة مركزية", "الوصول إلى وحدات الجهة كلها من دون عبور جهة أخرى", "اختبارات Tenant وOrgScope", "ناجح ضمن الوحدة"],
        ["TS-02", "مدير فرع", "إرجاع الجذر وكل التفرعات فقط", "OrgScopeService وHrOrgScopeVisibility", "ناجح ضمن الوحدة"],
        ["TS-03", "فرع شقيق", "رفض وحدة خارج مجموعة الفرع", "اختبارات السماح والمنع", "ناجح ضمن الوحدة"],
        ["TS-04", "جهة أخرى", "رفض tenant مختلف حتى مع معرف وحدة صالح", "Tenant query filters", "ناجح مع تحقق RLS حي مؤجل"],
        ["TS-05", "عهد الأصول", "تصفية الاستعلام بالنطاق التنظيمي", "اختبار الحارس المعماري", "فاشل ويحتاج إصلاحاً"],
        ["TS-06", "RLS", "رفض الاستعلام داخل PostgreSQL عند tenant مخالف", "Testcontainers/PostgreSQL", "غير منفذ لغياب Docker"],
    ]
    add_table(doc, rows, "حالات التحقق من نموذج الجهة والمركز والفروع", 3, 7)
    add_callout(doc, "قرار القبول:", "لا يعد عزل الجهات والفروع مقبولاً للإنتاج قبل نجاح TS-05 وتشغيل TS-06 وتوسيع العينة إلى الاستعلامات الحساسة في الخدمات الأربع عشرة.", fill="FFF7E6")


def add_results_section(doc: Document) -> None:
    add_heading(doc, "3.14 النتائج", 2)
    requirement_count = len(catalog_functional_rows())
    capability_count = max(0, len(capability_truth_rows()) - 1)
    add_para(doc, f"حقق المشروع منصة موارد بشرية متعددة الجهات تربط الحوكمة المركزية بالتنفيذ المقيد في الفروع. يضم خط الأساس 14 خدمة مجال و15 سياق بيانات و274 DbSet و72 واجهة لعقود الرسائل، ويوثق {requirement_count} متطلباً وظيفياً و{capability_count} قدرة في سجل الحقيقة. وتجمع الحزمة المرجعية الحالية حالات الاستخدام وDFD للجهة والحجر والهيكل المؤرخ والسياسات وسير العمل والرواتب والخزينة واللوجستيات والملفات، بينما تحتفظ الحزمة القطاعية السابقة بـ247 حالة بوصفها أثراً تحليلياً. تثبت هذه الأرقام اتساع السطح، لكن نتائج الاختبار وحدود القبول هي التي تحدد مستوى الثقة ولا تسمح بتحويل الاتساع إلى ادعاء جاهزية.")
    rows = [
        ["البعد", "النتيجة المثبتة", "ما لا تدعيه الدراسة"],
        ["الحوكمة التنظيمية", "Tenant للجهة وOrgScope للمركز والفروع", "قبول مؤسسي حي أو اكتمال RLS قبل تشغيله"],
        ["حوكمة نقل البيانات", "قواعد إسناد وحجر ومصالحة للسجلات غير المنسوبة", "تشغيل ترحيل مؤسسة حية وتوقيع تقارير قبل وبعد"],
        ["الهيكل والملاك", "OrgUnitVersion ودليل وظائف ووظائف ممولة وخطة قوى عاملة", "اعتماد مالي خارجي أو قبول ميداني للملاك"],
        ["حوكمة القواعد", "سياسات حية وإصدارات ومحاكاة ودفتر إنفاذ", "اعتماد كل سياسة قانونياً أو إدارياً"],
        ["الرواتب والخزينة", "إصدارات ولقطات وتشغيل قابل للاستئناف وتمويل ومتأخرات وأوامر وحجوزات", "محرك قانوني مصدق أو قنوات دفع ومصارف حية أو أداء مقاس"],
        ["العزل المعماري", "خدمات وقواعد مستقلة وعقود وأحداث", "انعدام التأخر أو التعارض في كل سيناريو فشل"],
        ["الجودة الوظيفية", "2416 اختبار وحدة و158 اختبار تكامل ناجحاً", "خلو خط الأساس من الفشل أو تغطية كاملة"],
        ["التجربة العربية", "واجهات RTL وتقارير PDF عربية بخط مضمّن", "دراسة قابلية استخدام مع مستخدمين فعليين"],
        ["التشغيل", "Compose وصور موحدة وHealth checks", "نشر إنتاجي حي أو أداء وRTO/RPO مقاسة"],
    ]
    add_table(doc, rows, "حدود تفسير نتائج المشروع", 3, 8)


def add_governance_section(doc: Document) -> None:
    add_heading(doc, "3.19 حوكمة البيانات وقرار الجاهزية", 2)
    add_para(doc, "تتعامل المنظومة مع بيانات الهوية والصحة والرواتب والشكاوى والملفات، ولذلك لا تكفي المصادقة وحدها. يجمع التصميم بين Keycloak وسياسات الدور والجهة والنطاق، وTenantAwareDbContext وRLS، وتشفير الرقم الوطني، ودورة فحص الملفات ومنح الوصول، ودفاتر التدقيق والإنفاذ. وتتطلب الجاهزية سياسة احتفاظ ومفاتيح وأسرار ونسخاً مجربة ومراجعة صلاحيات واختبارات عزل حية.")
    rows = [
        ["القيد أو الخطر", "الحالة الحالية", "إجراء قبل الإنتاج"],
        ["حارس نطاق عهد الأصول", "اختبار أمني فاشل", "تطبيق OrgScope على الاستعلام وإعادة مجموعة الاختبار"],
        ["Mass Assignment في اللوجستيات", "14 إجراء كشفها الاختبار", "استخدام أوامر وDTO محددة وAllow-list للربط"],
        ["بدائل Policy/Code في التكامل", "تهيئة الاختبار ناقصة", "تسجيل fakes موحدة واختبار حالات السماح والمنع والفشل"],
        ["دليل STR-005 المؤقت", "انتهت صلاحيته المقصودة للاختبار", "تثبيت الساعة أو تجديد fixture من دون تغيير قاعدة العمل"],
        ["RLS وTestcontainers", "غير مشغلة لغياب Docker daemon", "تشغيلها في CI على PostgreSQL فعلي"],
        ["إسناد الجهة والحجر", "النموذج ودليل التشغيل موجودان بلا ترحيل قبول حي", "Dry Run ونسخة احتياطية ومصالحة وتوقيع أصحاب البيانات"],
        ["الرواتب والخزينة", "المسارات والسجلات الداخلية موجودة", "اعتماد قانوني ومالي واختبار أداء وتكامل قناة دفع وتسوية حي"],
        ["ثقة الخدمات وفحص الملفات", "جرد هويات ونطاقات ودورة حجر وفحص موجودة", "اختبار رفض شامل وتشغيل ماسح مؤسسي وانتهاء المنح"],
        ["قياس الأداء والقبول", "غير منفذ", "اختبار حمل معلن وPilot ببيانات تجريبية معتمدة"],
        ["الاحتفاظ والاستعادة", "آليات ووثائق موجودة دون تمرين قبول كامل", "سياسة زمنية وتمرين استعادة رحلة عمل"],
    ]
    add_table(doc, rows, "القيود المعلنة وخطة الإغلاق", 3, 9)
    add_callout(doc, "قرار الجاهزية:", "المنصة أساس مؤسسي واسع وليست جاهزة بعد للنشر الإنتاجي العام. يصبح القرار قابلاً للمراجعة بعد إغلاق الاختبارات الفاشلة وتشغيل العزل الحي والأداء والاستعادة والقبول الميداني.", fill="FFF7E6")


def add_database_design_section(doc: Document) -> None:
    """Document the database topology and the integration rationale."""
    add_heading(doc, "3.5.1 منهجية تصميم قواعد البيانات وملكية البيانات", 3)
    add_para(doc, "اختيرت في HRManagement منهجية قاعدة بيانات مملوكة لكل خدمة (Database per Service) ضمن سياقات محدودة (Bounded Contexts). والسبب ليس توزيع الجداول لمجرد التوزيع، بل فصل دورة حياة مجالات مختلفة الحساسية والتغير: بيانات الموظف والهيكل التنظيمي في HR.Core، العقود والإجازات في HR.Contracts، الرواتب في Payroll، والملفات الصحية في HR.Health. يتيح هذا الفصل امتلاك كل خدمة لنموذجها وترحيلاتها وصلاحياتها، ويقلل أثر تغيير مخطط مجال على مجال آخر، ويدعم النشر المستقل والتوسع التدريجي ومبدأ أقل صلاحية.")
    add_para(doc, "لا يعني هذا الاختيار أن كل خدمة تعمل بمعزل عن بقية المنظومة. فالخدمات لا تنفذ استعلامات مباشرة إلى قواعد بعضها ولا تنشئ مفاتيح أجنبية عابرة للحدود. بدلاً من ذلك تستخدم EmployeeId وOrgUnitId والمعرفات المرجعية الثابتة، وقد تحتفظ الخدمة التابعة بلقطة تنظيمية مثل OrgUnitSnapshot عندما تحتاج قراءة مستقرة. وتبقى ملكية السجل الأصلي في الخدمة التي تعرف قواعده، بينما تتولى الخدمات الأخرى آثارها المحلية وفق عقود التكامل.")
    add_para(doc, "اعتمد التصميم PostgreSQL موفراً افتراضياً لقواعد البيانات التطبيقية عبر Npgsql وترحيلات مستقلة لكل سياق. ويستخدم SQLite في بعض اختبارات التكامل أو التشغيل المحلي الخفيف عند تمرير سلسلة اتصال من نمط Data Source=. يخدم هذا البديل أغراض الاختبار، في حين يبقى PostgreSQL المحرك المستهدف للنشر، مع ضرورة اعتماد سياسات الترحيل والنسخ الاحتياطي والاستعادة في البيئة الفعلية.")
    begin_landscape_section(doc)
    add_figure(doc, "database_landscape", max_width=9.25, max_height=6.15)
    try:
        inventory = json.loads(DATABASE_INVENTORY_SOURCE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        inventory = []
    rows = [["الخدمة", "السياق وقاعدة البيانات", "DbSet / جداول الأعمال", "جداول البنية"]]
    for item in inventory:
        rows.append([
            item.get("service_ar", item.get("service", "-")),
            f"{item.get('context', '-')}\n{item.get('database', '-')}",
            f"{item.get('dbset_count', '-')} / {item.get('business_table_count', '-')}",
            ", ".join(item.get("infrastructure_tables", [])) or "-",
        ])
    add_table(doc, rows, "توصيف قواعد البيانات وسياقات ملكية البيانات", 3, 8)
    add_para(doc, "يلخص الجدول نموذج البيانات النهائي لكل سياق، ويبين مجموعات الكيانات وجداول الأعمال وجداول البنية المشتركة. ويمثل هذا التوصيف التصميم المنطقي المنفذ، في حين يتطلب إثبات جاهزية النشر تنفيذ الترحيلات واختبارات الاستعادة في بيئة تشغيل معتمدة.")
    doc.add_page_break()
    add_figure(doc, "database_hrcore", max_width=9.25, max_height=6.15)
    doc.add_page_break()
    add_figure(doc, "database_contracts", max_width=9.25, max_height=6.15)
    doc.add_page_break()
    add_figure(doc, "database_payroll", max_width=9.25, max_height=6.15)
    begin_portrait_section(doc)

    add_heading(doc, "3.5.2 آلية التواصل بين قواعد البيانات", 3)
    add_para(doc, "تتواصل قواعد البيانات بصورة غير مباشرة عبر طبقة التكامل القائمة على RabbitMQ وMassTransit. تنشر الخدمة المالكة للحقيقة حدثاً بعد تثبيت معاملتها المحلية، مثل IEmployeeCreated أو IContractActivated أو ITransferExecuted أو IPayrollProcessed. تستهلك الخدمات المعنية الحدث وتكتب أثرها في قاعدتها الخاصة. وبذلك يصبح التواصل موجهاً بالحقائق والأحداث، لا بربط الجداول أو تنفيذ Join عابر لقواعد البيانات.")
    add_para(doc, "تدعم بعض السياقات نمطي Outbox وInbox لتقليل احتمال فقد الرسالة أو تكرار أثرها: يسجل المنتج رسالة الخروج ضمن معاملته، ثم ينشرها عامل النقل، بينما يحتفظ المستهلك بحالة الرسالة أو معرفها قبل تطبيق الأثر. ومع ذلك فإن وجود الكيانات البرمجية لا يكفي لإثبات سلامة تشغيلية كاملة؛ إذ يلزم في بيئة الإنتاج اختبار فشل الوسيط، إعادة المحاولة، التعارض، والتسوية بعد الانقطاع.")
    add_para(doc, "توجد حدود تكامل أخرى واضحة: يحتفظ File.Storage ببيانات وصفية للملف بينما يخزن MinIO المحتوى الثنائي، ويؤدي Keycloak دور مخزن الهوية الخارجي، وتملك Gateway قاعدة مستقلة لتدقيق طلبات البوابة. أما الأرشيف والتحليلات فيستهلكان أحداثاً أو معرفات مرجعية دون امتلاك السجلات المصدرية للموظف أو العقد أو الراتب.")
    add_figure(doc, "database_integration")
    add_callout(doc, "قرار التصميم وحدوده:", "تحقق قاعدة بيانات لكل خدمة استقلالية المجال وتقليل الصلاحيات، لكنه يفرض اتساقاً نهائياً وحاجة إلى مراقبة الرسائل والتسوية. لذلك يوصى قبل الإنتاج بإضافة اختبارات تكامل للحالات الفاشلة، وقياس زمن انتشار الأحداث، وتمارين نسخ احتياطي واستعادة لكل قاعدة.", fill="FFF7E6")
    add_heading(doc, "3.5.3 تقويم الطوبولوجيا بحسب رحلات الأعمال", 3)
    add_para(doc, "طبقت الدراسة النقدية لقاعدة بيانات لكل خدمة على الرحلات التي تحدد نجاح الأعمال، لا على عدد الجداول. وتفصل بين الملكية المنطقية، التي تحققت عبر السياقات والصلاحيات والترحيلات، وبين استقلال مجال الفشل الفيزيائي الذي تحدده بيئة النشر ولم يثبت في التشغيل الحالي. كما تميز بين نجاح المعاملة المحلية واكتمال الأثر العابر للخدمات ضمن مهلة مقبولة.")
    reliability_rows = [
        ["الرحلة", "البيانات المالكة", "التبعيات الإلزامية", "إمكان الاستمرار الجزئي", "اختبار القبول المطلوب"],
        ["التعيين", "HR.Core وHR.Contracts", "العقد وقائمة التحقق", "يمكن تأخير التحليلات والتدريب مع إظهار الأثر المعلق", "فصل الوسيط ثم إعادة النشر ومنع التكرار"],
        ["النقل", "HR.Movements ثم HR.Core", "اعتماد الحركة وتحديث الإسناد", "تتوقف قرارات النطاق حتى اتساق الوحدة", "منع ظهور الموظف في فرعين واستعادة الإسقاط"],
        ["الفصل", "HR.Contracts", "قرار الإنهاء", "تستمر القراءة مع تنبيه آثار البطاقة والراتب والعهد", "تسوية كل مستهلك وإثبات الإغلاق"],
        ["الرواتب", "Payroll", "الإصدار واللقطة والحساب والاعتماد", "يمكن فصل الدفع عن الترحيل", "استئناف job وتسوية الأوامر والمتأخرات والعكس"],
        ["الملفات", "File.Storage وMinIO", "الحجر والفحص والمنحة", "تبقى المعاملة بلا ملف متاح حتى النظافة", "فشل الماسح ومنع التنزيل وانتهاء المنحة"],
    ]
    begin_landscape_section(doc)
    add_table(doc, reliability_rows, "تطبيق دراسة الاعتمادية على رحلات النظام الحرجة", 3, 9)
    begin_portrait_section(doc)
    add_para(doc, "لا تتضمن النتائج الحالية حقن أعطال أو قياس احتمال نجاح رحلة أو زمن استعادة. لذلك تبقى فرضيات انخفاض نطاق العطل أو زيادة كلفة التبعيات قابلة للاختبار، ولا تعرض كحقيقة مقاسة. ويلزم عند النشر توثيق منصة استضافة كل قاعدة، وارتباط مجالات الفشل، وتزامن النسخ، وترتيب استعادة الأحداث والملفات.")
    add_heading(doc, "3.5.4 خرائط ERD المختصرة للخدمات المساندة", 3)
    add_para(doc, "بعد عرض ERD المختصر للخدمات الجوهرية HR.Core وHR.Contracts وPayroll، تعرض الأشكال التالية الخرائط المنطقية المختصرة لبقية سياقات البيانات. تركز الرسوم على أسماء الجداول والمفاتيح والعلاقات الداخلية، بينما يقدم ملحق (ح) النموذج الكامل لكل قاعدة بيانات.")
    for key in DATABASE_COMPACT_ERD_MAP:
        if key not in {"database_hrcore", "database_contracts", "database_payroll"}:
            add_figure(doc, key)


def add_complete_database_erdds(doc: Document) -> None:
    """Add the complete one-context-per-database ERD package in landscape pages."""
    try:
        manifest = json.loads(DATABASE_FULL_ERD_MANIFEST.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        manifest = []
    if not manifest:
        return
    doc.add_section(WD_SECTION.NEW_PAGE)
    section = doc.sections[-1]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.right_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    continue_page_numbering(section)
    add_running_header(section)
    add_page_number_footer(section)
    add_heading(doc, "ملحق (ح): مخططات ERD الكاملة لكل قواعد البيانات", 2)
    add_para(doc, "يضم هذا الملحق مخططاً مستقلاً لكل سياق من سياقات البيانات الخمسة عشر. تعكس المخططات نموذج البيانات النهائي، وتعرض الجداول والمفاتيح الأساسية وحقول المفاتيح الأجنبية والعلاقات الداخلية. أما EmployeeId وOrgUnitId وأشباههما عندما تشير إلى خدمة أخرى فتظهر كمعرفات تكاملية، لأن المعمارية لا تنشئ مفاتيح أجنبية عابرة لقواعد الخدمات.")
    for item, figure_key in zip(manifest, DATABASE_FULL_ERD_KEYS):
        add_heading(doc, f"{item['service_ar']} ({item['service']}) - {item['database']}", 3)
        add_para(doc, f"السياق: {item['context']} | الكيانات والجداول: {item['entity_count']} | العلاقات الداخلية: {item['relationship_count']}", size=11.5, after=4)
        add_figure(doc, figure_key, max_width=9.55, max_height=5.9)
    doc.add_section(WD_SECTION.NEW_PAGE)
    portrait = doc.sections[-1]
    configure_section(portrait)
    continue_page_numbering(portrait)
    add_running_header(portrait)
    add_page_number_footer(portrait)


def parse_and_render_source(doc: Document) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    active = False
    current_chapter = 0
    current_heading = ""
    table_counts = {1: 0, 2: 0, 3: 0}
    table_title_offsets: dict[str, int] = {}
    paragraph_parts: list[str] = []
    code_lines: list[str] = []
    in_code = False
    code_language = ""
    skip_level: int | None = None
    formal_models_added = False
    database_section_added = False

    def flush_paragraph() -> None:
        nonlocal paragraph_parts
        if paragraph_parts:
            text = strip_markdown(" ".join(part.strip() for part in paragraph_parts if part.strip()))
            if text:
                add_para(doc, text)
            paragraph_parts = []

    def inject_after_chapter_one() -> None:
        add_evaluation_framework(doc)

    def inject_after_chapter_three() -> None:
        add_governance_section(doc)

    def normalize_heading_number(heading: str) -> str:
        return heading

    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()

        if line.startswith("# الفصل الأول"):
            active = True
            current_chapter = 1
            begin_chapter(doc, "الفصل الأول", "الإطار العام للمشروع", first=True)
            index += 1
            continue
        if line.startswith("# الفصل الثاني"):
            if active and current_chapter == 1:
                flush_paragraph()
                inject_after_chapter_one()
            active = True
            current_chapter = 2
            begin_chapter(doc, "الفصل الثاني", "الأساس النظري والدراسة المقارنة")
            index += 1
            continue
        if line.startswith("# الفصل الثالث"):
            if active and current_chapter == 2:
                flush_paragraph()
                add_related_work(doc)
            active = True
            current_chapter = 3
            begin_chapter(doc, "الفصل الثالث", "التحليل والتصميم والتنفيذ والتحقق")
            index += 1
            continue
        if line.startswith("# الخاتمة"):
            if active and current_chapter == 3:
                flush_paragraph()
                inject_after_chapter_three()
            begin_chapter(doc, "الخاتمة", "")
            active = True
            index += 1
            continue
        if line.startswith("# قائمة الاختصارات"):
            active = False
            index += 1
            continue
        if line.startswith("# المراجع"):
            break

        if not active:
            if line.startswith("# الخاتمة"):
                active = True
            else:
                index += 1
                continue

        heading_match = re.match(r"^(#{2,3})\s+(.+)$", line)
        if heading_match:
            level = len(heading_match.group(1))
            heading = strip_markdown(heading_match.group(2))
            if skip_level is not None:
                if level <= skip_level:
                    skip_level = None
                else:
                    index += 1
                    continue
            flush_paragraph()
            current_heading = line

            if heading.startswith("3.6 ") and not database_section_added:
                add_database_design_section(doc)
                database_section_added = True

            if heading.startswith("3.13 الاختبارات"):
                add_test_evidence(doc)
                skip_level = level
                index += 1
                continue
            if heading.startswith("3.14 النتائج"):
                add_results_section(doc)
                skip_level = level
                index += 1
                continue
            if heading.startswith("1.9 المتطلبات"):
                add_heading(doc, heading, 2)
                add_requirements_catalog(doc)
                skip_level = level
                index += 1
                continue
            if heading.startswith("3.5 تصميم قواعد البيانات") and not formal_models_added:
                add_formal_design_models(doc)
                formal_models_added = True

            add_heading(doc, normalize_heading_number(heading), 2 if level == 2 else 3)
            for prefix, note in CITATION_NOTES.items():
                if line.startswith(prefix):
                    add_para(doc, note, after=6)
                    break
            for prefix, figure_key in FIGURE_AFTER_HEADING.items():
                if line.startswith(prefix):
                    add_figure(doc, figure_key)
                    break
            index += 1
            continue

        if skip_level is not None:
            index += 1
            continue

        if line.startswith("```"):
            flush_paragraph()
            if not in_code:
                in_code = True
                code_language = line[3:].strip().lower()
                code_lines = []
            else:
                if code_language != "mermaid":
                    add_code(doc, "\n".join(code_lines))
                in_code = False
                code_lines = []
            index += 1
            continue

        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-{3,}", lines[index + 1].strip()):
            flush_paragraph()
            table_lines = [line]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = markdown_rows(table_lines)
            table_counts[current_chapter] += 1
            titles = TABLE_TITLES.get(current_heading.split(" ", 1)[0] + " " + current_heading.split(" ", 2)[1] if " " in current_heading else current_heading, [])
            table_title_offsets.setdefault(current_heading, 0)
            title_index = table_title_offsets[current_heading]
            table_title_offsets[current_heading] += 1
            if title_index < len(titles):
                title = titles[title_index]
            else:
                title = f"بيانات القسم {strip_markdown(current_heading.lstrip('#').strip())}"
            add_table(doc, rows, title, current_chapter, table_counts[current_chapter])
            continue

        number_match = re.match(r"^\d+\.\s+(.+)$", line)
        if number_match:
            flush_paragraph()
            number_match = re.match(r"^(\d+)\.\s+(.+)$", line)
            p = doc.add_paragraph()
            bidi(p, WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.right_indent = Cm(0.45)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, rtl_list_item(number_match.group(1), strip_markdown(number_match.group(2))), 13.0)
            index += 1
            continue

        if line.startswith("- "):
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            bidi(p, WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, strip_markdown(line[2:]), 13.0)
            index += 1
            continue

        if line.startswith("> "):
            flush_paragraph()
            add_callout(doc, "ملاحظة:", strip_markdown(line[2:]))
            index += 1
            continue

        if not line or line == "---":
            flush_paragraph()
            index += 1
            continue

        paragraph_parts.append(raw)
        index += 1

    flush_paragraph()


def add_back_matter(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "قائمة الاختصارات", 1)
    rows = [
        ["الاختصار", "المعنى"],
        ["API", "Application Programming Interface"],
        ["DDD", "Domain-Driven Design"],
        ["EF Core", "Entity Framework Core"],
        ["ESS", "Employee Self-Service"],
        ["HRMS", "Human Resources Management System"],
        ["JWT", "JSON Web Token"],
        ["OIDC", "OpenID Connect"],
        ["RBAC", "Role-Based Access Control"],
        ["RTL", "Right To Left"],
        ["YARP", "Yet Another Reverse Proxy"],
    ]
    add_table(doc, rows, "الاختصارات المستخدمة", 3, 7, label_override="الجدول (خ-1)")

    doc.add_page_break()
    add_heading(doc, "المراجع", 1)
    references = [
        "[1] E. Evans, Domain-Driven Design: Tackling Complexity in the Heart of Software. Boston, MA, USA: Addison-Wesley, 2003.",
        "[2] S. Newman, Building Microservices, 2nd ed. Sebastopol, CA, USA: O'Reilly Media, 2021.",
        "[3] J. Lewis and M. Fowler, \"Microservices: A definition of this new architectural term,\" martinfowler.com, Mar. 25, 2014. [Online]. Available: https://martinfowler.com/articles/microservices.html. [Accessed: Aug. 16, 2026].",
        "[4] Microsoft, \"ASP.NET Core documentation.\" [Online]. Available: https://learn.microsoft.com/aspnet/core/. [Accessed: Aug. 16, 2026].",
        "[5] Microsoft, \"Entity Framework Core documentation.\" [Online]. Available: https://learn.microsoft.com/ef/core/. [Accessed: Aug. 16, 2026].",
        "[6] Microsoft, \"YARP reverse proxy documentation.\" [Online]. Available: https://learn.microsoft.com/aspnet/core/fundamentals/servers/yarp/. [Accessed: Aug. 16, 2026].",
        "[7] MassTransit, \"MassTransit documentation: Consumers, retry, and transactional outbox.\" [Online]. Available: https://masstransit.io/documentation. [Accessed: Aug. 16, 2026].",
        "[8] Broadcom, \"RabbitMQ documentation.\" [Online]. Available: https://www.rabbitmq.com/docs. [Accessed: Aug. 16, 2026].",
        "[9] Keycloak, \"Keycloak documentation.\" [Online]. Available: https://www.keycloak.org/documentation. [Accessed: Aug. 16, 2026].",
        "[10] PostgreSQL Global Development Group, \"PostgreSQL documentation: Row security policies.\" [Online]. Available: https://www.postgresql.org/docs/current/ddl-rowsecurity.html. [Accessed: Aug. 16, 2026].",
        "[11] Docker, Inc., \"Docker Compose documentation.\" [Online]. Available: https://docs.docker.com/compose/. [Accessed: Aug. 16, 2026].",
        "[12] OpenTelemetry Authors, \"OpenTelemetry documentation.\" [Online]. Available: https://opentelemetry.io/docs/. [Accessed: Aug. 16, 2026].",
        "[13] Systems and Software Engineering - Systems and Software Quality Requirements and Evaluation (SQuaRE) - Product Quality Model, ISO/IEC 25010:2023, 2023.",
        "[14] V. C. Hu et al., Guide to Attribute Based Access Control (ABAC) Definition and Considerations, NIST SP 800-162, National Institute of Standards and Technology, 2014.",
        "[15] فريق مشروع HRManagement، \"تحليل المتطلبات ومصفوفة أحداث التكامل ونطاق الإصدار الأول،\" وثيقة داخلية، 2026.",
        "[16] فريق مشروع HRManagement، \"التصميم المعماري وتقرير التنفيذ والتحقق البرمجي،\" وثيقة داخلية، 2026.",
        "[17] كلية الهندسة المعلوماتية، جامعة حمص، \"منهجية إعداد مشاريع التخرج،\" دليل أكاديمي، 2026.",
        "[18] M. J. Kavanagh and R. D. Johnson, Human Resource Information Systems: Basics, Applications, and Future Directions, 4th ed. Thousand Oaks, CA, USA: SAGE, 2020.",
        "[19] J. H. Marler and J. W. Boudreau, \"An evidence-based review of HR analytics,\" The International Journal of Human Resource Management, vol. 28, no. 1, pp. 3-26, 2017.",
        "[20] D. L. Stone, D. L. Deadrick, K. M. Lukaszewski, and R. Johnson, \"The influence of technology on the future of human resource management,\" Human Resource Management Review, vol. 25, no. 2, pp. 216-231, 2015.",
        "[21] M. Richards and N. Ford, Fundamentals of Software Architecture. Sebastopol, CA, USA: O'Reilly Media, 2020.",
        "[22] M. Kleppmann, Designing Data-Intensive Applications. Sebastopol, CA, USA: O'Reilly Media, 2017.",
        "[23] Digital Identity Guidelines, NIST SP 800-63-3, National Institute of Standards and Technology, 2017.",
        "[24] Information Security, Cybersecurity and Privacy Protection - Information Security Management Systems - Requirements, ISO/IEC 27001:2022, 2022.",
        "[25] M. Swanson et al., Contingency Planning Guide for Federal Information Systems, NIST SP 800-34 Rev. 1, National Institute of Standards and Technology, 2010.",
        "[26] Human Resource Management - Guidelines for Internal and External Human Capital Reporting, ISO 30414:2018, 2018.",
        "[27] OWASP Foundation, \"Application Security Verification Standard 4.0.3,\" 2021. [Online]. Available: https://owasp.org/www-project-application-security-verification-standard/. [Accessed: Aug. 16, 2026].",
        "[28] OWASP Foundation, \"OWASP Top 10: The ten most critical web application security risks,\" 2021. [Online]. Available: https://owasp.org/www-project-top-ten/. [Accessed: Aug. 16, 2026].",
        "[29] L. Bass, P. Clements, and R. Kazman, Software Architecture in Practice, 4th ed. Boston, MA, USA: Addison-Wesley, 2021.",
        "[30] C. Richardson, Microservices Patterns. Shelter Island, NY, USA: Manning, 2018.",
        "[31] C. Pahl and P. Jamshidi, \"Microservices: A systematic mapping study,\" in Proc. CLOSER, Rome, Italy, 2016, pp. 137-146.",
        "[32] B. Beyer, C. Jones, J. Petoff, and N. R. Murphy, Site Reliability Engineering. Sebastopol, CA, USA: O'Reilly Media, 2016.",
        "[33] N. Forsgren, J. Humble, and G. Kim, Accelerate. Portland, OR, USA: IT Revolution Press, 2018.",
        "[34] Information Security, Cybersecurity and Privacy Protection - Information Security Controls, ISO/IEC 27002:2022, 2022.",
        "[35] D. Hardt, The OAuth 2.0 Authorization Framework, IETF RFC 6749, Oct. 2012.",
        "[36] N. Sakimura, J. Bradley, M. Jones, B. de Medeiros, and C. Mortimore, \"OpenID Connect Core 1.0,\" OpenID Foundation, Nov. 2014.",
        "[37] Security and Privacy Controls for Information Systems and Organizations, NIST SP 800-53 Rev. 5, National Institute of Standards and Technology, 2020.",
        "[38] فريق مشروع HRManagement، \"ADR-001: نموذج الجهة والوحدة التنظيمية وعزل Tenant،\" قرار معماري داخلي، 2026.",
        "[39] فريق مشروع HRManagement، \"دراسة اعتمادية قواعد البيانات بحسب رحلات الأعمال،\" دراسة داخلية، 2026.",
        "[40] فريق مشروع HRManagement، \"سجل السياسات وتوليد الرموز وحزم التهيئة المؤسسية،\" وثائق تصميم داخلية، 2026.",
        "[41] فريق مشروع HRManagement، \"جرد سياقات البيانات ومخططات الترحيل المؤرخة،\" دليل تصميم داخلي، 2026.",
        "[42] فريق مشروع HRManagement، \"خط أساس المتطلبات وسجل الحقيقة ومصفوفة حالات الاستخدام والتتبع،\" وثيقة داخلية، 2026.",
        "[43] SAP, \"SAP SuccessFactors Employee Central.\" [Online]. Available: https://www.sap.com/products/hcm/successfactors-employee-central.html. [Accessed: Aug. 16, 2026].",
        "[44] Oracle, \"Oracle Fusion Cloud Human Resources.\" [Online]. Available: https://www.oracle.com/human-capital-management/human-resources/. [Accessed: Aug. 16, 2026].",
        "[45] Workday, Inc., \"Human Capital Management software.\" [Online]. Available: https://www.workday.com/en-us/products/human-capital-management/overview.html. [Accessed: Aug. 16, 2026].",
        "[46] Odoo S.A., \"Odoo Employees features.\" [Online]. Available: https://www.odoo.com/app/employees-features. [Accessed: Aug. 16, 2026].",
        "[47] OrangeHRM, Inc., \"OrangeHRM human resource management software.\" [Online]. Available: https://orangehrm.com/. [Accessed: Aug. 16, 2026].",
        "[48] Oracle, \"Oracle Fusion Cloud Payroll.\" [Online]. Available: https://www.oracle.com/human-capital-management/payroll/. [Accessed: Aug. 16, 2026].",
    ]
    for ref in references:
        p = doc.add_paragraph()
        bidi(p, WD_ALIGN_PARAGRAPH.RIGHT)
        p.paragraph_format.right_indent = Cm(0.45)
        p.paragraph_format.space_after = Pt(5)
        add_inline(p, ref, 12.2)

    begin_landscape_section(doc)
    add_heading(doc, "الملاحق", 1)
    add_heading(doc, "ملحق (أ): مصفوفة تتبع المتطلبات", 2)
    fr_rows = [["المتطلب", "المجال", "الخدمة المالكة", "دليل القبول أو التحقق"]]
    for row in catalog_functional_rows():
        if len(row) >= 7:
            fr_rows.append([row[0], row[1], row[5], row[6]])
    add_table(doc, fr_rows, "مصفوفة التتبع الكاملة للمتطلبات الوظيفية", 3, 8, label_override="الجدول (أ-1)")

    add_capability_truth_register(doc)
    add_approved_usecase_package(doc, 1)
    add_dfd_appendix(doc)
    add_requirement_cards(doc)
    add_quality_cards(doc)
    add_service_profiles(doc)

    add_heading(doc, "ملحق (ز-2): ملخص الخدمات ومجالات الملكية", 3)
    services = [
        ["الخدمة", "الملكية الوظيفية"],
        ["HR.Core", "الجهة والموظف والوحدات والتوظيف والسياسات والرموز والتهيئة وESS"],
        ["HR.Contracts", "العقود والمباشرة والإجازات والسلف والفصل"],
        ["Payroll", "إصدارات القواعد واللقطات والدورات القابلة للاستئناف والتمويل والمتأخرات وأوامر الدفع والتسوية والحجوزات"],
        ["HR.Health", "السجلات واللجان والصحة المهنية"],
        ["HR.Movements", "النقل والانتداب"],
        ["HR.Training", "التدريب والتعلم الإلكتروني"],
        ["IDCard", "تصميم وإصدار وإبطال البطاقات"],
        ["Records.Archive", "الأرشفة الرقمية والفيزيائية"],
        ["Admin.Logistics", "العهد والأصول والمخزون والتوريد وبلاغات الفقد"],
        ["IT.Infra", "الأجهزة والدعم والنسخ الاحتياطي والحوادث"],
        ["Dev.Analytics", "التقارير ولوحات المؤشرات"],
        ["Compliance", "الشكاوى والتدقيق والإجراءات التأديبية"],
        ["Auth.Directory", "الدليل المحلي والأدوار وسجلات المصادقة"],
        ["File.Storage", "حجر الملفات والفحص والتنقية والمراجع ومنح الوصول وتدقيقه"],
    ]
    add_table(doc, services, "دليل الخدمات ومجالات ملكية البيانات", 3, 9, label_override="الجدول (ز-1)")

    add_complete_database_erdds(doc)

    add_verification_protocol(doc, new_page=False)

    add_heading(doc, "ملحق (ط-2): إعادة تشغيل اختبارات التحقق", 3)
    add_para(doc, "تعاد حزم الوحدة والتكامل والنهاية إلى النهاية في وضع Release، وتسجل نتائج النجاح والفشل والتخطي كل على حدة. سجل التشغيل الحالي 2416 اختبار وحدة ناجحاً و5 فاشلة، و158 اختبار تكامل ناجحاً و7 فاشلة و47 متخطاة، واختبارين E2E ناجحين و55 متخطاة.")
    add_code(doc, "dotnet test src/Tests/HRManagement.UnitTests/HRManagement.UnitTests.csproj --no-restore --configuration Release --nologo\ndotnet test src/Tests/HRManagement.IntegrationTests/HRManagement.IntegrationTests.csproj --no-restore --configuration Release --nologo\ndotnet test src/Tests/HRManagement.E2ETests/HRManagement.E2ETests.csproj --no-restore --configuration Release --nologo")
    add_para(doc, "تحتاج Testcontainers واختبارات PostgreSQL RLS إلى Docker daemon. لا تحول الاختبارات المتخطاة إلى ناجحة، ولا تستخدم نتائج نسخة سابقة مكان الخط الحالي.")

    add_heading(doc, "ملحق (ي): مطابقة بنية الأطروحة مع المنهجية المرجعية", 2)
    methodology = [
        ["بند المنهجية", "موضعه في الأطروحة"],
        ["الغلاف العربي والإنكليزي", "الصفحات الأولى"],
        ["الشكر والملخصان", "التمهيد"],
        ["فهرس المحتويات وقوائم الأشكال والجداول", "التمهيد"],
        ["الفصل الأول", "المقدمة والمشكلة والأهداف والمنهجية"],
        ["الفصل الثاني", "الدراسة النظرية والتوثيق المرجعي"],
        ["الفصل الثالث", "التطبيق العملي والاختبارات والنتائج"],
        ["الخاتمة والاختصارات", "بعد الفصل الثالث"],
        ["المراجع والملاحق", "نهاية الوثيقة"],
    ]
    add_table(doc, methodology, "مطابقة البنية مع منهجية إعداد مشاريع التخرج", 3, 10, label_override="الجدول (ي-1)")
    add_submission_integrity_appendix(doc)
    add_execution_log_appendix(doc)


def set_update_fields(doc: Document) -> None:
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Source thesis was not found: {SOURCE}")
    TABLE_COUNTERS.update({1: 0, 2: 0, 3: 0})
    FIGURE_COUNTERS.clear()
    make_figures()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_section(doc.sections[0])
    apply_styles(doc)
    doc.core_properties.title = "HRManagement Graduation Thesis"
    doc.core_properties.subject = "Human resources platform for central and branch administrations"
    doc.core_properties.author = "HRManagement Graduation Team"
    doc.core_properties.keywords = "HRManagement, human resources, microservices, branches"
    add_front_matter(doc)
    parse_and_render_source(doc)
    add_back_matter(doc)
    materialize_figure_catalog(doc)
    materialize_table_catalog(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
